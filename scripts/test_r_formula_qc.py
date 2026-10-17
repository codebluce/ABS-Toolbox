"""Independent R gate adversarial tests: ZIP caches and source identities, no aggregate."""
from datetime import datetime
import hashlib
from pathlib import Path
import tempfile
import unittest
from unittest import mock
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET

from openpyxl import Workbook
import latest_bookkeeping_qc as latest
from lab import monthly_bookkeeping_qc as monthly
import dashboard_provenance as provenance

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
HEADERS = [f'字段{i}' for i in range(1, 26)]
for col, name in {4:'资产类型',5:'项目名称',6:'发行场所',10:'规模',12:'簿记时间',
                  15:'期限',16:'分层情况',17:'比例',18:'对应金额（亿）',19:'计划管理人',
                  20:'成本',21:'认购机构',22:'认购份额',23:'申购利率',24:'穿透机构',25:'申购规模'}.items():
    HEADERS[col-1] = name


def row(project='项目一', tier='优先A', q=.81):
    result = [None] * 25
    for col, value in {4:'赊销白条',5:project,6:'上交所',10:10,12:datetime(2026,10,16),
                       15:'12',16:tier,17:q,18:10*q,19:'管理人',20:.02,
                       21:'配售机构',22:1,23:.02,24:'申购机构',25:1}.items():
        result[col-1] = value
    return result


class RFormulaQCTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / '2026年ABS发行台账-1016-定稿.xlsx'

    def source(self, formula='=J2*Q2', cache=8.1, rows=None, merges=(), operands=None):
        book = Workbook(); ws = book.active; ws.title = '总表'
        ws.append(HEADERS)
        for values in rows or [row(), row()]:
            ws.append(values)
        ws['R2'] = formula
        for coord, value in (operands or {}).items():
            ws[coord] = value
        for area in merges:
            ws.merge_cells(area)
        book.save(self.path); book.close()
        if cache is not None:
            with ZipFile(self.path) as archive:
                entries = {name:archive.read(name) for name in archive.namelist()}
            xml = ET.fromstring(entries['xl/worksheets/sheet1.xml'])
            cell = xml.find(".//m:c[@r='R2']", NS)
            value = cell.find('m:v', NS)
            if value is None:
                value = ET.SubElement(cell, '{'+NS['m']+'}v')
            value.text = str(cache)
            entries['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
            with ZipFile(self.path, 'w', ZIP_DEFLATED) as archive:
                for name, data in entries.items():
                    archive.writestr(name, data)
        return self.path

    def pass_readers(self):
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        for reader in (latest.read_source, monthly.read_source):
            with self.subTest(reader=reader.__module__):
                records = reader(self.path)
                self.assertEqual(records[0]['对应金额（亿）'], 8.1)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(), before)

    def fail_readers(self, code):
        for reader in (latest.read_source, monthly.read_source):
            with self.subTest(reader=reader.__module__):
                with self.assertRaisesRegex(ValueError, code):
                    reader(self.path)
        report = monthly.verify_source(self.path, raise_on_fail=False)
        self.assertEqual(report['status'], 'FAIL')
        self.assertRegex(str(report['issues']), code)
        # No generated data needed to prove source guard hard-blocks.
        report = latest.check_data({}, self.path, raise_on_fail=False)
        self.assertEqual(report['status'], 'FAIL')
        self.assertRegex(str(report['issues']), code)

    def test_wrong_project_equal_cache_or_missing_cache_fails(self):
        for formula in ('=J3*Q2', '=J2*Q3', '=$J$3*$Q$3'):
            for cache in (8.1, None):
                self.source(formula, cache, [row(), row(project='其他项目')])
                self.fail_readers('R_FORMULA_REFERENCE')

    def test_wrong_tier_equal_cache_or_missing_cache_fails(self):
        for cache in (8.1, None):
            self.source('=J2*Q3', cache, [row(), row(tier='次级')])
            self.fail_readers('R_FORMULA_REFERENCE')

    def test_same_project_any_j_same_tier_q_and_absolute_relative(self):
        for formula in ('=J3*Q2', '=$J$3*Q$2', '=Q2*$J3', '= j3 * q2 '):
            self.source(formula, rows=[row(), row(tier='次级')])
            self.pass_readers()

    def test_valid_same_tier_q_anchor(self):
        self.source('=$J$3*$Q$3')
        self.pass_readers()

    def test_constant_valid_even_when_not_jq(self):
        self.source('=8.1', rows=[row(q=.2), row(q=.2)])
        self.pass_readers()

    def test_constant_cache_mismatch(self):
        self.source('=8.1', cache=8.2)
        self.fail_readers('R_FORMULA_CACHE')

    def test_direct_cache_mismatch_and_tolerance(self):
        self.source(cache=8.10011)
        self.fail_readers('R_FORMULA_CACHE')
        self.source(cache=8.1001)
        for reader in (latest.read_source, monthly.read_source):
            self.assertEqual(reader(self.path)[0]['对应金额（亿）'], 8.1001)

    def test_zero_q_is_valid(self):
        self.source(cache=0, rows=[row(q=0), row(q=0)])
        for reader in (latest.read_source, monthly.read_source):
            self.assertEqual(reader(self.path)[0]['对应金额（亿）'], 0)

    def test_missing_valid_cache_warn_without_fallback(self):
        for formula in ('=J2*Q2', '=8.1'):
            self.source(formula, cache=None)
            for reader in (latest.read_source, monthly.read_source):
                records = reader(self.path)
                self.assertIsNone(records[0]['对应金额（亿）'])
                self.assertEqual(records[0]['_uncachedLayerSize'], 'R2')
            report = monthly.verify_source(self.path, raise_on_fail=False)
            self.assertEqual(report['status'], 'PASS')
            self.assertIn('MONTHLY_LAYER_SIZE_UNCACHED', str(report['issues']))
            self.assertEqual(report['checks']['rFormulas']['uncached'], 1)

    def test_complex_cross_sheet_range_and_wrong_columns_block(self):
        for formula in ("='总表'!J2*Q2", '=SUM(J2:J3)*Q2', '=J2*Q2+0', '=J2*J3'):
            self.source(formula)
            self.fail_readers('R_FORMULA_(UNSUPPORTED|REFERENCE)')

    def test_numeric_operands_and_identity_required(self):
        for operands in ({'J2':'10'}, {'Q2':'0.81'}, {'Q2':True}, {'Q2':1.1},
                         {'J2':-1}, {'Q2':None}, {'J2':'=10'}, {'Q2':'=.81'}):
            self.source(operands=operands)
            # Key-field formula J can also be rejected by original cache guard.
            for reader in (latest.read_source, monthly.read_source):
                with self.assertRaises(ValueError):
                    reader(self.path)
        self.source(operands={'E2':None})
        self.fail_readers('R_FORMULA_IDENTITY')

    def test_explicit_merge_metadata_expanded_before_identity_check(self):
        self.source('=J3*Q3', merges=('D2:D3','E2:E3','F2:F3','J2:J3',
                                     'L2:L3','P2:P3','Q2:Q3'))
        self.pass_readers()
        self.source('=J3*Q3', merges=('D2:D3','E2:E3','F2:F3','J2:J3','L2:L3','Q2:Q3'),
                    rows=[row(), row(tier='次级')])
        self.fail_readers('R_FORMULA_REFERENCE')

    def test_independent_of_production_helpers_and_aggregates(self):
        self.source()
        with mock.patch('ledger_formula_tools.validate_r_formula', side_effect=AssertionError('production proof')), \
             mock.patch('latest_bookkeeping_panel.compute_data', side_effect=AssertionError('production aggregate')):
            self.pass_readers()

    def test_fingerprints_cover_helper_and_repair_tool(self):
        self.assertIn('ledger_formula_tools.py', provenance.PRODUCERS)
        self.assertNotIn('repair_r_formulas.py', provenance.PRODUCERS)
        dependencies = monthly.dependency_fingerprints()
        for name in ('ledger_formula_tools.py', 'repair_r_formulas.py'):
            self.assertIn('scripts/'+name, dependencies)
        original = latest.rules_sha256()
        read = Path.read_bytes
        def changed(path):
            data = read(path)
            return data+b'changed' if path.name == 'ledger_formula_tools.py' else data
        with mock.patch.object(Path, 'read_bytes', changed):
            self.assertNotEqual(original, latest.rules_sha256())


if __name__ == '__main__':
    unittest.main()
