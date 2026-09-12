#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""增量合并·既有项目字段同步 单元测试

覆盖 v2.6.0 新增的 Step 2.5：以最新原始台账为准回填既有项目的项目级字段。
核心不变量：
  1. 项目级字段（月份/场所/管理人/规模/成本等）按原始表回填到该项目**每一行**；
  2. WXY(U/V/W/X/Y) 与分层层字段(16-20) 绝不被触碰；
  3. 原始表未出现的项目保持原样；
  4. 无差异时不产生变更记录。

运行：PYTHONUTF8=1 .venv/bin/python -m unittest scripts.test_increment_merge_field_sync -v
"""
import datetime
import os
import sys
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from increment_merge import (  # noqa: E402
    PROJECT_LEVEL_COLS, _field_equal, sync_existing_project_fields, get_all_projects,
)

# 1-indexed 列位（与 PROJECT_LEVEL_COLS 一致）
C_MONTH, C_ASSET, C_VENUE, C_MGR, C_UW = 2, 4, 6, 7, 8
C_CUST, C_SIZE, C_TERM, C_DATE, C_CD, C_COST, C_MULT = 9, 10, 11, 12, 13, 14, 15
C_NAME = 5
C_LAYER, C_LAYER_PCT, C_LAYER_AMT, C_RATING, C_LAYER_COST = 16, 17, 18, 19, 20
C_U, C_V, C_W, C_X, C_Y = 21, 22, 23, 24, 25

D1 = datetime.datetime(2026, 8, 10)
D2 = datetime.datetime(2026, 9, 7)


def make_sheet(rows):
    """rows: list[dict{col: value}]，索引即 1-based 列位；前两行为表头。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in (1, 2):
        for c in range(1, 26):
            ws.cell(row=r, column=c).value = f'h{c}'
    for i, row in enumerate(rows, start=3):
        for c, v in row.items():
            ws.cell(row=i, column=c).value = v
    return ws


def proj_rows(name, n_rows, overrides=None):
    """构造一个项目的 n 行，项目级字段在每行重复（与真实台账一致）。

    overrides 为 列位 -> 值 的 dict（列位是 int，不能用关键字参数）。
    """
    base = {
        C_MONTH: '8月', C_ASSET: '保理', C_NAME: name, C_VENUE: '上交所',
        C_MGR: '财通资管', C_UW: '平安证券', C_CUST: '光大北分', C_SIZE: 10,
        C_TERM: '3+3', C_DATE: D1, C_CD: 0.01485, C_COST: 0.017, C_MULT: 2.12,
    }
    base.update(overrides or {})
    rows = []
    for i in range(n_rows):
        r = dict(base)
        r[C_LAYER] = '优先级'
        r[C_LAYER_PCT] = 0.99
        r[C_LAYER_AMT] = 9.9
        r[C_RATING] = 'AAA'
        r[C_LAYER_COST] = 0.0155
        r[C_U] = f'机构{i}' if i % 2 == 0 else None
        r[C_V] = 1.0
        r[C_W] = 0.015
        r[C_X] = f'机构{i}'
        r[C_Y] = 2.0
        rows.append(r)
    return rows


def snapshot(ws, rows_range):
    return [tuple(ws.cell(row=r, column=c).value for c in range(1, 26))
            for r in rows_range]


class TestFieldEqual(unittest.TestCase):
    def test_none_and_blank_string_are_equal(self):
        self.assertTrue(_field_equal(None, None))
        self.assertTrue(_field_equal('', None))
        self.assertTrue(_field_equal('   ', None))

    def test_string_whitespace_insensitive(self):
        self.assertTrue(_field_equal('上交所 ', '上交所'))
        self.assertFalse(_field_equal('上交所', '深交所'))

    def test_datetime_and_numeric(self):
        self.assertTrue(_field_equal(D1, D1))
        self.assertFalse(_field_equal(D1, D2))
        self.assertTrue(_field_equal(0.017, 0.0170000000001))
        self.assertFalse(_field_equal(0.017, 0.0171))

    def test_numeric_vs_string_not_equal(self):
        self.assertFalse(_field_equal('10', 10))


class TestSyncExistingProjectFields(unittest.TestCase):
    def setUp(self):
        # 已加工台账（含 WXY）：项目A 3行、项目B 2行、仅A表有项目C
        ws_rows = (proj_rows('项目A', 3) + proj_rows('项目B', 2)
                   + proj_rows('仅加工表有', 1))
        self.ws_a = make_sheet(ws_rows)
        # 原始台账：项目A 修订2处、项目B 无差异
        raw_rows = (proj_rows('项目A', 2, {C_VENUE: '深交所', C_COST: 0.019})
                    + proj_rows('项目B', 2))
        self.ws_b = make_sheet(raw_rows)

    def test_detects_only_real_changes(self):
        changes, touched, _protected = sync_existing_project_fields(
            self.ws_a, self.ws_b,
            get_all_projects(self.ws_a), get_all_projects(self.ws_b))
        self.assertEqual(
            sorted((c['project'], c['field']) for c in changes),
            [('项目A', 'ALL-IN成本'), ('项目A', '发行场所')])
        self.assertEqual(touched, {'项目A'})

    def test_backfills_every_row_of_group(self):
        sync_existing_project_fields(
            self.ws_a, self.ws_b,
            get_all_projects(self.ws_a), get_all_projects(self.ws_b))
        for r in (3, 4, 5):  # 项目A 的 3 行
            self.assertEqual(self.ws_a.cell(row=r, column=C_VENUE).value, '深交所')
            self.assertEqual(self.ws_a.cell(row=r, column=C_COST).value, 0.019)

    def test_wxy_and_layer_columns_untouched(self):
        pa = get_all_projects(self.ws_a)['项目A']
        before = snapshot(self.ws_a, range(pa['start'], pa['end'] + 1))
        sync_existing_project_fields(
            self.ws_a, self.ws_b,
            get_all_projects(self.ws_a), get_all_projects(self.ws_b))
        after = snapshot(self.ws_a, range(pa['start'], pa['end'] + 1))
        untouched = [C_LAYER, C_LAYER_PCT, C_LAYER_AMT, C_RATING, C_LAYER_COST,
                     C_U, C_V, C_W, C_X, C_Y]
        for old_row, new_row in zip(before, after):
            for c in untouched:
                self.assertEqual(old_row[c - 1], new_row[c - 1],
                                 f'列{c} 不应被字段同步改动')

    def test_identical_project_has_no_change(self):
        changes, _touched, _protected = sync_existing_project_fields(
            self.ws_a, self.ws_b,
            get_all_projects(self.ws_a), get_all_projects(self.ws_b))
        self.assertNotIn('项目B', {c['project'] for c in changes})

    def test_project_absent_from_raw_untouched(self):
        pa = get_all_projects(self.ws_a)['仅加工表有']
        before = snapshot(self.ws_a, range(pa['start'], pa['end'] + 1))
        sync_existing_project_fields(
            self.ws_a, self.ws_b,
            get_all_projects(self.ws_a), get_all_projects(self.ws_b))
        self.assertEqual(before, snapshot(self.ws_a, range(pa['start'], pa['end'] + 1)))

    def test_change_record_shape(self):
        changes, _touched, _protected = sync_existing_project_fields(
            self.ws_a, self.ws_b,
            get_all_projects(self.ws_a), get_all_projects(self.ws_b))
        venue = next(c for c in changes if c['field'] == '发行场所')
        self.assertEqual(venue['old'], '上交所')
        self.assertEqual(venue['new'], '深交所')
        self.assertEqual(venue['rows'], '3-5')

    def test_datetime_field_synced_and_serialized(self):
        ws_rows = proj_rows('项目D', 2) + []
        ws_a = make_sheet(ws_rows)
        ws_b = make_sheet(proj_rows('项目D', 1, {C_DATE: D2}))
        changes, _touched, _protected = sync_existing_project_fields(
            ws_a, ws_b, get_all_projects(ws_a), get_all_projects(ws_b))
        self.assertEqual([c['field'] for c in changes], ['簿记时间'])
        self.assertEqual(changes[0]['old'], '2026-08-10')
        self.assertEqual(changes[0]['new'], '2026-09-07')
        for r in (3, 4):
            self.assertEqual(ws_a.cell(row=r, column=C_DATE).value, D2)

    def test_project_level_cols_do_not_include_wxy_or_layer(self):
        for forbidden in (1, C_NAME, C_LAYER, C_LAYER_PCT, C_LAYER_AMT,
                          C_RATING, C_LAYER_COST, C_U, C_V, C_W, C_X, C_Y):
            self.assertNotIn(forbidden, PROJECT_LEVEL_COLS,
                             f'列{forbidden} 不应纳入字段同步范围')


class TestKnownYearTypoProtected(unittest.TestCase):
    """源表把 2026 簿记年份误录为 2025 时，不得覆盖台账中已纠错的值。

    金采7-12 即此类：原始台账三期均为 2025-05-12，而定稿存的是纠错后的
    2026-05-12。若无条件回填，会把纠错结果退回错误年份。
    """

    def _sheets(self, raw_date, processed_date):
        ws_a = make_sheet(proj_rows('金采7-12', 2, {C_DATE: processed_date}))
        ws_b = make_sheet(proj_rows('金采7-12', 2, {C_DATE: raw_date}))
        return ws_a, ws_b

    def test_2026_to_2025_typo_not_backfilled(self):
        keep = datetime.datetime(2026, 5, 12)
        typo = datetime.datetime(2025, 5, 12)
        ws_a, ws_b = self._sheets(typo, keep)
        changes, _touched, protected = sync_existing_project_fields(
            ws_a, ws_b, get_all_projects(ws_a), get_all_projects(ws_b))
        self.assertEqual(changes, [])
        self.assertEqual(len(protected), 1)
        self.assertEqual(protected[0]['kept'], '2026-05-12')
        self.assertEqual(protected[0]['raw'], '2025-05-12')
        # 台账保持纠错后的年份
        for r in (3, 4):
            self.assertEqual(ws_a.cell(row=r, column=C_DATE).value, keep)

    def test_different_month_day_still_syncs(self):
        """月-日不同则属真实修订，应正常回填（年份问题交由下游纠错）。"""
        ws_a, ws_b = self._sheets(datetime.datetime(2025, 6, 15),
                                  datetime.datetime(2026, 5, 12))
        changes, _touched, protected = sync_existing_project_fields(
            ws_a, ws_b, get_all_projects(ws_a), get_all_projects(ws_b))
        self.assertEqual(len(changes), 1)
        self.assertEqual(changes[0]['field'], '簿记时间')
        self.assertEqual(protected, [])

    def test_reverse_direction_syncs_normally(self):
        """台账 2025、源表 2026：方向与已知笔误相反，应正常回填。"""
        ws_a, ws_b = self._sheets(datetime.datetime(2026, 5, 12),
                                  datetime.datetime(2025, 5, 12))
        changes, _touched, protected = sync_existing_project_fields(
            ws_a, ws_b, get_all_projects(ws_a), get_all_projects(ws_b))
        self.assertEqual(len(changes), 1)
        self.assertEqual(changes[0]['new'], '2026-05-12')
        self.assertEqual(protected, [])

    def test_non_date_field_not_treated_as_typo(self):
        """保护仅针对簿记时间列，其他字段正常回填。"""
        ws_a = make_sheet(proj_rows('项目E', 1, {C_VENUE: '上交所', C_DATE: D1}))
        ws_b = make_sheet(proj_rows('项目E', 1, {C_VENUE: '深交所', C_DATE: D1}))
        changes, _touched, protected = sync_existing_project_fields(
            ws_a, ws_b, get_all_projects(ws_a), get_all_projects(ws_b))
        self.assertEqual([c['field'] for c in changes], ['发行场所'])
        self.assertEqual(protected, [])


if __name__ == '__main__':
    unittest.main(verbosity=2)
