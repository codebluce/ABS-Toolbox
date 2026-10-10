"""门禁对抗回归：坏数据/陈旧产物必须失败，失败保留上一版。"""
import copy
from pathlib import Path
import tempfile
import re
import unittest
from unittest import mock

from openpyxl import Workbook

import latest_bookkeeping_panel as panel
import latest_bookkeeping_qc as qc
import update_latest_bookkeeping as updater

HEADERS = ['资产类型','项目名称','簿记时间','发行场所','分层情况','成本','认购机构',
           '认购份额','申购利率','穿透机构','申购规模','规模','计划管理人','期限']
BASE = ['赊销白条','项目一','2026-09-21','上交所','优先A',0.0177,'配售机构',1,
        0.0176,'申购机构',2,1,'管理人','21+3']


class GateTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name)/'2026年ABS发行台账-0924-定稿.xlsx'

    def write_source(self, rows=None, merges=(), active_other=False):
        book = Workbook();sheet=book.active;sheet.title='总表'
        sheet.append(HEADERS)
        for row in rows or [BASE]: sheet.append(row)
        for area in merges:sheet.merge_cells(area)
        if active_other:book.create_sheet('其他').append(['不是台账']);book.active=1
        book.save(self.path);book.close()

    def test_real_fields_and_active_sheet_not_used(self):
        self.write_source(active_other=True)
        data=panel.compute_data(self.path)
        self.assertEqual(data['qc']['status'],'PASS')
        self.assertEqual(data['assets'][0]['projects'][0]['tiers'][0]['allocations'][0]['cost'],0.0177)

    def test_missing_cost_never_inherits_previous_layer(self):
        other=BASE.copy();other[4]='次级';other[5]=None;other[7]=0.2
        first=BASE.copy();first[7]=0.8
        self.write_source([first,other])
        data=panel.compute_data(self.path)
        tier=data['assets'][0]['projects'][0]['tiers'][1]
        self.assertEqual(tier['costs'],[])
        self.assertEqual(tier['allocations'][0]['cost'],'')

    def test_each_allocation_keeps_its_own_cost(self):
        other=BASE.copy();other[5]=0.02;other[6]='另一机构';other[7]=0.5
        first=BASE.copy();first[11]=1.5;other[11]=1.5
        self.write_source([first,other])
        data=panel.compute_data(self.path)
        values=[r['cost'] for r in data['assets'][0]['projects'][0]['tiers'][0]['allocations']]
        self.assertEqual(values,[0.0177,0.02])
        self.assertIn('MULTIPLE_COSTS',[i['code'] for i in data['qc']['issues']])

    def test_bad_latest_date_is_not_silently_dropped(self):
        other=BASE.copy();other[1]='新项目';other[2]='无效日期'
        self.write_source([BASE,other])
        with self.assertRaises(qc.BookkeepingGateError):panel.compute_data(self.path)

    def test_orphan_bid_negative_amount_and_rate_units_block(self):
        for index,value,code in [(9,None,'ORPHAN_BID'),(7,-1,'NEGATIVE_AMOUNT'),(8,1.77,'RATE_UNIT')]:
            with self.subTest(code=code):
                changed=BASE.copy();changed[index]=value;self.write_source([changed])
                with self.assertRaises(qc.BookkeepingGateError) as ctx:panel.compute_data(self.path)
                self.assertIn(code,[i['code'] for i in ctx.exception.report['issues']])

    def test_merged_project_fields_expand_but_investors_block(self):
        other=BASE.copy();other[6]='另一机构'
        first=BASE.copy();first[7]=0.5;other[7]=0.5
        self.write_source([first,other],merges=('A2:A3','B2:B3','C2:C3','D2:D3','E2:E3','F2:F3'))
        self.assertEqual(panel.compute_data(self.path)['qc']['status'],'PASS')
        self.write_source([BASE,other],merges=('G2:G3',))
        with self.assertRaises(qc.BookkeepingGateError):panel.compute_data(self.path)

    def test_tampering_latest_scope_details_and_hash_block(self):
        self.write_source();data=panel.compute_data(self.path)
        for mutation in ('size','detail','hash','rule','date'):
            wrong=copy.deepcopy(data)
            project=wrong['assets'][0]['projects'][0]
            if mutation=='size':project['size']=100
            if mutation=='detail':project['tiers'][0]['bids'][0]['bidSize']=200
            if mutation=='hash':wrong['provenance']['sourceSha256']='old'
            if mutation=='rule':wrong['provenance']['rulesSha256']='old'
            if mutation=='date':wrong['assets'][0]['date']='2026-09-01'
            self.assertEqual(qc.check_data(wrong,self.path,raise_on_fail=False)['status'],'FAIL')

    def test_source_edit_with_same_filename_stales_html(self):
        self.write_source();html=panel.render_body(panel.compute_data(self.path))
        changed=BASE.copy();changed[10]=3;self.write_source([changed])
        with self.assertRaises(qc.BookkeepingGateError):qc.verify_html(html,self.path)

    def test_update_preserves_other_modules_and_is_idempotent(self):
        self.write_source();data=panel.compute_data(self.path)
        html='<html><head><style>/*existing*/</style></head><body><div class="sub-tabs-pane" data-module="pricing"></div><div class="panel-container"><div id="other">保留其他面板</div></div></body></html>'
        updated=updater.replace_panel(html,data)
        self.assertIn('<div id="other">保留其他面板</div>',updated)
        self.assertEqual(qc.verify_html(updated,self.path)['status'],'PASS')
        self.assertEqual(updater.replace_panel(updated,data),updated)

    def test_bad_html_never_overwrites_previous_and_cleans_temp(self):
        self.write_source();target=Path(self.temp.name)/'previous.html';target.write_text('上一版')
        with self.assertRaises(ValueError):updater.atomic_write(target,'坏面板',self.path)
        self.assertEqual(target.read_text(),'上一版')
        self.assertEqual(list(target.parent.glob('.bookkeeping-*')),[])

    def test_oversized_placement_fails_but_overbidding_allowed(self):
        changed=BASE.copy();changed[7]=100000;self.write_source([changed])
        with self.assertRaises(qc.BookkeepingGateError) as ctx:panel.compute_data(self.path)
        self.assertIn('ALLOCATION_PROJECT_OVER',[i['code'] for i in ctx.exception.report['issues']])
        changed=BASE.copy();changed[10]=100000;self.write_source([changed])
        self.assertEqual(panel.compute_data(self.path)['qc']['status'],'PASS')

    def test_uncached_formulas_blocked(self):
        for column,value in [(5,'=0.0177'),(7,'=1'),(10,'=2')]:
            with self.subTest(column=column):
                changed=BASE.copy();changed[column]=value;self.write_source([changed])
                with self.assertRaises(qc.BookkeepingGateError) as ctx:panel.compute_data(self.path)
                self.assertIn('计算缓存',str(ctx.exception))

    def test_natural_day_selection_retains_both_times(self):
        morning=BASE.copy();morning[2]='2026-09-21 09:00:00'
        afternoon=BASE.copy();afternoon[1]='下午项目';afternoon[2]='2026-09-21 15:00:00'
        self.write_source([morning,afternoon])
        data=panel.compute_data(self.path)
        self.assertEqual([p['name'] for p in data['assets'][0]['projects']],['项目一','下午项目'])
        self.assertEqual(data['qc']['status'],'PASS')

    def test_bool_rates_and_local_header_block(self):
        for column,value in [(5,True),(8,False),(5,'成本')]:
            changed=BASE.copy();changed[column]=value;self.write_source([changed])
            with self.assertRaises(qc.BookkeepingGateError):panel.compute_data(self.path)

    def test_real_runtime_verifies_actual_rows_and_exports(self):
        from latest_bookkeeping_runtime import verify_runtime
        self.write_source()
        result=verify_runtime(panel.compute_data(self.path))
        self.assertEqual(result['status'],'PASS')
        self.assertEqual(result['projects'],1)
        self.assertEqual(result['details'],2)
        self.assertEqual(result['exports'],2)

    def test_layer_oversized_and_optional_formula_warn(self):
        book=Workbook();sheet=book.active;sheet.title='总表'
        sheet.append(HEADERS+['对应金额（亿）']);sheet.append(BASE+[0.5]);book.save(self.path);book.close()
        with self.assertRaises(qc.BookkeepingGateError) as ctx:panel.compute_data(self.path)
        self.assertIn('ALLOCATION_LAYER_OVER',[i['code'] for i in ctx.exception.report['issues']])
        book=Workbook();sheet=book.active;sheet.title='总表'
        sheet.append(HEADERS+['对应金额（亿）']);sheet.append(BASE+['=1']);book.save(self.path);book.close()
        data=panel.compute_data(self.path)
        self.assertEqual(data['qc']['status'],'PASS')
        self.assertIn('LAYER_SIZE_UNCACHED',[i['code'] for i in data['qc']['issues']])

    def test_controlled_pending_final_failure_preserves_target(self):
        bad=BASE.copy();bad[7]=1000;self.write_source([bad])
        with mock.patch.object(updater,'ROOT',Path(self.temp.name)):
            target=Path(self.temp.name)/'deliverables/ledger/03_final/2026年ABS发行台账-0924-定稿.xlsx'
            target.parent.mkdir(parents=True);target.write_text('上一版台账')
            with self.assertRaises(qc.BookkeepingGateError):updater.validate_pending_final(self.path,target)
            self.assertEqual(target.read_text(),'上一版台账')

    def test_complete_local_header_skipped_and_bid_only_header_blocked(self):
        header=BASE.copy()
        for index, label in [(5,'成本'),(6,'认购机构'),(7,'认购份额'),(8,'申购利率'),(9,'穿透机构'),(10,'申购规模')]:header[index]=label
        self.write_source([BASE,header])
        data=panel.compute_data(self.path)
        tier=data['assets'][0]['projects'][0]['tiers'][0]
        self.assertEqual(tier['costs'],[0.0177])
        self.assertEqual(len(tier['allocations']),1)
        self.assertEqual(len(tier['bids']),1)
        bad=BASE.copy();bad[6]=None;bad[7]=None;bad[8]='申购利率'
        self.write_source([bad])
        with self.assertRaises(qc.BookkeepingGateError) as ctx:panel.compute_data(self.path)
        self.assertIn('LOCAL_HEADER',[i['code'] for i in ctx.exception.report['issues']])

    def test_project_tenor_preserves_x_plus_y_and_zero(self):
        from latest_bookkeeping_runtime import verify_runtime
        for term in ['21+3', '0+47']:
            with self.subTest(term=term):
                changed=BASE.copy();changed[-1]=term;self.write_source([changed])
                data=panel.compute_data(self.path)
                self.assertEqual(data['assets'][0]['projects'][0]['tenor'],term)
                self.assertEqual(verify_runtime(data)['status'],'PASS')

    def test_project_tenor_conflict_or_tampering_blocks(self):
        first=BASE.copy();first[7]=0.5
        second=first.copy();second[-1]='22+3'
        self.write_source([first,second])
        with self.assertRaises(qc.BookkeepingGateError) as ctx:panel.compute_data(self.path)
        self.assertIn('期限',str(ctx.exception))
        self.write_source();data=panel.compute_data(self.path)
        data['assets'][0]['projects'][0]['tenor']='22+3'
        self.assertEqual(qc.check_data(data,self.path,raise_on_fail=False)['status'],'FAIL')

    def test_project_tenor_missing_warns_and_formula_without_cache_blocks(self):
        from latest_bookkeeping_runtime import verify_runtime
        changed=BASE.copy();changed[-1]=None;self.write_source([changed])
        data=panel.compute_data(self.path)
        self.assertEqual(data['assets'][0]['projects'][0]['tenor'],'')
        self.assertIn('TENOR_MISSING',[i['code'] for i in data['qc']['issues']])
        self.assertEqual(verify_runtime(data)['status'],'PASS')
        changed[-1]='=21+3';self.write_source([changed])
        with self.assertRaises(qc.BookkeepingGateError) as ctx:panel.compute_data(self.path)
        self.assertIn('计算缓存',str(ctx.exception))

    def test_three_period_gate_preserves_each_period_and_rejects_omission(self):
        from latest_bookkeeping_runtime import verify_runtime
        rows=[]
        for stamp,name in [('2026-09-21','最新甲'),('2026-09-21','最新乙'),
                           ('2026-09-18','次新'),('2026-09-10','第三期'),('2026-09-01','第四期')]:
            changed=BASE.copy();changed[1]=name;changed[2]=stamp;rows.append(changed)
        self.write_source(rows)
        data=panel.compute_data(self.path)
        asset=data['assets'][0]
        self.assertEqual(len(asset['projects']),4)
        self.assertEqual(asset['dates'],['2026-09-21','2026-09-18','2026-09-10'])
        self.assertEqual(verify_runtime(data)['projects'],4)
        wrong=copy.deepcopy(data);wrong['assets'][0]['projects'].pop()
        self.assertEqual(qc.check_data(wrong,self.path,raise_on_fail=False)['status'],'FAIL')
        wrong=copy.deepcopy(data);wrong['assets'][0]['dates'].pop()
        self.assertEqual(qc.check_data(wrong,self.path,raise_on_fail=False)['status'],'FAIL')

    def test_invalid_project_and_layer_amounts_cannot_disable_upper_bounds(self):
        for field,value in [('规模','10亿元'),('规模',True),('规模',-1),('规模','NaN'),
                            ('对应金额（亿）','0.3亿元'),('对应金额（亿）',False),
                            ('对应金额（亿）',-1),('对应金额（亿）','Infinity')]:
            with self.subTest(field=field,value=value):
                if field=='规模':
                    changed=BASE.copy();changed[11]=value;self.write_source([changed])
                else:
                    book=Workbook();sheet=book.active;sheet.title='总表'
                    sheet.append(HEADERS+['对应金额（亿）']);sheet.append(BASE+[value])
                    book.save(self.path);book.close()
                with self.assertRaises(qc.BookkeepingGateError) as ctx:panel.compute_data(self.path)
                expected='PROJECT_AMOUNT_FORMAT' if field=='规模' else 'LAYER_AMOUNT_FORMAT'
                self.assertIn(expected,[item['code'] for item in ctx.exception.report['issues']])

    def test_truly_missing_project_and_layer_amounts_warn(self):
        changed=BASE.copy();changed[11]=None;self.write_source([changed])
        result=panel.compute_data(self.path)
        self.assertEqual(result['qc']['status'],'PASS')
        codes=[item['code'] for item in result['qc']['issues']]
        self.assertIn('PROJECT_SIZE_MISSING',codes)
        self.assertIn('LAYER_SIZE_MISSING',codes)

    def test_full_provenance_embedded_source_is_html_safe(self):
        import dashboard_provenance as provenance
        self.write_source()
        source=self.path.with_name('<b>来源&台账.xlsx')
        source.write_bytes(self.path.read_bytes())
        document=provenance.add_full_generation_provenance('<html><body></body></html>',source,source_hash=provenance.sha256(source))
        self.assertNotIn('<b>来源',document)
        self.assertIn('\\u003c',document)
        self.assertEqual(provenance.verify_full_generation(document,source)['source'],source.name)

    def test_partial_refresh_cannot_resign_full_generation_provenance(self):
        import deploy_github_pages as deploy
        import dashboard_provenance as provenance
        from gen_integrated_dashboard import build_integrated_html
        self.write_source();data=panel.compute_data(self.path)
        doc=build_integrated_html([('pricing','bookkeeping',panel.render_body(data))],panel.BOOKKEEPING_CSS)
        doc=doc.replace('</body>','<script id="stale-other-panel">配售原值1</script></body>')
        doc=provenance.add_full_generation_provenance(doc,self.path,source_hash=qc.file_sha256(self.path))
        source=Path(self.temp.name)/'complete.html';source.write_text(doc,encoding='utf-8')
        deploy.verify_dashboard_artifact(source,self.path)
        changed=BASE.copy();changed[7]=0.5;self.write_source([changed])
        target=Path(self.temp.name)/'partial.html'
        updater.refresh(self.path,source,target)
        content=target.read_text(encoding='utf-8')
        self.assertIn('配售原值1',content)
        self.assertEqual(qc.extract_data(content)['assets'][0]['projects'][0]['tiers'][0]['allocations'][0]['share'],0.5)
        with self.assertRaisesRegex(RuntimeError,'其他台账面板来源陈旧'):
            deploy.verify_dashboard_artifact(target,self.path)

    def test_inlined_component_tampering_blocked(self):
        self.write_source()
        html = panel.render_body(panel.compute_data(self.path))
        tampered = html.replace("view = 'bids'", "view = 'allocations'")
        with self.assertRaisesRegex(ValueError, '实际HTML/JS'):
            qc.verify_html(tampered,self.path)

    def test_future_booking_date_and_amount_string_blocked(self):
        for index,value,code in [(2,'2026-09-30','DATE_RANGE'),(7,'1亿元','AMOUNT_FORMAT')]:
            with self.subTest(code=code):
                changed=BASE.copy();changed[index]=value;self.write_source([changed])
                with self.assertRaises(qc.BookkeepingGateError) as ctx:panel.compute_data(self.path)
                self.assertIn(code,[i['code'] for i in ctx.exception.report['issues']])

    def test_mobile_reuses_payload_and_handles_new_profile_aliases(self):
        import json
        import gen_mobile_dashboard as mobile
        self.write_source()
        data=panel.compute_data(self.path)
        desktop=Path(self.temp.name)/'desktop.html'
        desktop.write_text('<html><style>'+panel.BOOKKEEPING_CSS+'</style><body>'+panel.render_body(data)+'</body></html>')
        with mock.patch.object(mobile,'build_prog_quick',return_value=[{'n':'机构'}]), \
             mock.patch.object(mobile,'build_asset',return_value={'kpis':[]}), \
             mock.patch.object(mobile,'build_peer',return_value={'groups':[],'trust':[],'top':[],'colors':{}}), \
             mock.patch.object(mobile,'build_meta',return_value={}), \
             mock.patch.object(mobile,'build_inst_stats',return_value=None):
            output=mobile.build_mobile_html(desktop,Path(self.temp.name)/'mobile.html')
        html=output.read_text()
        self.assertIn('bookkeepingHtml',html)
        self.assertIn('m-bookkeeping-frame',html)
        self.assertIn('\\u003c',html)
        raw=re.search(r'window.ABS_MOBILE_DATA=(.*?);</script>',html,re.S).group(1)
        payload=json.loads(raw)['bookkeepingHtml']
        self.assertEqual(qc.extract_data(payload),data)
        mobile.verify_mobile_bookkeeping(desktop.read_text(), html)
        broken=html.replace("if (INJ.bookkeepingHtml) items.push", "if (false) items.push")
        with self.assertRaisesRegex(ValueError,'实际导航应用脚本'):
            mobile.verify_mobile_bookkeeping(desktop.read_text(),broken)
        with self.assertRaises(ValueError):mobile.verify_mobile_bookkeeping(desktop.read_text(), '<html>旧手机</html>')

    def test_generator_always_registers_and_publish_blocks_old_html(self):
        from gen_integrated_dashboard import main
        self.write_source()
        # 真正main在其他模块计算前运行必选簿记门禁；故意让它FAIL，验证旧产物不变。
        target=Path(self.temp.name)/'previous.html';target.write_text('上一版')
        broken=BASE.copy();broken[8]=3;self.write_source([broken])
        with mock.patch('sys.argv',['generate',str(self.path),str(target)]):
            with self.assertRaises(qc.BookkeepingGateError):main()
        self.assertEqual(target.read_text(),'上一版')
        import deploy_github_pages as deploy
        old=Path(self.temp.name)/'old.html';old.write_text('<div class="panel"></div><script>function selectModule(){}function selectSub(){}</script>')
        with self.assertRaises(RuntimeError):deploy.verify_dashboard_artifact(old,self.path)


if __name__=='__main__':unittest.main()
