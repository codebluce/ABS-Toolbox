"""将经用户确认的华泰24项候选补入同业底稿；日期按含设立日四交易日推算。"""
from pathlib import Path
from datetime import date, datetime, timedelta
from decimal import Decimal
from copy import copy
import json
import hashlib
import argparse
import openpyxl
from peer_issuance_panel import parse_dynamics, normalize_product_key

ROOT = Path(__file__).resolve().parents[1]

def closed_days():
    # 境内证券交易日：调休周末不开市；覆盖本批推算所需的2025年底及2026假期。
    ranges = [('2026-01-01','2026-01-03'),('2026-02-15','2026-02-23'),
              ('2026-04-04','2026-04-06'),('2026-05-01','2026-05-05'),
              ('2026-06-19','2026-06-21'),('2026-09-25','2026-09-27'),
              ('2026-10-01','2026-10-07')]
    result=set()
    for begin,end in ranges:
        d=date.fromisoformat(begin)
        while d<=date.fromisoformat(end):
            result.add(d);d+=timedelta(days=1)
    return result

def infer_booking(established):
    """设立日计第一个交易日，向前到第四个；10/16→10/13。"""
    d=established;count=0;closed=closed_days()
    if d.weekday()>=5 or d in closed:
        raise ValueError(f'设立日不是交易日，须人工确认: {d}')
    while True:
        if d.weekday()<5 and d not in closed:
            count+=1
            if count==4:return d
        d-=timedelta(days=1)

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--comparison',type=Path,required=True)
    parser.add_argument('--source',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    compared=json.loads(args.comparison.read_text())
    candidates=compared['results']['2026']['missing_same']
    if len(candidates)!=24:raise ValueError('确认清单应为24项')
    original=parse_dynamics(args.source);keys={normalize_product_key(r['product']) for r in original}
    if len(keys)!=len(original):raise ValueError('基线产品名称重复')
    w=openpyxl.load_workbook(args.source);s=w.worksheets[0]
    s.title='2026同业发行底稿'
    for c,title in enumerate(['录入来源','簿记日口径','原始设立日','来源工作表','来源行号','资产映射说明'],19):s.cell(1,c,title)
    for r in range(2,s.max_row+1):
        if s.cell(r,2).value:
            if not s.cell(r,19).value:s.cell(r,19,args.source.name)
            s.cell(r,20,'原有来源簿记日（未改写）')
    cross=w.create_sheet('跨年候选2025')
    for c in range(1,25):cross.cell(1,c,s.cell(1,c).value)
    raw=w.create_sheet('华泰分层原值')
    raw.append(['项目','华泰工作表','原始行','评级','分层','规模亿元','票面利率百分数','期限年'])
    entries=[]
    end=max(i for i in range(2,s.max_row+1) if s.cell(i,2).value)
    for h in candidates:
        key=normalize_product_key(h['product'])
        if key in keys:raise ValueError(f'候选已经存在:{h["product"]}')
        keys.add(key);est=date.fromisoformat(h['date']);booking=infer_booking(est)
        party=h['party'];asset=h['asset']
        # 不把缺失基础资产名硬归到已有特定贷款产品；保留待细分标签。
        if party=='京东':base='京东'+('金条白取混池' if '金条/白取' in asset else asset);atype='消金提现类' if '金条/白取' in asset else '小微toB'
        elif party=='度小满':base='度小满广泰资产（待细分）';atype='消金提现类'
        elif party=='金城银行':base='金城银行贷款资产（待细分）';atype='小微toB'
        elif party=='平安普惠':base='平安普惠经营贷';atype='小微toB'
        else:base='蚂蚁经营贷';atype='小微toB'
        aaa=[t for t in h['tiers'] if t['rating'].strip().startswith('AAA')]
        aa=[t for t in h['tiers'] if t['rating'].strip().startswith('AA+')]
        def joined(items,k,suffix=''):
            return '、'.join(str(t[k])+suffix for t in items if t[k] is not None)
        def percent(items):
            return '、'.join(f"{Decimal(t['amount'])/Decimal(h['amount'])*100:.4f}%" for t in items if t['amount'] is not None)
        values=[f'{booking.year}.{booking.month}.{booking.day}',h['product'],'ABS' if h['sheet'].startswith('交易所-') else 'ABN',h['venue'],h['originator'],base,atype,float(h['amount']),datetime.combine(booking,datetime.min.time()),datetime.combine(est,datetime.min.time()),joined(aaa,'rate','%') or None,joined(aaa,'term') or None,percent(aaa) or None,joined(aa,'rate','%') or None,joined(aa,'term') or None,percent(aa) or None,None,h['manager'],'华泰公告发行明细（用户确认补录）','推算：设立日含当日第4交易日，非公告实际簿记日',est.isoformat(),h['sheet'],h['row'],'按资产方归属；未披露具体资产部分待细分']
        target=s if booking.year==2026 else cross
        if target is s:end+=1;row=end
        else:row=target.max_row+1
        for c,v in enumerate(values,1):
            cell=target.cell(row,c,v)
            if target is s:cell._style=copy(s.cell(2,c)._style)
        for c in [9,10]:target.cell(row,c).number_format='yyyy/mm/dd'
        for t in h['tiers']:raw.append([h['product'],h['sheet'],t['row'],t['rating'],t['layer'],float(t['amount']) if t['amount'] is not None else None,float(t['rate']) if t['rate'] is not None else None,float(t['term']) if t['term'] is not None else None])
        entries.append({'project':h['product'],'amount':h['amount'],'established':h['date'],'booking_inferred':booking.isoformat(),'target_sheet':target.title,'AAA_missing':not aaa or not joined(aaa,'rate'),'source_sheet':h['sheet'],'source_row':h['row']})
    notes=w.create_sheet('数据口径与来源')
    notes.append(['项目','说明'])
    note='本底稿涵盖现面板全部已录2026项目；24项华泰补录中23项计入2026，1项倒推至2025年并单列。补录簿记日按用户示例含设立日向前第4交易日推算，非公告实证；跳过周末与境内证券休市日。原有记录日期不改写。9月28日至10月11日仍未取得完整同业周表，补录不代表市场资料已完整。缺少AAA评级或利率不推断填充。'
    notes.append(['面板说明',note]);notes.append(['基线',args.source.name]);notes.append(['算法','2026-10-16→2026-10-13；设立日计第1交易日，向前3个交易日']);notes.append(['休市区间',json.dumps(sorted(str(x) for x in closed_days()))])
    notes.append(['交易日依据','上证公告〔2025〕45号、〔2026〕22号；https://www.sse.com.cn/disclosure/announcement/general/c/c_20260915_10832273.shtml'])
    notes.append(['华泰输入SHA256',next(v for k,v in compared['provenance']['input_sha256'].items() if '公告发行' in k)])
    for sheet in [s,cross,raw]:sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    args.output.parent.mkdir(parents=True,exist_ok=True);w.save(args.output);w.close()
    record={'source':str(args.source),'output':str(args.output),'entries':entries,'total_count':len(entries),'total_amount':str(sum(Decimal(e['amount']) for e in entries)),'2026_count':sum(e['target_sheet']==s.title for e in entries),'2026_amount':str(sum(Decimal(e['amount']) for e in entries if e['target_sheet']==s.title)),'note':note}
    args.output.with_suffix('.supplement.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n')
    current=parse_dynamics(args.output);assert len(current)==len(original)+record['2026_count']
    print(json.dumps({k:v for k,v in record.items() if k!='entries'},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
