"""最新簿记硬门禁：独立源表核对、业务规则和产物完整性；无网络/发布操作。"""
from __future__ import annotations

from contextvars import ContextVar
from datetime import datetime
from decimal import Decimal, InvalidOperation
import hashlib
from html.parser import HTMLParser
import json
import math
from pathlib import Path
import re
import shutil
import subprocess

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

SCHEMA_VERSION = 'latest-bookkeeping/v2'
RECENT_PERIODS = 3
VENUES = {'上交所', '深交所', '银行间'}
EXCLUDED_VENUES = {'非标', '平层非标', '结构化非标', '保登', '中保登'}
ORDER = ('赊销白条', '金条白取混发', '信托金条', '白条取现', '信托企业主贷',
         '小贷企业主贷', '保理', '金采', '京企贷', '动产质押', '汽车融资租赁')
ALIASES = {'信托白取': '白条取现', '金条&白取混发': '金条白取混发'}
FIELDS = ('规模', '发行场所', '资产类型', '项目名称', '簿记时间', '分层情况', '成本',
          '认购机构', '认购份额', '申购利率', '穿透机构', '申购规模', '期限')
_MONTHLY_GATE_ACTIVE = ContextVar('monthly_bookkeeping_gate_active', default=False)


class BookkeepingGateError(ValueError):
    def __init__(self, report):
        self.report = report
        failures = [f"{item['code']}: {item['message']}" for item in report['issues'] if item['level'] == 'FAIL']
        super().__init__('最新簿记门禁失败: ' + '; '.join(failures[:8]))


def component_self_check():
    """每次更新前运行真实JS组件的最小契约自检，坏表头/导出/色标不放行。"""
    node = shutil.which('node')
    if node is None:
        raise ValueError('最新簿记门禁需要node执行组件自检，未找到node')
    result = subprocess.run([node, str(Path(__file__).with_name('test_latest_bookkeeping_panel.js'))],
                            text=True, capture_output=True)
    if result.returncode:
        raise ValueError('最新簿记组件自检失败: ' + result.stdout + result.stderr)


def file_sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def rules_sha256():
    root = Path(__file__).resolve().parent
    digest = hashlib.sha256()
    for name in ('latest_bookkeeping_panel.py', 'latest_bookkeeping_panel.js',
                 'latest_bookkeeping_panel.css', 'latest_bookkeeping_panel_body.html', 'latest_bookkeeping_qc.py',
                 'ledger_formula_tools.py', 'repair_r_formulas.py'):
        digest.update(name.encode())
        digest.update((root / name).read_bytes())
    return digest.hexdigest()


def _text(value):
    return '' if value is None else str(value).strip()


def _number(value):
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def _date(value, epoch):
    if value is None or value == '':
        return ''
    try:
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            value = from_excel(value, epoch)
        if isinstance(value, datetime):
            return value.strftime('%Y-%m-%d')
        return datetime.fromisoformat(str(value).strip().replace('/', '-')).strftime('%Y-%m-%d')
    except (ValueError, TypeError, OverflowError):
        return ''



_R_PRODUCT = re.compile(r'^=\s*(\$?([JQ])\$?([1-9][0-9]*))\s*\*\s*(\$?([JQ])\$?([1-9][0-9]*))\s*$', re.I)
_R_CONSTANT = re.compile(r'^=\s*([+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[Ee][+-]?[0-9]+)?)\s*$')
_R_EPSILON = Decimal('0.0001')


def validate_source_r_formulas(sheet, formulas, positions, merged, records):
    """Independent source proof; explicit merges only, no producer aggregation/fill.

    Literal R keeps existing business checks. Direct J*Q requires complete source
    identities and plain numeric source operands; no Excel evaluation. Constants
    compare with their literal, never with J*Q. Missing R caches remain WARN and
    are never filled here: incremental inputs must arrive with prepared raw caches.
    """
    rcol = positions.get('对应金额（亿）')
    if rcol is None:
        return
    anchors = {}
    for area in formulas.merged_cells.ranges:
        for rr in range(area.min_row, area.max_row + 1):
            for cc in range(area.min_col, area.max_col + 1):
                anchors[rr, cc] = (area.min_row, area.min_col)

    def raw(rr, cc):
        return merged.get((rr, cc), sheet.cell(rr, cc).value)

    def identity(rr, tier=False):
        fields = ('资产类型', '簿记时间', '项目名称', '发行场所')
        key = tuple(_date(raw(rr, positions[f]), sheet.parent.epoch) if f == '簿记时间'
                    else _text(raw(rr, positions[f])) for f in fields)
        return key + ((_text(raw(rr, positions['分层情况'])),) if tier else ())

    def numeric(value, label):
        if isinstance(value, (bool, str)) or value is None:
            raise ValueError(f'{label}: R_FORMULA_OPERAND 原始值必须是非负有限数值')
        try:
            result = Decimal(str(value))
        except (InvalidOperation, ValueError, TypeError) as error:
            raise ValueError(f'{label}: R_FORMULA_OPERAND 数值非法') from error
        if not result.is_finite() or result < 0:
            raise ValueError(f'{label}: R_FORMULA_OPERAND 数值非法')
        return result

    for record in records:
        rr = record['_row']
        ar, ac = anchors.get((rr, rcol), (rr, rcol))
        expression = formulas.cell(ar, ac)
        if expression.data_type != 'f':
            continue
        label = f'总表R{rr}'
        formula = expression.value
        constant = _R_CONSTANT.fullmatch(formula or '')
        product = _R_PRODUCT.fullmatch(formula or '')
        if constant:
            expected = Decimal(constant[1])
            if not expected.is_finite() or expected < 0:
                raise ValueError(f'{label}: R_FORMULA_CONSTANT 常量须为非负有限数值')
            record['_rFormulaKind'] = 'constant'
        elif product:
            refs = {product[2].upper(): int(product[3]), product[5].upper(): int(product[6])}
            if set(refs) != {'J', 'Q'}:
                raise ValueError(f'{label}: R_FORMULA_REFERENCE 须直接引用一个J和一个Q')
            if positions.get('规模') != 10 or positions.get('对应金额（亿）') != 18:
                raise ValueError(f'{label}: R_FORMULA_SCHEMA J/R列位置不符合源表定义')
            key, layer = identity(rr), identity(rr, True)
            if not all(layer):
                raise ValueError(f'{label}: R_FORMULA_IDENTITY 项目/档位身份不完整')
            if not all(2 <= ref <= sheet.max_row for ref in refs.values()):
                raise ValueError(f'{label}: R_FORMULA_REFERENCE 引用不在源表数据范围')
            if identity(refs['J']) != key or identity(refs['Q'], True) != layer:
                raise ValueError(f'{label}: R_FORMULA_REFERENCE J跨项目或Q跨项目/档位 ({formula})')
            operands = {}
            for col, cc in (('J', 10), ('Q', 17)):
                target = refs[col]
                er, ec = anchors.get((target, cc), (target, cc))
                if formulas.cell(er, ec).data_type == 'f':
                    raise ValueError(f'{label}: R_FORMULA_OPERAND {col}{target}本身为公式，缓存不能证明其来源')
                operands[col] = numeric(raw(target, cc), f'{label}/{col}{target}')
            if operands['Q'] > 1:
                raise ValueError(f'{label}: R_FORMULA_OPERAND Q比例超出[0,1]')
            expected = operands['J'] * operands['Q']
            record['_rFormulaKind'] = 'direct-jq'
        else:
            raise ValueError(f'{label}: R_FORMULA_UNSUPPORTED 复杂/跨sheet/范围公式阻断，不推算 ({formula})')
        cached = raw(rr, rcol)
        if cached is None:
            record['_uncachedLayerSize'] = f'R{rr}'
            continue
        amount = numeric(cached, label)
        if abs(amount - expected) > _R_EPSILON:
            raise ValueError(f'{label}: R_FORMULA_CACHE 缓存{amount}与已验证值{expected}差异超过0.0001')


def r_formula_evidence(records):
    formulas = [r for r in records if r.get('_rFormulaKind')]
    return {'reader': 'independent source parser', 'formulas': len(formulas),
            'directJQ': sum(r['_rFormulaKind'] == 'direct-jq' for r in formulas),
            'constants': sum(r['_rFormulaKind'] == 'constant' for r in formulas),
            'uncached': sum(bool(r.get('_uncachedLayerSize')) for r in formulas),
            'epsilonBillions': '0.0001',
            'missingCachePolicy': 'WARN; QC never fills R; incremental input requires prepared raw caches',
            'unsupportedPolicy': 'FAIL; no evaluation'}


def read_source(path):
    """只展开明确合并的项目/层字段，不跨空行、项目或档位向下猜测填充。

    U/V/W/X/Y原始坐标独立保留。机构/金额合并无法逐笔确认，硬阻断要求源表明确。
    """
    book = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    try:
        if '总表' not in book.sheetnames:
            raise ValueError('台账必须包含“总表”工作表')
        sheet = book['总表']
        headers = [_text(cell.value) for cell in sheet[1]]
        missing = set(FIELDS) - set(headers)
        if missing or len(headers) != len(set(headers)):
            raise ValueError(f'总表列缺失或重复: {sorted(missing)}')
        positions = {name: i + 1 for i, name in enumerate(headers)}
        formula_sheet = formulas['总表']
        critical = set(FIELDS) | {'规模', '计划管理人'}
        for name in critical:
            if name not in positions:
                continue
            column = positions[name]
            for r in range(2, formula_sheet.max_row + 1):
                expression = formula_sheet.cell(r, column)
                cached = sheet.cell(r, column)
                if expression.data_type == 'f' and (cached.value is None or cached.data_type == 'e'):
                    raise ValueError(f'总表{expression.coordinate}公式没有有效计算缓存，请用Excel重算后保存再生成')
        merged = {}
        for area in sheet.merged_cells.ranges:
            if area.max_col >= positions['认购机构']:
                raise ValueError(f'机构/金额/报价列存在合并单元格，须先明确逐笔原值: {area}')
            for r in range(area.min_row, area.max_row + 1):
                for c in range(area.min_col, area.max_col + 1):
                    merged[(r, c)] = sheet.cell(area.min_row, area.min_col).value
        records = []
        for index, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            values = {name: merged.get((index, col), cells[col - 1]) for name, col in positions.items()}
            if not any(value is not None for value in values.values()):
                continue
            complete_detail_header = all(_text(values.get(name)) == name for name in ('成本','认购机构','认购份额','申购利率','穿透机构','申购规模'))
            if complete_detail_header or _text(values.get('项目名称')) == '项目名称':
                continue
            values['_row'] = index
            values['_uncachedLayerSize'] = None
            values['_date'] = _date(values.get('簿记时间'), book.epoch)
            records.append(values)
        validate_source_r_formulas(sheet, formula_sheet, positions, merged, records)
        return records
    finally:
        book.close()
        formulas.close()


def check_data(data, source_path, *, raise_on_fail=True):
    """独立按原始坐标逐笔对账，避免生产聚合函数自己验证自己。"""
    report = {'schema': SCHEMA_VERSION, 'checkedAt': datetime.now().astimezone().isoformat(timespec='seconds'),
              'source': data.get('source', Path(source_path).name), 'issues': [], 'checks': {}, 'status': 'PASS'}
    def issue(level, code, message):
        report['issues'].append({'level': level, 'code': code, 'message': message})
    try:
        records = read_source(source_path)
        report['checks']['rFormulas'] = r_formula_evidence(records)
    except (ValueError, OSError) as error:
        issue('FAIL', 'SOURCE_SCHEMA', str(error))
        records = []
    eligible = []
    for row in records:
        venue = _text(row.get('发行场所'))
        if venue not in VENUES:
            if venue not in EXCLUDED_VENUES:
                issue('WARN', 'UNKNOWN_VENUE', f"总表第{row['_row']}行场所‘{venue or '未披露'}’未纳入")
            continue
        project = _text(row.get('项目名称'))
        asset = _text(row.get('资产类型'))
        if not project or not asset or not row['_date']:
            issue('FAIL', 'IDENTITY_OR_DATE', f"总表第{row['_row']}行缺少资产/项目/有效簿记日期")
            continue
        match = re.fullmatch(r'(\d{4})年ABS发行台账-(\d{4})-定稿.xlsx', data.get('source', Path(source_path).name))
        if match:
            cutoff = ''.join(match.groups())
            if row['_date'].replace('-', '') > cutoff or row['_date'][:4] != match.group(1):
                issue('FAIL', 'DATE_RANGE', f"总表第{row['_row']}行簿记日期{row['_date']}不在该定稿业务年份/截止日内")
        eligible.append(row)
    if not eligible:
        issue('FAIL', 'EMPTY_SCOPE', '上交所/深交所/银行间没有有效项目')
    source_assets = {}
    for row in eligible:
        raw = _text(row['资产类型'])
        asset = ALIASES.get(raw, raw)
        source_assets.setdefault(asset, []).append(row)
    asset_names = sorted(source_assets, key=lambda name: ORDER.index(name) if name in ORDER else len(ORDER))
    if [a.get('name') for a in data.get('assets', [])] != asset_names:
        issue('FAIL', 'ASSET_SCOPE_ORDER', '资产集合、合并口径或排列顺序不一致')
    total_records = 0
    for asset in data.get('assets', []):
        source_rows = source_assets.get(asset.get('name'), [])
        if not source_rows:
            continue
        ordered = sorted(source_rows, key=lambda row: row['_date'], reverse=True)
        expected_projects = list(dict.fromkeys((row['_date'], _text(row['项目名称']), _text(row['发行场所'])) for row in ordered))[:RECENT_PERIODS]
        recent_dates = list(dict.fromkeys(key[0] for key in expected_projects))
        recent = recent_dates[0]
        keys = set(expected_projects)
        selected = [row for row in ordered if (row['_date'], _text(row['项目名称']), _text(row['发行场所'])) in keys]
        actual_projects = [(p.get('date'), p.get('name'), p.get('venue')) for p in asset.get('projects', [])]
        if asset.get('date') != recent or asset.get('dates') != recent_dates or actual_projects != expected_projects:
            issue('FAIL', 'LATEST_PROJECTS', f"{asset.get('name')}不是最新最多3个项目，或仍含旧项目")
        for project in asset.get('projects', []):
            rows = [row for row in selected if _text(row['项目名称']) == project.get('name') and row['_date'] == project.get('date') and _text(row['发行场所']) == project.get('venue')]
            if not rows:
                continue
            if project.get('date') not in recent_dates or project.get('venue') not in VENUES:
                issue('FAIL', 'PROJECT_SCOPE', f"{project.get('name')}日期或发行场所错误")
            for field, key in (('发行场所','venue'), ('计划管理人','manager'), ('期限','tenor')):
                source_values = {_text(row.get(field)) for row in rows}
                if len(source_values) > 1 or project.get(key) not in source_values:
                    issue('FAIL', 'PROJECT_FIELD_CONFLICT', f"{project.get('name')}的{field}不一致")
            if project.get('tenor') in {'', '-'}:
                issue('WARN', 'TENOR_MISSING', f"{project.get('name')}项目期限未披露，不推算或补造")
            # 项目J和辅助R金额非空非法不能被当成None关闭上限；真正缺失明确告警。
            for field, code in (('规模', 'PROJECT_AMOUNT_FORMAT'), ('对应金额（亿）', 'LAYER_AMOUNT_FORMAT')):
                for row in rows:
                    raw = _text(row.get(field))
                    value = _number(row.get(field))
                    if raw not in {'', '-'} and (value is None or isinstance(row.get(field), bool) or value < 0):
                        issue('FAIL', code, f"总表第{row['_row']}行{field}须为非负有限数值，不能以非法文本或布尔值关闭上限")
            if not any(_number(row.get('规模')) is not None for row in rows):
                issue('WARN', 'PROJECT_SIZE_MISSING', f"{project.get('name')}发行规模未披露，无法核验项目配售上限")
            sizes = {_number(row.get('规模')) for row in rows}
            if len(sizes) > 1 or project.get('size') not in sizes:
                issue('FAIL', 'PROJECT_SIZE', f"{project.get('name')}发行规模不一致")
            expected_layers = list(dict.fromkeys(_text(row.get('分层情况')) or '未分层' for row in rows))
            expected_layers = list(dict.fromkeys('未分层' if name == '-' else name for name in expected_layers))
            if [t.get('name') for t in project.get('tiers', [])] != expected_layers:
                issue('FAIL', 'LAYERS', f"{project.get('name')}档位集合错误")
            total_share = 0
            for tier in project.get('tiers', []):
                layer_rows = [row for row in rows if (_text(row.get('分层情况')) if _text(row.get('分层情况')) not in {'','-'} else '未分层') == tier.get('name')]
                expected_costs = []
                expected_allocation, expected_bids = [], []
                for row in layer_rows:
                    cost = _number(row.get('成本'))
                    value = cost if cost is not None else _text(row.get('成本'))
                    if value != '' and value not in expected_costs:
                        expected_costs.append(value)
                    if cost is None and _text(row.get('成本')) == '':
                        issue('WARN', 'COST_MISSING', f"总表第{row['_row']}行成本未披露")
                    for field in ('认购份额', '申购规模'):
                        val = _number(row.get(field))
                        raw = _text(row.get(field))
                        if isinstance(row.get(field), bool):
                            issue('FAIL', 'AMOUNT_FORMAT', f"总表第{row['_row']}行{field}不能是布尔值")
                        if val is None and raw not in {'', '-', '认购份额', '申购规模'}:
                            issue('FAIL', 'AMOUNT_FORMAT', f"总表第{row['_row']}行{field}不是明确的数值或缺失标记")
                        if val is not None and val < 0:
                            issue('FAIL', 'NEGATIVE_AMOUNT', f"总表第{row['_row']}行{field}为负数")
                    if not _text(row.get('分层情况')):
                        issue('WARN', 'LAYER_MISSING', f"总表第{row['_row']}行档位未披露，保留为未分层")
                    for field in ('成本','申购利率'):
                        value = _number(row.get(field))
                        raw = _text(row.get(field))
                        if isinstance(row.get(field), bool):
                            issue('FAIL', 'RATE_FORMAT', f"总表第{row['_row']}行{field}不能是布尔值")
                        if raw in {'成本', '申购利率'}:
                            issue('FAIL', 'LOCAL_HEADER', f"总表第{row['_row']}行正文出现利率表头标记")
                        if value is None and raw not in {'', '-', '自持', '不出表', '成本', '申购利率'}:
                            issue('FAIL', 'RATE_FORMAT', f"总表第{row['_row']}行{field}不是数值利率或约定标记")
                        if value is not None and not 0 <= value <= 1:
                            issue('FAIL', 'RATE_UNIT', f"总表第{row['_row']}行{field}非0~1小数利率")
                    institution = _text(row.get('认购机构'))
                    share = _number(row.get('认购份额'))
                    if institution and institution not in {'-', '认购机构'}:
                        expected_allocation.append((institution, share, cost if cost is not None else _text(row.get('成本'))))
                        if share is None:
                            issue('WARN', 'ALLOCATION_AMOUNT_MISSING', f"总表第{row['_row']}行配售规模未披露")
                    elif share is not None:
                        issue('FAIL', 'ORPHAN_ALLOCATION', f"总表第{row['_row']}行有配售规模无机构")
                    bidder = _text(row.get('穿透机构'))
                    bid_size = _number(row.get('申购规模'))
                    quote = _number(row.get('申购利率'))
                    if bidder and bidder not in {'-', '穿透机构'}:
                        expected_bids.append((bidder, quote if quote is not None else _text(row.get('申购利率')), bid_size))
                        if bid_size is None or row.get('申购利率') is None:
                            issue('WARN', 'BID_FIELD_MISSING', f"总表第{row['_row']}行申购字段未披露")
                    elif quote is not None or bid_size is not None:
                        issue('FAIL', 'ORPHAN_BID', f"总表第{row['_row']}行有报价/申购规模无机构")
                actual_allocation = [(r.get('institution'),r.get('share'),r.get('cost')) for r in tier.get('allocations', [])]
                actual_bids = [(r.get('institution'),r.get('rate'),r.get('bidSize')) for r in tier.get('bids', [])]
                if actual_allocation != expected_allocation or actual_bids != expected_bids or tier.get('costs') != expected_costs:
                    issue('FAIL', 'DETAIL_RECONCILIATION', f"{project.get('name')}/{tier.get('name')}逐笔机构、成本、规模或利率与源表不一致")
                if any(r.get('rate') is not None or r.get('bidSize') is not None for r in tier.get('allocations', [])) or any(r.get('share') is not None for r in tier.get('bids', [])):
                    issue('FAIL', 'DUAL_TRACK_PAIRING', '最终配售与申购记录被错误关联')
                layer_share = sum(r.get('share') or 0 for r in tier.get('allocations', []))
                total_share += layer_share
                layer_sizes = {_number(row.get('对应金额（亿）')) for row in layer_rows} - {None}
                uncalculated = [row['_uncachedLayerSize'] for row in layer_rows if row.get('_uncachedLayerSize')]
                if not layer_sizes and not uncalculated:
                    issue('WARN', 'LAYER_SIZE_MISSING', f"{project.get('name')}/{tier.get('name')}档位金额未披露，无法核验档位上限")
                if uncalculated:
                    issue('WARN', 'LAYER_SIZE_UNCACHED', f"{project.get('name')}/{tier.get('name')}辅助档位金额公式无缓存（{uncalculated[0]}等），无法校验档位上限，仍执行项目配售上限校验")
                if len(layer_sizes) == 1 and layer_share - next(iter(layer_sizes)) > 0.0001:
                    issue('FAIL', 'ALLOCATION_LAYER_OVER', f"{project.get('name')}/{tier.get('name')}配售{layer_share:.6g}亿超过档位金额{next(iter(layer_sizes))}亿")
                if len(layer_sizes) > 1:
                    issue('FAIL', 'LAYER_SIZE_CONFLICT', f"{project.get('name')}/{tier.get('name')}档位金额冲突")
                total_records += len(actual_allocation) + len(actual_bids)
                if len(set(map(repr, actual_allocation))) < len(actual_allocation) or len(set(map(repr, actual_bids))) < len(actual_bids):
                    issue('WARN', 'DUPLICATE_DETAILS', f"{project.get('name')}/{tier.get('name')}存在相同明细，保留原始记录待复核")
                if len(expected_costs) > 1:
                    issue('WARN', 'MULTIPLE_COSTS', f"{project.get('name')}/{tier.get('name')}存在多条成本口径，逐笔成本必须保留")
                if not actual_bids:
                    issue('WARN', 'NO_BIDS', f"{project.get('name')}/{tier.get('name')}未披露申购明细，不回退历史")
            if project.get('size') is not None and total_share - project['size'] > 0.0001:
                issue('FAIL', 'ALLOCATION_PROJECT_OVER', f"{project.get('name')}配售合计{total_share:.6g}亿超过发行规模{project['size']}亿")
            elif project.get('size') is not None and abs(total_share - project['size']) > 0.0001:
                issue('WARN', 'ALLOCATION_SIZE_DIFF', f"{project.get('name')}已披露配售合计{total_share:.6g}亿与发行规模{project['size']}亿不同")
    source_dates = [r['_date'] for r in records if r['_date']]
    if source_dates and data.get('asOf') != max(source_dates):
        issue('FAIL', 'AS_OF_DATE', '台账最新簿记日期与源表不一致')
    meta = data.get('provenance', {})
    if meta.get('schema') != SCHEMA_VERSION or meta.get('sourceSha256') != file_sha256(source_path) or meta.get('rulesSha256') != rules_sha256():
        issue('FAIL', 'PROVENANCE', '来源哈希或面板规则指纹不匹配')
    report['checks'] = {'assets':len(data.get('assets',[])), 'projects':sum(len(a['projects']) for a in data.get('assets',[])), 'details':total_records}
    report['status'] = 'FAIL' if any(i['level']=='FAIL' for i in report['issues']) else 'PASS'
    if raise_on_fail and report['status'] == 'FAIL':
        raise BookkeepingGateError(report)
    return report


class _EmbeddedParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.payloads, self.collecting, self.buffer = [], False, []
    def handle_starttag(self, tag, attrs):
        if tag == 'script' and dict(attrs).get('id') == 'lb-data':
            self.collecting, self.buffer = True, []
    def handle_data(self, data):
        if self.collecting:
            self.buffer.append(data)
    def handle_endtag(self, tag):
        if tag == 'script' and self.collecting:
            self.payloads.append(''.join(self.buffer))
            self.collecting = False


def extract_data(content):
    parser = _EmbeddedParser()
    parser.feed(content)
    if len(parser.payloads) != 1:
        raise ValueError(f'最新簿记数据块应唯一，实际{len(parser.payloads)}个')
    return json.loads(parser.payloads[0])


def has_monthly_markers(content):
    """数据块、DOM、资源或renderer身份任一存在，都不能按旧组件放行。"""
    pattern = r'lb-monthly|monthly[-_]bookkeeping|monthly_bid_institutions'
    if re.search(pattern, content, flags=re.I):
        return True
    class Markers(HTMLParser):
        found = False
        def handle_starttag(self, tag, attrs):
            if any(re.search(pattern, value or '', flags=re.I) for _, value in attrs):
                self.found = True
        handle_startendtag = handle_starttag
    parser = Markers()
    parser.feed(content)
    return parser.found


def _load_monthly_qc():
    # 旧组件默认路径不加载实验模块及其额外依赖。
    from lab import monthly_bookkeeping_qc
    return monthly_bookkeeping_qc


def verify_html(content, source_path=None):
    if _MONTHLY_GATE_ACTIVE.get():
        raise ValueError('月度门禁不得回调旧verify_html；最近项目请调用check_data，避免递归')
    if has_monthly_markers(content):
        if source_path is None:
            raise ValueError('月度组件必须显式提供源台账并通过专门门禁')
        monthly_qc = _load_monthly_qc()
        _, monthly = monthly_qc.extract_data(content)
        if (not isinstance(monthly, dict) or
                not isinstance(monthly.get('provenance'), dict) or
                monthly['provenance'].get('schema') != monthly_qc.SCHEMA_VERSION):
            raise ValueError('孤立月度payload或不支持的月度schema/renderer，停止旧组件校验')
        token = _MONTHLY_GATE_ACTIVE.set(True)
        try:
            report = monthly_qc.verify_html(content, source_path, raise_on_fail=True)
            if (report.get('status') != 'PASS' or report.get('releaseStatus') == 'BLOCK' or
                    any(item.get('level') == 'FAIL' for item in report.get('issues', []))):
                raise ValueError('月度专门门禁未通过，停止旧组件校验')
            return report
        finally:
            _MONTHLY_GATE_ACTIVE.reset(token)
    return _verify_legacy_html(content, source_path)


def _verify_legacy_html(content, source_path=None):
    """原组件HTML/JS/CSS和源表合同保持严格；不接受月度组件。"""
    data = extract_data(content)
    if content.count('id="latest-bookkeeping"') != 1 or '<h2>簿记明细</h2>' not in content:
        raise ValueError('最新簿记面板容器或标题缺失/重复')
    if '<button type="button" data-view="bids">申购明细</button><button type="button" data-view="allocations">最终配售</button>' not in content:
        raise ValueError('申购明细/最终配售顺序不符')
    meta = data.get('provenance', {})
    if meta.get('schema') != SCHEMA_VERSION or meta.get('rulesSha256') != rules_sha256():
        raise ValueError('最新簿记面板规则已更新，请重新生成看板')
    # 不能只相信JSON里的规则哈希：实际内联脚本/模板同样必须是当前代码。
    from latest_bookkeeping_panel import render_body, BOOKKEEPING_CSS
    if render_body(data) not in content and render_body(data, preview=True) not in content:
        raise ValueError('最新簿记实际HTML/JS与当前组件不一致，请重新生成')
    if '<html' in content and BOOKKEEPING_CSS not in content:
        raise ValueError('最新簿记样式与当前组件不一致')
    if source_path is not None:
        return check_data(data, source_path)
    return {'status':'PASS', 'checks':{'schema':SCHEMA_VERSION}, 'issues':[]}
