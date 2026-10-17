"""月表独立门禁的对抗测试；真实xlsx输入、Demo产物与实际inlineJS/CSS。"""
from __future__ import annotations

import copy
from datetime import datetime
import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook

LAB = Path(__file__).resolve().parent
SCRIPTS = LAB.parent
for directory in (SCRIPTS, LAB):
    if str(directory) not in sys.path:
        sys.path.insert(0, str(directory))
import monthly_bookkeeping_qc as gate
import monthly_bid_institutions_demo as demo
import latest_bookkeeping_panel as panel
import latest_bookkeeping_qc as recent_qc

# 原始 U/V/W/X/Y 固定物理位置，与生产或核对函数的字段推导无关。
HEADERS = ['序号', '月份', '类别', '资产类型', '项目名称', '发行场所', '计划管理人',
           '联席承销商', '托管行', '规模', '期限', '簿记时间', '国股CD', 'ALL-IN成本',
           '认购倍数', '分层情况', '分层占比', '对应金额（亿）', '评级', '成本',
           '认购机构', '认购份额', '申购利率', '穿透机构', '申购规模']


def source_row(project='五月甲', date='2026-05-06', asset='赊销白条', venue='上交所',
               allocated='建设银行', allocation=2, bidder='中国建设银行', size=3,
               rate=0.02, cost=0.02):
    row = [None] * 25
    values = {'资产类型': asset, '项目名称': project, '发行场所': venue,
              '计划管理人': '管理人', '规模': 20, '期限': '21+3', '簿记时间': date,
              '分层情况': '优先A', '成本': cost, '认购机构': allocated, '认购份额': allocation,
              '申购利率': rate, '穿透机构': bidder, '申购规模': size}
    for name, value in values.items():
        row[HEADERS.index(name)] = value
    return row


def replace_payload(content, identifier, value):
    encoded = json.dumps(value, ensure_ascii=False, allow_nan=False).replace('<', '\\u003c').replace('&', '\\u0026')
    pattern = r'(<script\b[^>]*\bid="' + re.escape(identifier) + r'"[^>]*>).*?(</script>)'
    updated, count = re.subn(pattern, lambda m: m[1] + encoded + m[2], content, flags=re.S)
    if count != 1:
        raise AssertionError(f'{identifier}数据块必须唯一')
    return updated


def with_provenance(monthly, ledger):
    # 测试只为旧构建器补足生成元数据；业务值全部来自真实生成产物。
    result = copy.deepcopy(monthly)
    result['provenance'] = gate.build_provenance(ledger)
    return result


class FixtureTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='monthly-gate-')
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / '2026年ABS发行台账-1016-定稿.xlsx'

    def write_source(self, rows=None, merges=(), active_other=False):
        book = Workbook()
        sheet = book.active
        sheet.title = '总表'
        sheet.append(HEADERS)
        for row in rows or [source_row()]:
            sheet.append(row)
        for area in merges:
            sheet.merge_cells(area)
        if active_other:
            book.create_sheet('其他').append(['不是台账'])
            book.active = 1
        book.save(self.path)
        book.close()

    def generate(self):
        return with_provenance(demo.compute_monthly(self.path), self.path)

    def html(self, monthly=None):
        return demo.build_document(panel.compute_data(self.path), monthly or self.generate())

    def assert_block(self, report, code=None):
        self.assertEqual(report['status'], 'FAIL', report)
        if code:
            self.assertIn(code, [i['code'] for i in report['issues'] if i['level'] == 'FAIL'], report)

    def test_real_generated_xlsx_data_and_html_pass(self):
        self.write_source(active_other=True)
        source_hash = gate.file_sha256(self.path)
        result = gate.verify_html(self.html(), self.path)
        self.assertEqual(result['status'], 'PASS')
        self.assertEqual(result['checks']['rawBids'], 1)
        self.assertEqual(result['checks']['latestRuntime']['projects'], 1)
        self.assertEqual(gate.file_sha256(self.path), source_hash)

    def test_hardcoded_five_project_arithmetic_hidden_allocations_and_zero_bid_month(self):
        rows = [source_row(), source_row(allocation=1, size=1),
                source_row(allocated='某信托', allocation=7, bidder='某信托', size=200),
                source_row(project='一月乙', date='2026-01-09', allocation=2, bidder=None, rate=None, size=None),
                source_row(project='三月丙', date='2026-03-09', allocation=1, bidder=None, rate=None, size=None),
                source_row(project='四月丁', date='2026-04-09', allocation=3, bidder=None, rate=None, size=None),
                source_row(project='十二月戊', date='2026-10-16', allocation=4, bidder=None, rate=None, size=None),
                source_row(cost='不出表', allocation=90, bidder=None, rate=None, size=None),
                source_row(allocated='某机构（自持）', allocation=90, bidder=None, rate=None, size=None),
                source_row(venue='非标', allocation=90, size=90)]
        for row in rows:
            row[HEADERS.index('规模')] = 200
        self.write_source(rows)
        monthly = self.generate()
        gate.verify_data(monthly, self.path)
        asset = monthly['assets']['赊销白条']
        bank = next(i for i in asset['institutions'] if i['name'] == '建设银行')
        self.assertEqual(asset['projectCount'], 5)
        self.assertEqual(bank['participationCount'], 3)
        self.assertEqual(bank['participationRate'], 3 / 5)
        self.assertEqual(bank['monthlySizes'][4], 4)
        self.assertEqual(bank['allocatedSize'], 13)
        self.assertEqual(asset['totalAllocatedSize'], 20)
        self.assertEqual(bank['allocationShare'], 13 / 20)
        self.assertEqual(asset['bubbleMaxSize'], 4)
        self.assertEqual(asset['fallbackCount'], 2)
        self.assertEqual(asset['counts'][3], 0)
        for field, value in [('totalAllocatedSize', 13), ('bubbleMaxSize', 200), ('projectCount', 3)]:
            wrong = copy.deepcopy(monthly)
            wrong['assets']['赊销白条'][field] = value
            self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False), 'MONTHLY_ASSET_METRIC')

    def test_entire_year_not_only_recent_three_periods_and_raw_fields(self):
        self.write_source([source_row(project=f'月{m}', date=f'2026-{m:02d}-01') for m in range(1, 11)])
        original = self.generate()
        gate.verify_data(original, self.path)
        for field, value in [('size', 100), ('rate', 0.9), ('rawInstitution', '伪造机构'),
                             ('project', '伪造项目'), ('date', '2026-02-01'), ('venue', '银行间'), ('row', 999)]:
            with self.subTest(field=field):
                wrong = copy.deepcopy(original)
                wrong['assets']['赊销白条']['institutions'][0]['months'][0][0][field] = value
                self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False), 'MONTHLY_RAW_RECONCILIATION')

    def test_fallback_per_project_and_not_after_march(self):
        self.write_source([source_row(project='一月有报价', date='2026-01-09'),
                           source_row(project='一月无报价', date='2026-01-09', bidder=None, size=None, rate=None),
                           source_row(project='四月无报价', date='2026-04-09', bidder=None, size=None, rate=None)])
        monthly = self.generate()
        gate.verify_data(monthly, self.path)
        inst = monthly['assets']['赊销白条']['institutions'][0]
        self.assertEqual(inst['monthlyTracks'][0], 'mixed')
        self.assertEqual(inst['monthlyBidSizes'][0], 3)
        self.assertEqual(inst['monthlyFallbackSizes'][0], 2)
        self.assertEqual({r['track'] for r in inst['months'][0]}, {'bids', 'allocations'})
        self.assertEqual(monthly['fallbackRows'], [3])
        wrong = copy.deepcopy(monthly)
        fallback = copy.deepcopy(inst['months'][0][-1])
        fallback.update(row=4, date='2026-04-09', project='四月无报价')
        wrong['assets']['赊销白条']['institutions'][0]['months'][3].append(fallback)
        wrong['fallbackRows'].append(4)
        self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False), 'MONTHLY_FALLBACK')

    def test_partial_bid_evidence_cannot_unlock_january_fallback(self):
        self.write_source([source_row(date='2026-01-09', rate=None)])
        monthly = self.generate()
        self.assertEqual(monthly['fallbackRows'], [])
        result = gate.verify_data(monthly, self.path, raise_on_fail=False)
        self.assertEqual(result['status'], 'FAIL')
        self.assertTrue(any(i['code'] == 'MONTHLY_BID_PARTIAL_OR_ZERO' and i.get('blocksRelease') for i in result['issues']))

    def test_institution_metrics_sort_scope_and_visible_bubbles_block(self):
        self.write_source([source_row(), source_row(bidder='某银行', allocated='某银行', allocation=4, size=10),
                           source_row(bidder='某信托', allocated='某信托', allocation=2, size=50)])
        original = self.generate()
        gate.verify_data(original, self.path)
        for field, value in [('participationCount', 2), ('participationRate', 0.1), ('monthCount', 11),
                             ('allocatedSize', 99), ('allocationShare', 0.1), ('group', '公募基金'),
                             ('monthlySizes', [99] * 12), ('monthlyTracks', ['allocations'] * 12)]:
            wrong = copy.deepcopy(original)
            wrong['assets']['赊销白条']['institutions'][0][field] = value
            self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False), 'MONTHLY_INSTITUTION_METRIC')
        wrong = copy.deepcopy(original)
        wrong['assets']['赊销白条']['institutions'].reverse()
        self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False), 'MONTHLY_INSTITUTION_SCOPE_SORT')
        for field, value in [('visibleGroups', list(gate.GROUPS)), ('visibleCounts', [99] * 12),
                             ('counts', [99] * 12), ('projectCounts', [99] * 12)]:
            wrong = copy.deepcopy(original)
            wrong['assets']['赊销白条'][field] = value
            self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False), 'MONTHLY_ASSET_METRIC')

    def test_numeric_boolean_nan_schema_and_unknown_asset_fail_closed(self):
        self.write_source()
        original = self.generate()
        for value in (True, float('nan'), float('inf'), 'NaN'):
            wrong = copy.deepcopy(original)
            wrong['assets']['赊销白条']['institutions'][0]['monthlySizes'][4] = value
            self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False))
        for value in ([], None, {'额外资产': {}}):
            wrong = copy.deepcopy(original)
            wrong['assets'] = value
            self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False))

    def test_rule_alias_classification_fingerprints_and_source_tamper_fail(self):
        self.write_source()
        original = self.generate()
        for field in ('sourceSha256', 'demoAliasesSha256', 'fundRulesSha256', 'classificationSha256'):
            wrong = copy.deepcopy(original)
            wrong[field] = '0' * 64
            self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False), 'MONTHLY_FINGERPRINT')
        for field in ('source', 'rulesSha256', 'scope', 'dependencies'):
            wrong = copy.deepcopy(original)
            wrong['provenance'][field] = 'tampered'
            self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False), 'MONTHLY_PROVENANCE')
        wrong = copy.deepcopy(original)
        wrong['provenance']['dependencies']['scripts/lab/monthly_bid_institution_aliases.json'] = '0' * 64
        self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False), 'MONTHLY_PROVENANCE')
        # 真正修改临时xlsx，同名旧产物即使改写来源哈希也不能绕过逐笔核对。
        self.write_source([source_row(size=99)])
        original['sourceSha256'] = gate.file_sha256(self.path)
        original['provenance'] = gate.build_provenance(self.path)
        self.assert_block(gate.verify_data(original, self.path, raise_on_fail=False), 'MONTHLY_RAW_RECONCILIATION')

    def test_actual_alias_dependency_change_stales_existing_artifact(self):
        self.write_source()
        aliases = Path(self.temp.name) / 'aliases.json'
        aliases.write_bytes(gate.DEMO_ALIASES.read_bytes())
        fingerprints = gate.dependency_fingerprints
        alias_key = gate.DEMO_ALIASES.relative_to(gate.ROOT).as_posix()

        def temporary_alias_fingerprints():
            result = fingerprints()
            result[alias_key] = gate.file_sha256(aliases)
            return result

        # 保留真实依赖计算和仓库根；只用临时别名字节替换同一个指纹项。
        with mock.patch.object(gate, 'dependency_fingerprints', side_effect=temporary_alias_fingerprints):
            original = self.generate()
            self.assertEqual(gate.verify_data(original, self.path)['status'], 'PASS')
            aliases.write_bytes(aliases.read_bytes() + b'\n')
            self.assert_block(gate.verify_data(original, self.path, raise_on_fail=False), 'MONTHLY_PROVENANCE')

    def test_independent_reader_not_using_production_iterators_or_compute(self):
        self.write_source([source_row(), source_row(allocation=1, size=1)])
        original = self.generate()
        with mock.patch.object(recent_qc, 'read_source', side_effect=AssertionError('不能复用生产读取')), \
             mock.patch.object(demo, 'compute_monthly', side_effect=AssertionError('不能自证')), \
             mock.patch.object(panel, 'compute_data', side_effect=AssertionError('不能自证')):
            self.assertEqual(gate.verify_data(original, self.path)['status'], 'PASS')

    def test_independent_merged_scope_and_raw_column_merges(self):
        self.write_source([source_row(), source_row(allocation=1, size=1)], merges=('D2:D3', 'E2:E3', 'F2:F3', 'L2:L3'))
        original = self.generate()
        self.assertEqual(gate.verify_data(original, self.path)['status'], 'PASS')
        self.write_source([source_row(), source_row()], merges=('U2:U3',))
        self.assert_block(gate.verify_data(original, self.path, raise_on_fail=False), 'MONTHLY_SCHEMA_OR_SOURCE')

    def test_key_formulas_missing_cache_block_but_optional_formula_warns(self):
        for column in ('V', 'W', 'Y'):
            self.write_source()
            original = self.generate()
            book = load_workbook(self.path)
            book['总表'][f'{column}2'] = '=1'
            book.save(self.path)
            book.close()
            self.assert_block(gate.verify_data(original, self.path, raise_on_fail=False), 'MONTHLY_SCHEMA_OR_SOURCE')
        self.write_source()
        book = load_workbook(self.path)
        book['总表']['R2'] = '=20'
        book.save(self.path)
        book.close()
        report = gate.verify_html(self.html(), self.path)
        self.assertEqual(report['status'], 'PASS')
        self.assertIn('LAYER_SIZE_UNCACHED', [i['code'] for i in report['latestWarnings']])
        self.assertEqual(report['monthlyBlocks'], [])

    def test_partial_wxy_headers_and_explicit_self_hold_preserved_as_warn(self):
        self.write_source([source_row(project='一月无申购', date='2026-01-09',
                                      bidder='穿透机构', size='申购规模', rate='申购利率'),
                           source_row(project='五月甲', bidder='某小贷', rate='自持')])
        monthly = self.generate()
        report = gate.verify_data(monthly, self.path)
        self.assertEqual(monthly['fallbackRows'], [2])
        self.assertEqual(report['status'], 'PASS')
        self.assertEqual(len(monthly['excluded']), 1)

    def test_real_inline_js_and_css_tampering_and_extra_override_block(self):
        self.write_source()
        html = self.html()
        gate.verify_html(html, self.path)
        for path, code in [(LAB / 'monthly_bid_institutions_demo.js', 'HTML_INLINE_JS'),
                           (SCRIPTS / 'latest_bookkeeping_panel.js', 'HTML_INLINE_JS'),
                           (LAB / 'monthly_bid_institutions_demo.css', 'HTML_INLINE_CSS'),
                           (SCRIPTS / 'latest_bookkeeping_panel.css', 'HTML_INLINE_CSS')]:
            with self.subTest(resource=path.name):
                resource = path.read_text(encoding='utf-8')
                changed = resource.replace('26*Math.sqrt', '260*Math.sqrt') if path.suffix == '.js' and '26*Math.sqrt' in resource else resource.replace('use strict', 'use sloppy', 1)
                tampered = html.replace(resource, changed, 1)
                if path.suffix == '.css':
                    changed = resource.replace('#latest-bookkeeping', '#tampered-bookkeeping', 1)
                    tampered = html.replace(resource, changed, 1)
                self.assert_block(gate.verify_html(tampered, self.path, raise_on_fail=False), code)
        for extra, code in [('<style>#latest-bookkeeping .lb-monthly{display:none}</style>', 'HTML_EXTRA_COMPONENT_CSS'),
                            ('<script>document.getElementById("lb-monthly-rows").innerHTML="";</script>', 'HTML_EXTRA_COMPONENT_JS')]:
            self.assert_block(gate.verify_html(html.replace('</body>', extra + '</body>'), self.path, raise_on_fail=False), code)

    def test_real_body_controls_and_resource_execution_attributes_block(self):
        self.write_source()
        html = self.html()
        for old, new, code in [
                ('<select id="lb-project">', '<select id="lb-project" disabled>', 'HTML_COMPONENT_STRUCTURE'),
                ('<thead id="lb-monthly-head">', '<div id="lb-monthly-head">', 'HTML_COMPONENT_STRUCTURE'),
                ('<div id="lb-source"', '<div hidden id="lb-source"', 'HTML_COMPONENT_STRUCTURE'),
                ('data-view="bids"', 'data-view="changed"', 'HTML_DETAIL_CONTROLS'),
                ('<style>', '<style media="print">', 'HTML_INLINE_CSS')]:
            with self.subTest(change=old):
                self.assertIn(old, html)
                self.assert_block(gate.verify_html(html.replace(old, new, 1), self.path, raise_on_fail=False), code)

    def test_duplicate_json_blocks_nonfinite_duplicate_keys_and_structure_block(self):
        self.write_source()
        html = self.html()
        for identifier in ('lb-data', 'lb-monthly-data'):
            block = re.search(r'<script[^>]*id="' + identifier + r'"[^>]*>.*?</script>', html, re.S).group()
            self.assert_block(gate.verify_html(html + block, self.path, raise_on_fail=False), 'HTML_GATE')
        wrong = html.replace('id="lb-monthly-head"', 'id="missing-head"')
        self.assert_block(gate.verify_html(wrong, self.path, raise_on_fail=False), 'HTML_COMPONENT_STRUCTURE')
        for payload in ('{"year":2026,"year":2025}', '{"year":NaN}'):
            wrong = re.sub(r'(<script id="lb-monthly-data" type="application/json">).*?(</script>)',
                           lambda m: m[1] + payload + m[2], html, flags=re.S)
            self.assert_block(gate.verify_html(wrong, self.path, raise_on_fail=False), 'HTML_GATE')

    def test_source_mismatch_old_latest_gate_and_wrapped_html(self):
        self.write_source()
        html = self.html()
        # 集成代理已新增路由，旧组件的整段子串验证依然独立严格。
        legacy = getattr(recent_qc, '_verify_legacy_html', recent_qc.verify_html)
        with self.assertRaises(ValueError):
            legacy(html, self.path)
        self.assertEqual(gate.verify_html(html, self.path)['status'], 'PASS')
        latest, _ = gate.extract_data(html)
        latest['source'] = '其他台账.xlsx'
        self.assert_block(gate.verify_html(replace_payload(html, 'lb-data', latest), self.path, raise_on_fail=False), 'HTML_SOURCE_MISMATCH')
        # 月表包装保持现有子数据门禁；篡改最近项目明细也必须FAIL。
        latest, _ = gate.extract_data(html)
        latest['assets'][0]['projects'][0]['tiers'][0]['bids'][0]['bidSize'] = 100
        self.assert_block(gate.verify_html(replace_payload(html, 'lb-data', latest), self.path, raise_on_fail=False), 'DETAIL_RECONCILIATION')

    def test_current_render_component_exactly_checks_wrapped_body(self):
        self.write_source()
        monthly, latest = self.generate(), panel.compute_data(self.path)
        component = demo.render_component(latest, monthly)
        html = '<html><head><style>' + demo.component_css() + '</style></head><body><nav>综合导航</nav><section>' + component + '</section></body></html>'
        self.assertEqual(gate.verify_html(html, self.path)['status'], 'PASS')
        changed = html.replace('最近项目簿记', '被篡改的标题', 1)
        self.assert_block(gate.verify_html(changed, self.path, raise_on_fail=False), 'HTML_CURRENT_RENDERER')

    def test_reaudit_inactive_wrappers_reject_but_controlled_tab_wrapper_passes(self):
        self.write_source()
        monthly = self.generate()
        demo._attach_qc(monthly, gate.verify_data(monthly, self.path))
        html = self.html(monthly)
        latest, _ = gate.extract_data(html)
        component = demo.render_component(latest, monthly)
        self.assertEqual(gate.verify_html(html, self.path)['status'], 'PASS')
        for wrapper in ('<div hidden>', '<template>', '<div inert>',
                        '<div style="display: none !important">',
                        '<div style="visibility:hidden">'):
            with self.subTest(wrapper=wrapper):
                end = '</template>' if wrapper == '<template>' else '</div>'
                altered = html.replace('<body>', '<body>' + wrapper).replace('</body>', end + '</body>')
                # 完整、正确、当前版本component保留，祖先使实际页面隐藏或不执行。
                self.assertTrue(component in altered, '反例必须保留完整当前组件')
                self.assertEqual(gate.verify_html(altered, self.path, raise_on_fail=False)['status'], 'FAIL')
                with self.assertRaises(ValueError):
                    recent_qc.verify_html(altered, self.path)
        valid = html.replace('<body>', '<body><div class="panel" id="controlled-tab">').replace('</body>', '</div></body>')
        self.assertEqual(gate.verify_html(valid, self.path)['status'], 'PASS')
        self.assertEqual(recent_qc.verify_html(valid, self.path)['status'], 'PASS')

    def test_reaudit_healthy_source_carrying_blocked_embedded_qc_is_blocked(self):
        self.write_source()
        healthy = self.generate()
        self.assertEqual(gate.verify_data(healthy, self.path)['status'], 'PASS')
        for metadata in (
                {'status': 'FAIL', 'releaseStatus': 'BLOCK', 'releaseBlocked': True,
                 'issues': [{'level': 'FAIL', 'code': 'PROBE_FAIL', 'message': '不可放行'}]},
                {'status': 'PASS', 'releaseStatus': 'BLOCK', 'releaseBlocked': True, 'issues': []},
                {'status': 'PASS', 'issues': [{'level': 'WARN', 'blocksRelease': True, 'code': 'PROBE', 'message': '阻断'}]}):
            monthly = copy.deepcopy(healthy)
            monthly['auditQc'] = metadata
            html = self.html(monthly)
            self.assert_block(gate.verify_html(html, self.path, raise_on_fail=False), 'HTML_EMBEDDED_QC_STATE')
            with self.assertRaises(gate.MonthlyBookkeepingGateError):
                gate.verify_html(html, self.path)
            with self.assertRaises(ValueError):
                recent_qc.verify_html(html, self.path)
        demo._attach_qc(healthy, gate.verify_data(healthy, self.path))
        self.assertEqual(gate.verify_html(self.html(healthy), self.path)['status'], 'PASS')

    def test_reaudit_selector_history_only_asset_passes_current_year_omission_blocks(self):
        # 普通工作簿名允许其他年份；定稿文件名限定年份的旧QC规则仍保持。
        self.path = Path(self.temp.name) / 'year-missing.xlsx'
        self.write_source([source_row(project='今年资产', date='2026-04-03'),
                           source_row(project='历史资产', date='2025-12-03', asset='京企贷')])
        monthly, latest = self.generate(), panel.compute_data(self.path)
        self.assertEqual(set(monthly['assets']), {'赊销白条'})
        self.assertEqual({a['name'] for a in latest['assets']}, {'赊销白条', '京企贷'})
        demo._attach_qc(monthly, gate.verify_data(monthly, self.path))
        html = demo.build_document(latest, monthly)
        self.assertEqual(gate.verify_html(html, self.path)['status'], 'PASS')
        self.assertEqual(recent_qc.verify_html(html, self.path)['status'], 'PASS')
        omitted = copy.deepcopy(latest)
        omitted['assets'] = [a for a in omitted['assets'] if a['name'] == '京企贷']
        self.assert_block(gate.verify_html(demo.build_document(omitted, monthly), self.path, raise_on_fail=False), 'HTML_ASSET_SELECTOR_SCOPE')
        missing = copy.deepcopy(monthly)
        missing['assets'] = {}
        self.assert_block(gate.verify_data(missing, self.path, raise_on_fail=False), 'MONTHLY_ASSET_SCOPE')

    def test_reaudit_payload_numeric_strings_reject_excel_numeric_strings_remain_valid(self):
        row = source_row(allocation=2, size=3)
        for name in ('规模', '成本', '认购份额', '申购利率', '申购规模'):
            row[HEADERS.index(name)] = str(row[HEADERS.index(name)])
        self.write_source([row, source_row(allocated='某信托', allocation=8, bidder='某信托', size=20)])
        monthly = self.generate()
        demo._attach_qc(monthly, gate.verify_data(monthly, self.path))
        self.assertEqual(gate.verify_html(self.html(monthly), self.path)['status'], 'PASS')
        def strings(value):
            if type(value) is float:
                return str(value)
            if isinstance(value, dict):
                return {key: strings(item) for key, item in value.items()}
            if isinstance(value, list):
                return [strings(item) for item in value]
            return value
        wrong = strings(monthly)
        # 真实JSON字符串数值与源金额相同，指纹和renderer也不变，必须仍阻断。
        self.assertEqual(wrong['provenance'], monthly['provenance'])
        self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False), 'MONTHLY_INSTITUTION_METRIC')
        html = self.html(wrong)
        self.assert_block(gate.verify_html(html, self.path, raise_on_fail=False), 'MONTHLY_INSTITUTION_METRIC')
        with self.assertRaises(ValueError):
            recent_qc.verify_html(html, self.path)

    def test_cli_pass_fail_reports_and_refuses_input_overwrite(self):
        self.write_source()
        path, report = Path(self.temp.name) / 'demo.html', Path(self.temp.name) / 'report.json'
        path.write_text(self.html(), encoding='utf-8')
        command = [sys.executable, '-B', str(LAB / 'monthly_bookkeeping_qc.py'),
                   '--dashboard', str(path), '--ledger', str(self.path), '--report', str(report)]
        good = subprocess.run(command, capture_output=True, text=True, timeout=30)
        self.assertEqual(good.returncode, 0, good.stdout + good.stderr)
        self.assertEqual(json.loads(report.read_text())['status'], 'PASS')
        content = path.read_text()
        _, monthly = gate.extract_data(content)
        monthly['assets']['赊销白条']['totalAllocatedSize'] = 100
        path.write_text(replace_payload(content, 'lb-monthly-data', monthly), encoding='utf-8')
        bad = subprocess.run(command, capture_output=True, text=True, timeout=30)
        self.assertEqual(bad.returncode, 1, bad.stdout + bad.stderr)
        self.assertEqual(json.loads(report.read_text())['status'], 'FAIL')
        saved_hash = gate.file_sha256(self.path)
        rejected = subprocess.run(command[:-1] + [str(self.path)], capture_output=True, text=True, timeout=30)
        self.assertEqual(rejected.returncode, 2)
        self.assertEqual(gate.file_sha256(self.path), saved_hash)

    def test_old_source_type_range_orphan_validation_cannot_hide_behind_recent_pass(self):
        legal_rows = [source_row(project='历史一月', date='2026-01-09')] + [
            source_row(project=f'近期{month}', date=f'2026-{month:02d}-01') for month in (8, 9, 10)]
        cases = [('申购利率', True), ('申购利率', 'Infinity'), ('申购利率', 1.1),
                 ('申购利率', -0.1), ('成本', True), ('成本', 'NaN'), ('成本', 1.1),
                 ('认购份额', True), ('认购份额', -1), ('认购份额', '错误金额'),
                 ('申购规模', True), ('申购规模', -1), ('申购规模', '#VALUE!'),
                 ('规模', True), ('规模', -1), ('对应金额（亿）', 'NaN'),
                 ('穿透机构', None), ('认购机构', None), ('穿透机构', 123)]
        for field, value in cases:
            with self.subTest(field=field, value=value):
                rows = copy.deepcopy(legal_rows)
                rows[0][HEADERS.index(field)] = value
                self.write_source(rows)
                # 旧最近三期没有历史异常：全年源门禁必须独立发现。
                self.assertEqual(panel.compute_data(self.path)['qc']['status'], 'PASS')
                result = gate.verify_source(self.path, raise_on_fail=False)
                self.assertEqual(result['status'], 'FAIL', result)
                self.assertEqual(result['releaseStatus'], 'BLOCK')

    def test_source_project_layer_limits_include_self_hold_and_preserve_duplicate_rows(self):
        rows = [source_row(project='历史一月', date='2026-01-09', allocation=2),
                source_row(project='历史一月', date='2026-01-09', allocated='京东自持', allocation=9,
                           bidder=None, size=None, rate=None, cost='自持')]
        for row in rows:
            row[HEADERS.index('规模')] = 10
            row[HEADERS.index('对应金额（亿）')] = 3
        rows += [source_row(project=f'近期{month}', date=f'2026-{month:02d}-01') for month in (8, 9, 10)]
        self.write_source(rows)
        monthly = self.generate()
        self.assertEqual(monthly['assets']['赊销白条']['totalAllocatedSize'], 8)
        result = gate.verify_data(monthly, self.path, raise_on_fail=False)
        self.assert_block(result, 'MONTHLY_ALLOCATION_PROJECT_OVER')
        self.assert_block(result, 'MONTHLY_ALLOCATION_LAYER_OVER')
        # 不把可疑重复自动去重；总额仍按原始两行4亿元。
        self.write_source([source_row(allocation=2), source_row(allocation=2)])
        result = gate.verify_source(self.path)
        self.assertIn('MONTHLY_DUPLICATE_ALLOCATION', [i['code'] for i in result['issues']])
        monthly = self.generate()
        self.assertEqual(monthly['assets']['赊销白条']['totalAllocatedSize'], 4)

    def test_unknown_allocation_retains_known_amounts_and_nulls_entire_asset_shares(self):
        self.write_source([source_row(allocation=2), source_row(allocated='建设银行', allocation=None),
                           source_row(allocated='某信托', allocation=None, bidder='某银行', size=5),
                           source_row(allocated='某银行', allocation=3, bidder='某银行', size=5),
                           source_row(allocated='京东自持', allocation=None, bidder=None, rate=None, size=None)])
        monthly = self.generate()
        result = gate.verify_data(monthly, self.path)
        self.assertEqual(result['status'], 'PASS')
        asset = monthly['assets']['赊销白条']
        bank = next(i for i in asset['institutions'] if i['name'] == '建设银行')
        self.assertEqual(bank['knownAllocatedSize'], 2)
        self.assertIsNone(bank['allocatedSize'])
        self.assertFalse(bank['allocationComplete'])
        self.assertEqual(bank['unknownAllocationRows'], [3])
        self.assertEqual(asset['totalKnownAllocatedSize'], 5)
        self.assertIsNone(asset['totalAllocatedSize'])
        self.assertFalse(asset['allocationComplete'])
        self.assertEqual(asset['unknownAllocationRows'], [3, 4])
        self.assertTrue(all(i['allocationShare'] is None for i in asset['institutions']))
        self.assertEqual([r['row'] for r in monthly['unknownAllocationRecords']], [3, 4])
        wrong = copy.deepcopy(monthly)
        wrong['assets']['赊销白条']['totalAllocatedSize'] = 5
        self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False), 'MONTHLY_ASSET_METRIC')
        wrong = copy.deepcopy(monthly)
        wrong['assets']['赊销白条']['institutions'][0]['allocationShare'] = 0
        self.assert_block(gate.verify_data(wrong, self.path, raise_on_fail=False), 'MONTHLY_INSTITUTION_METRIC')

    def test_raw_zero_partial_and_no_bid_evidence_have_distinct_fallback_rules(self):
        for updates in ({'申购规模': 0}, {'申购利率': None}, {'申购规模': None}):
            row = source_row(project='一月', date='2026-01-09')
            for field, value in updates.items():
                row[HEADERS.index(field)] = value
            self.write_source([row])
            monthly = self.generate()
            self.assertEqual(monthly['fallbackRows'], [])
            result = gate.verify_data(monthly, self.path, raise_on_fail=False)
            self.assertEqual(result['releaseStatus'], 'BLOCK')
            self.assertTrue(any(i['code'] == 'MONTHLY_BID_PARTIAL_OR_ZERO' and i.get('blocksRelease') for i in result['issues']))
        self.write_source([source_row(date='2026-01-09', bidder=None, rate=None, size=None)])
        monthly = self.generate()
        self.assertEqual(monthly['fallbackRows'], [2])
        gate.verify_data(monthly, self.path)

    def test_new_public_fund_dependencies_do_not_restore_ambiguous_bare_aliases(self):
        self.write_source([source_row(bidder='华夏'), source_row(bidder='华夏基金'),
                           source_row(bidder='前海开源'), source_row(bidder='国投资管（投行）'),
                           source_row(bidder='华泰资管（自营）')])
        monthly = self.generate()
        gate.verify_data(monthly, self.path)
        institutions = {i['name']: i['group'] for i in monthly['assets']['赊销白条']['institutions']}
        self.assertEqual(institutions['华夏'], '其他')
        self.assertEqual(institutions['华夏基金'], '公募基金')
        self.assertEqual(institutions['前海开源基金'], '公募基金')
        self.assertEqual(institutions['国投资管（投行）'], '券商')
        self.assertEqual(institutions['华泰资管（自营）'], '券商')
        dependencies = gate.dependency_fingerprints()
        self.assertIn('scripts/lab/monthly_bid_classification.json', dependencies)
        self.assertFalse(any('20261017_公募历史获配' in path for path in dependencies))


class BaselineArtifactTests(unittest.TestCase):
    """合成历史超限源表；不依赖审计目录或纠错后会变动的真实台账。"""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='monthly-history-')
        self.addCleanup(self.temp.cleanup)
        self.ledger = Path(self.temp.name) / '2026年ABS发行台账-1016-定稿.xlsx'
        # 独立人工期望：一月三行共12.7，超过J10和R0.3；重复行必须保留。
        rows = [source_row(project='历史一月', date='2026-01-09', allocation=amount,
                           bidder=None, rate=None, size=None) for amount in (6.2, 6.2, 0.3)]
        for row in rows:
            row[9] = 10
            row[17] = 0.3
        for month in (8, 9, 10):
            row = source_row(project=f'近期{month}', date=f'2026-{month:02d}-01')
            row[9] = 2
            row[17] = 2
            rows.append(row)
        book = Workbook()
        sheet = book.active
        sheet.title = '总表'
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(row)
        book.save(self.ledger)
        book.close()
        self.source_hash = gate.file_sha256(self.ledger)

    def tearDown(self):
        self.assertEqual(gate.file_sha256(self.ledger), self.source_hash)

    def test_missing_new_provenance_preserves_latest_qc(self):
        latest, monthly = panel.compute_data(self.ledger), demo.compute_monthly(self.ledger)
        demo._attach_qc(monthly, gate.verify_data(monthly, self.ledger, raise_on_fail=False))
        del monthly['provenance']
        report = gate.verify_html(demo.build_document(latest, monthly), self.ledger, raise_on_fail=False)
        self.assertEqual(report['status'], 'FAIL')
        self.assertIn('MONTHLY_PROVENANCE', [i['code'] for i in report['monthlyBlocks']])
        self.assertEqual(report['latestQC']['status'], 'PASS')
        self.assertEqual(report['latestQC']['checks']['projects'], 3)
        self.assertEqual(report['latestWarnings'], [])

    def test_whole_year_records_preserved_but_source_conflicts_block(self):
        monthly = demo.compute_monthly(self.ledger)
        report = gate.verify_data(monthly, self.ledger, raise_on_fail=False)
        self.assertEqual(report['status'], 'FAIL')
        failures = [i['code'] for i in report['issues'] if i['level'] == 'FAIL']
        self.assertCountEqual(failures, ['MONTHLY_ALLOCATION_LAYER_OVER', 'MONTHLY_ALLOCATION_PROJECT_OVER'])
        self.assertEqual(report['checks']['projects'], 4)
        self.assertEqual(report['checks']['rawBids'], 3)
        self.assertEqual(report['checks']['fallbackAllocations'], 3)
        self.assertEqual(report['checks']['allocationRecords'], 6)
        asset = monthly['assets']['赊销白条']
        self.assertEqual(asset['totalAllocatedSize'], 18.7)
        self.assertEqual(monthly['fallbackRows'], [2, 3, 4])
        self.assertEqual([r['row'] for r in monthly['allocationRecords']], [2, 3, 4, 5, 6, 7])
        for mutation, expected_code in [('size', 'MONTHLY_RAW_RECONCILIATION'),
                                         ('denominator', 'MONTHLY_ASSET_METRIC'),
                                         ('alias', 'MONTHLY_FINGERPRINT')]:
            wrong = copy.deepcopy(monthly)
            asset = wrong['assets']['赊销白条']
            if mutation == 'size':
                bid = next(r for i in asset['institutions'] for m in i['months'] for r in m if r['track'] == 'bids')
                bid['size'] += 1
            elif mutation == 'denominator':
                asset['totalAllocatedSize'] += 1
            else:
                wrong['demoAliasesSha256'] = '0' * 64
            result = gate.verify_data(wrong, self.ledger, raise_on_fail=False)
            self.assertEqual(result['status'], 'FAIL', mutation)
            self.assertIn(expected_code, [i['code'] for i in result['issues'] if i['level'] == 'FAIL'])

    def test_source_independent_arithmetic_and_structured_rows(self):
        report = gate.verify_source(self.ledger, raise_on_fail=False)
        self.assertEqual(report['releaseStatus'], 'BLOCK')
        layers = [i for i in report['issues'] if i['code'] == 'MONTHLY_ALLOCATION_LAYER_OVER']
        projects = [i for i in report['issues'] if i['code'] == 'MONTHLY_ALLOCATION_PROJECT_OVER']
        self.assertEqual(len(layers), 1)
        self.assertEqual(len(projects), 1)
        project = projects[0]
        self.assertEqual(project['asset'], '赊销白条')
        self.assertEqual(project['project'], '历史一月')
        self.assertEqual(project['date'], '2026-01-09')
        self.assertEqual(project['rows'], [2, 3, 4])
        self.assertEqual(project['coordinates']['allocationAmount'], ['V2', 'V3', 'V4'])
        self.assertIn('12.7 > J10', project['message'])
        self.assertEqual(layers[0]['rows'], [2, 3, 4])
        self.assertEqual(layers[0]['coordinates']['layerSize'], ['R2', 'R3', 'R4'])
        self.assertIn('12.7 > R0.3', layers[0]['message'])
        duplicates = [i for i in report['issues'] if i['code'] == 'MONTHLY_DUPLICATE_ALLOCATION']
        self.assertEqual(len(duplicates), 1)
        self.assertEqual(duplicates[0]['level'], 'WARN')
        self.assertIn('相同U/V/成本行[2, 3]', duplicates[0]['message'])
        totals = report['checks']['sourceConservation']['projectTotals']
        self.assertEqual([p['uvIncludingSelfHeld'] for p in totals], ['12.7', '2', '2', '2'])

    def test_source_preview_renders_block_and_never_passes_release(self):
        latest, monthly = panel.compute_data(self.ledger), demo.compute_monthly(self.ledger)
        self.assertEqual(latest['qc']['status'], 'PASS')
        data_report = gate.verify_data(monthly, self.ledger, raise_on_fail=False)
        demo._attach_qc(monthly, data_report)
        content = demo.build_document(latest, monthly)
        report = gate.verify_html(content, self.ledger, raise_on_fail=False)
        self.assertEqual(report['releaseStatus'], 'BLOCK')
        self.assertEqual(report['summary']['failCount'], 2, report['issues'])
        self.assertEqual(report['latestWarnings'], [])
        self.assertEqual(report['checks']['latestRuntime']['projects'], 3)
        self.assertIn('data-qc-status="FAIL"', content)
        with self.assertRaises(gate.MonthlyBookkeepingGateError):
            gate.verify_html(content, self.ledger)
        bad = content.replace('data-qc-status="FAIL"', 'data-qc-status="PASS"')
        altered = gate.verify_html(bad, self.ledger, raise_on_fail=False)
        self.assertIn('HTML_RELEASE_BLOCK_BANNER', [i['code'] for i in altered['issues']])


if __name__ == '__main__':
    unittest.main()
