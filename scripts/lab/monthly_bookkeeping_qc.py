"""2026 月度簿记 Demo 独立门禁；不修改台账，不自动接入正式 generator。

生成端可写入 monthly['provenance'] = build_provenance(ledger)。核对端不调用
compute_monthly/compute_data；独立 openpyxl 读取与 Decimal/集合对账。
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
import hashlib
from html.parser import HTMLParser
import json
from pathlib import Path
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

SCRIPTS = Path(__file__).resolve().parents[1]
ROOT = SCRIPTS.parent
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

SCHEMA_VERSION = 'monthly-bookkeeping-demo/v1'
YEAR = 2026
GROUPS = ('理财子', '银行', '公募基金', '券商', '其他')
VENUES = frozenset({'上交所', '深交所', '银行间'})
THREE_GROUP_ASSETS = frozenset({'赊销白条', '金条白取混发', '信托金条', '白条取现',
                               '信托企业主贷', '小贷企业主贷', '保理', '金采'})
ASSET_ALIASES = {'信托白取': '白条取现', '金条&白取混发': '金条白取混发'}
CLASSIFICATION = ROOT / 'data/机构分配字典.json'
FUND_RULES = Path(__file__).with_name('monthly_bid_classification.json')
DEMO_ALIASES = Path(__file__).with_name('monthly_bid_institution_aliases.json')
SCOPE = {'year': YEAR, 'venues': sorted(VENUES), 'fallbackMonths': [1, 2, 3],
         'allocationDenominator': '全年U/V全部板块，排除明确自持或不出表',
         'threeGroupAssets': sorted(THREE_GROUP_ASSETS)}
FIELDS = ('资产类型', '项目名称', '簿记时间', '发行场所', '成本', '认购机构',
          '认购份额', '申购利率', '穿透机构', '申购规模', '规模', '分层情况')
DETAIL_HEADERS = ('成本', '认购机构', '认购份额', '申购利率', '穿透机构', '申购规模')


class MonthlyBookkeepingGateError(ValueError):
    def __init__(self, report):
        self.report = report
        super().__init__('月度 Demo 门禁失败: ' + '; '.join(
            f"{i['code']}: {i['message']}" for i in report['issues'] if i['level'] == 'FAIL')[:2000])


def file_sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def dependency_fingerprints():
    """包括规则、分类、两层别名及真正内联的资源；不执行公募底稿。"""
    paths = [SCRIPTS / name for name in (
        'abs_common.py', 'institution_identity.py', 'latest_bookkeeping_panel.py',
        'latest_bookkeeping_panel.js', 'latest_bookkeeping_panel.css',
        'latest_bookkeeping_panel_body.html', 'latest_bookkeeping_qc.py',
        'latest_bookkeeping_runtime.py', 'verify_latest_bookkeeping_runtime.js',
        'lab/monthly_bid_institutions_demo.py', 'lab/monthly_bid_institutions_demo.js',
        'lab/monthly_bid_institutions_demo.css', 'lab/monthly_bookkeeping_qc.py')]
    paths += [CLASSIFICATION, FUND_RULES, DEMO_ALIASES,
              ROOT / 'data/institution_alias_registry.json']
    return {path.relative_to(ROOT).as_posix(): file_sha256(path) for path in sorted(set(paths))}


def rules_sha256():
    return hashlib.sha256(json.dumps(dependency_fingerprints(), sort_keys=True,
                                    separators=(',', ':')).encode()).hexdigest()


def build_provenance(ledger):
    dependencies = dependency_fingerprints()
    return {'schema': SCHEMA_VERSION, 'source': Path(ledger).name,
            'sourceSha256': file_sha256(ledger), 'scope': json.loads(json.dumps(SCOPE)),
            'dependencies': dependencies,
            'rulesSha256': hashlib.sha256(json.dumps(dependencies, sort_keys=True,
                                separators=(',', ':')).encode()).hexdigest()}


def _report(ledger):
    return {'schema': SCHEMA_VERSION, 'checkedAt': datetime.now().astimezone().isoformat(timespec='seconds'),
            'source': str(Path(ledger).resolve()), 'status': 'PASS', 'issues': [], 'checks': {},
            'integration': 'Demo专门门禁；正式generator未自动启用'}


def _issue(report, code, message, *, level='FAIL', origin='monthly', blocks_release=False):
    item = {'level': level, 'origin': origin, 'code': code, 'message': message}
    source_rows = re.search(r'总表(\[[\d, ]+\]|\d+)', message)
    if source_rows:
        rows = [int(r) for r in re.findall(r'\d+', source_rows[1])]
        item['rows'] = rows
        item['sheet'] = '总表'
        item['coordinates'] = {label: [f'{column}{r}' for r in rows] for label, column in (
            ('projectSize', 'J'), ('layerSize', 'R'), ('allocationInstitution', 'U'),
            ('allocationAmount', 'V'), ('bidRate', 'W'), ('bidInstitution', 'X'), ('bidSize', 'Y'))}
    context = re.match(r'(.+?)/(\d{4}-\d{2}-\d{2})/(.+?)/(上交所|深交所|银行间)(?:/([^ ]+))? 总表', message)
    if context:
        item.update({'asset': context[1], 'date': context[2], 'project': context[3], 'venue': context[4]})
        if context[5]:
            item['layer'] = context[5]
    if blocks_release:
        item['blocksRelease'] = True
    report['issues'].append(item)


def _finish(report, raise_on_fail):
    report['monthlyBlocks'] = [i for i in report['issues'] if i['origin'] == 'monthly' and (i['level'] == 'FAIL' or i.get('blocksRelease'))]
    report['latestWarnings'] = [i for i in report['issues'] if i['origin'] == 'latest' and i['level'] == 'WARN']
    report['status'] = 'FAIL' if any(i['level'] == 'FAIL' or i.get('blocksRelease') for i in report['issues']) else 'PASS'
    report['releaseStatus'] = 'BLOCK' if report['status'] == 'FAIL' else 'PASS'
    report['releaseBlocked'] = report['status'] == 'FAIL'
    report['summary'] = {'failCount': sum(i['level'] == 'FAIL' for i in report['issues']),
                         'warnCount': sum(i['level'] == 'WARN' for i in report['issues'])}
    if raise_on_fail and report['status'] == 'FAIL':
        raise MonthlyBookkeepingGateError(report)
    return report


def _text(value):
    return '' if value is None else str(value).strip()


def _decimal(value):
    if isinstance(value, bool):
        return None
    try:
        number = Decimal(str(value).strip())
        import math
        return number if number.is_finite() and math.isfinite(float(number)) else None
    except (InvalidOperation, ValueError, TypeError, OverflowError):
        return None


def _date(value, epoch):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        value = from_excel(value, epoch)
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    return datetime.fromisoformat(_text(value).replace('/', '-')).strftime('%Y-%m-%d')


def read_source(ledger):
    """独立展开明确合并区域；只读总表，绝不跨空行猜测填充 U/V/W/X/Y。"""
    book = load_workbook(ledger, data_only=True)
    formula_book = None
    try:
        formula_book = load_workbook(ledger, data_only=False)
        if '总表' not in book.sheetnames:
            raise ValueError('台账缺少总表')
        sheet, formulas = book['总表'], formula_book['总表']
        headers = [_text(cell.value) for cell in sheet[1]]
        if len(headers) != len(set(headers)) or set(FIELDS) - set(headers):
            raise ValueError('总表字段缺失或重复')
        positions = {name: index for index, name in enumerate(headers, 1)}
        merged = {}
        for area in sheet.merged_cells.ranges:
            if any(positions[name] <= area.max_col and positions[name] >= area.min_col
                   for name in ('认购机构', '认购份额', '申购利率', '穿透机构', '申购规模')):
                raise ValueError(f'原始机构/金额/报价列有合并区域: {area}')
            for row in range(area.min_row, area.max_row + 1):
                for col in range(area.min_col, area.max_col + 1):
                    merged[row, col] = sheet.cell(area.min_row, area.min_col).value
        rows = []
        for r in range(2, sheet.max_row + 1):
            values = {}
            for name, col in positions.items():
                cached, expression = sheet.cell(r, col), formulas.cell(r, col)
                # 辅助层金额公式沿用最近三期 QC 的 WARN，不作为月度源字段。
                if name in FIELDS and expression.data_type == 'f' and (
                        cached.value is None or cached.data_type == 'e'):
                    raise ValueError(f'{expression.coordinate}关键公式没有有效计算缓存')
                if (name in FIELDS or name == '对应金额（亿）') and cached.data_type == 'e':
                    raise ValueError(f'{cached.coordinate}关键字段含Excel错误')
                values[name] = merged.get((r, col), cached.value)
            if not any(v is not None for v in values.values()):
                continue
            if _text(values.get('项目名称')) == '项目名称' or all(
                    _text(values.get(name)) == name for name in DETAIL_HEADERS):
                continue
            values['_row'] = r
            values['_uncachedLayerSize'] = None
            if '对应金额（亿）' in positions:
                col = positions['对应金额（亿）']
                if formulas.cell(r, col).data_type == 'f' and sheet.cell(r, col).value is None:
                    values['_uncachedLayerSize'] = sheet.cell(r, col).coordinate
            try:
                values['_date'] = _date(values.get('簿记时间'), book.epoch)
            except (ValueError, TypeError, OverflowError):
                values['_date'] = ''
            rows.append(values)
        return rows
    finally:
        book.close()
        if formula_book is not None:
            formula_book.close()


def _aliases():
    result = {}
    for item in json.loads(DEMO_ALIASES.read_text(encoding='utf-8'))['groups']:
        for alias in item['aliases']:
            if not alias or alias in result or alias == item['canonical']:
                raise ValueError('Demo别名必须唯一、非空、非自身映射')
            result[alias] = (item['canonical'], item['reason'])
    if set(result) & {name for name, _ in result.values()}:
        raise ValueError('Demo别名必须单跳')
    return result


def _identity(raw, aliases):
    from abs_common import normalize_investor_name
    formatted = re.sub(r'(股份有限公司|有限责任公司|有限公司)(?=[（(]|$)', '', raw.strip())
    if '信托' not in formatted:
        formatted = formatted.replace('(', '（').replace(')', '）')
    name = normalize_investor_name(formatted)
    return aliases.get(name, (name, '沿用项目机构别名与格式规则'))


def _classification_lookup():
    from abs_common import normalize_investor_name
    from institution_identity import registry
    source = json.loads(CLASSIFICATION.read_text(encoding='utf-8'))
    result = {}
    for category, names in source['category_to_institutions'].items():
        for name in names:
            for member in re.split(r'[/&：:]', name):
                member = normalize_investor_name(member.strip())
                if '理财' in member and '信托' not in member:
                    group = '理财子'
                elif category in {'全国大行', '国股行', '城商行', '外资行'} or category.startswith('ABN主承销商'):
                    group = '银行'
                elif category == '公募基金':
                    group = '公募基金'
                elif category == '券商' or category.startswith(('ABS管理人', 'ABS销售机构')):
                    group = '券商'
                else:
                    group = '其他'
                result.setdefault(member, group)
    classification = _json(FUND_RULES.read_text(encoding='utf-8'))
    funds = classification.get('publicFunds')
    if classification.get('schema') != 'monthly-bid-classification/v1' or not isinstance(funds, list) or len(funds) != 16 or any(
            not isinstance(f, str) or f != f.strip() or not f.endswith('基金') for f in funds) or len(set(funds)) != len(funds):
        raise ValueError('月表公募分类规则必须有正确schema及16个完整唯一publicFunds名称')
    if classification.get('publicFundAliases') != {'前海开源': '前海开源基金'} or '前海开源基金' not in funds:
        raise ValueError('月表公募分类别名只允许前海开源')
    origin = classification.get('source', {})
    if not isinstance(origin, dict) or not _text(origin.get('path')) or origin.get('constant') != 'FUNDS' or not re.fullmatch(r'[0-9a-f]{64}', _text(origin.get('sha256'))):
        raise ValueError('公募规则必须保留原FUNDS溯源，但不依赖原分析脚本')
    for fund in funds:
        result[normalize_investor_name(fund)] = '公募基金'
    result.update({alias: '公募基金' for alias in classification['publicFundAliases']})
    for bare in ('平安', '华夏', '中银', '建信'):
        if result.get(bare) == '公募基金':
            del result[bare]
    for item in registry()['aliases']:
        if '证券' in item['alias'] and not any(word in item['alias'] for word in ('银行', '信托')):
            result.setdefault(item['canonical'], '券商')
    overrides = classification.get('explicitCategoryOverrides', {})
    if not isinstance(overrides, dict) or any(not isinstance(name, str) or group not in GROUPS for name, group in overrides.items()):
        raise ValueError('explicitCategoryOverrides必须是精确名称到板块的映射')
    for name, group in overrides.items():
        result[name] = group
    for name in ('国投资管（投行）', '华泰资管（自营）'):
        if overrides.get(name) != '券商':
            raise ValueError(f'已确认分类覆盖缺失：{name}=券商')
    return result


def _classify(name, lookup):
    if name in lookup:
        return lookup[name], '机构分配字典精确匹配'
    if any(word in name for word in ('信托', '私募', '利曦', '保险', '养老', '财务', '保理', '小贷', '员工理财')):
        return '其他', '源名称行业标识'
    if '理财' in name:
        return '理财子', '源名称理财标识'
    if any(word in name for word in ('银行', '农商', '农金市')):
        return '银行', '源名称银行标识'
    if '基金' in name and lookup.get(name.removesuffix('管理')) == '公募基金':
        return '公募基金', '机构分配字典公募名称匹配'
    if any(word in name for word in ('证券', '中信建投', '建投衍生品', '中金', '中国国际金融',
                                     '国泰海通', '国联民生', '申万宏源', '中信投顾')):
        return '券商', '源名称券商标识'
    return '其他', '分类未确认，保留其他'


def _same(actual, expected):
    if isinstance(expected, Decimal):
        # Excel源单元格可为数值字符串；输出JSON必须与前端number契约一致。
        if type(actual) not in (int, float):
            return False
        value = _decimal(actual)
        return value is not None and abs(value - expected) <= max(Decimal('1e-12'), abs(expected) * Decimal('1e-12'))
    if isinstance(expected, dict):
        return isinstance(actual, dict) and set(actual) == set(expected) and all(_same(actual[k], v) for k, v in expected.items())
    if isinstance(expected, (list, tuple)):
        return isinstance(actual, (list, tuple)) and len(actual) == len(expected) and all(_same(a, b) for a, b in zip(actual, expected))
    return type(actual) is type(expected) and actual == expected


def _expect(report, actual, expected, code, context):
    if not _same(actual, expected):
        _issue(report, code, context)


def _check_source_fields(row, report):
    complete_header = all(_text(row.get(f)) == f for f in ('申购利率', '穿透机构', '申购规模'))
    for field in ('认购份额', '申购规模', '规模', '对应金额（亿）', '申购利率', '成本'):
        value, raw = _decimal(row.get(field)), _text(row.get(field))
        if raw in ('', '-') or (field in ('认购份额', '申购规模', '申购利率', '成本') and raw == field) or (
                field in ('申购利率', '成本') and raw in ('自持', '不出表')):
            continue
        if value is None or value < 0 or (field in ('申购利率', '成本') and value > 1):
            _issue(report, 'MONTHLY_SOURCE_VALUE', f"总表{row['_row']} {field}非法，不能静默丢弃")
    for field in ('认购机构', '穿透机构'):
        if _text(row.get(field)) not in ('', '-', field) and not isinstance(row.get(field), str):
            _issue(report, 'MONTHLY_SOURCE_INSTITUTION', f"总表{row['_row']} {field}不是名称字符串")
    if not complete_header and any(_text(row.get(f)) == f for f in ('申购利率', '穿透机构', '申购规模')):
        _issue(report, 'MONTHLY_LOCAL_HEADER', f"总表{row['_row']} WXY仅部分字段为表头占位，不是完整局部表头")
    bidder = _text(row.get('穿透机构'))
    if bidder in ('', '-', '穿透机构') and any(_text(row.get(f)) not in ('', '-', f) for f in ('申购利率', '申购规模')):
        _issue(report, 'MONTHLY_ORPHAN_BID', f"总表{row['_row']} W/Y有值但X机构缺失")
    if _text(row.get('认购机构')) in ('', '-', '认购机构') and _text(row.get('认购份额')) not in ('', '-', '认购份额'):
        _issue(report, 'MONTHLY_ORPHAN_ALLOCATION', f"总表{row['_row']} V有金额但U机构缺失")
    if bidder not in ('', '-', '穿透机构') and not any(mark in bidder + _text(row.get('申购利率')) for mark in ('自持', '不出表')):
        rate, size = _decimal(row.get('申购利率')), _decimal(row.get('申购规模'))
        if rate is None or size is None or size == 0:
            _issue(report, 'MONTHLY_BID_PARTIAL_OR_ZERO', f"总表{row['_row']} 非自持WXY部分披露/明确零，禁止UV回退且阻断正式放行",
                   level='WARN', blocks_release=True)


def _bid_evidence(row):
    fields = ('申购利率', '穿透机构', '申购规模')
    return not all(_text(row.get(f)) == f for f in fields) and any(_text(row.get(f)) not in ('', '-') for f in fields)


def _source_scope(rows, ledger, report):
    scoped = []
    stamp = re.fullmatch(r'(\d{4})年ABS发行台账-(\d{4})-定稿.xlsx', Path(ledger).name)
    cutoff = '-'.join((stamp[1], stamp[2][:2], stamp[2][2:])) if stamp else None
    for row in rows:
        if _text(row.get('发行场所')) not in VENUES:
            continue
        date = row['_date']
        if not date or not _text(row.get('项目名称')) or not _text(row.get('资产类型')):
            _issue(report, 'MONTHLY_SOURCE_SCOPE', f"总表{row['_row']} 场内项目缺资产/项目/有效日期")
            continue
        if cutoff and (date > cutoff or date[:4] != stamp[1]):
            _issue(report, 'MONTHLY_SOURCE_DATE_RANGE', f"总表{row['_row']} 日期{date}超出台账业务年份/截止日{cutoff}")
        if not date.startswith(f'{YEAR}-'):
            continue
        asset = ASSET_ALIASES.get(_text(row['资产类型']), _text(row['资产类型']))
        row['_key'] = asset, date, _text(row['项目名称']), _text(row['发行场所'])
        row['_month'] = int(date[5:7]) - 1
        scoped.append(row)
        _check_source_fields(row, report)
    if not scoped:
        _issue(report, 'MONTHLY_EMPTY_SCOPE', '全年范围内没有有效项目')
    return scoped


def _source_conservation(scoped, report):
    by_project = defaultdict(list)
    for row in scoped:
        by_project[row['_key']].append(row)
    epsilon = Decimal('0.0001')
    checked_layers = 0
    details = []
    for key, rows in by_project.items():
        amounts = {_decimal(r.get('规模')) for r in rows} - {None}
        total = sum((_decimal(r.get('认购份额')) or Decimal(0) for r in rows), Decimal(0))
        summary = {'project': list(key), 'rows': [r['_row'] for r in rows],
                   'uvIncludingSelfHeld': str(total), 'projectSizes': sorted(map(str, amounts))}
        details.append(summary)
        context = f"{key[0]}/{key[1]}/{key[2]}/{key[3]} 总表{summary['rows']}"
        if len(amounts) > 1:
            _issue(report, 'MONTHLY_PROJECT_SIZE_CONFLICT', context + f' J发行规模冲突{sorted(map(str, amounts))}')
        elif not amounts:
            _issue(report, 'MONTHLY_PROJECT_SIZE_MISSING', context + ' J缺失，无法核验项目上限', level='WARN')
        else:
            amount = next(iter(amounts))
            if total - amount > epsilon:
                _issue(report, 'MONTHLY_ALLOCATION_PROJECT_OVER', context + f' 全部U/V（含自持）{total} > J{amount}')
            elif abs(total - amount) > epsilon:
                _issue(report, 'MONTHLY_ALLOCATION_SIZE_DIFF', context + f' 已披露U/V{total}与J{amount}不同', level='WARN')
        for field in ('期限', '计划管理人'):
            if len({_text(r.get(field)) for r in rows}) > 1:
                _issue(report, 'MONTHLY_PROJECT_FIELD_CONFLICT', context + f' {field}冲突')
        layers = defaultdict(list)
        for row in rows:
            layer = _text(row.get('分层情况'))
            layers[layer if layer not in ('', '-') else '未分层'].append(row)
        for layer, layer_rows in layers.items():
            checked_layers += 1
            layer_amounts = {_decimal(r.get('对应金额（亿）')) for r in layer_rows} - {None}
            allocated = sum((_decimal(r.get('认购份额')) or Decimal(0) for r in layer_rows), Decimal(0))
            coordinates = [r['_row'] for r in layer_rows]
            uncalculated = [r['_uncachedLayerSize'] for r in layer_rows if r.get('_uncachedLayerSize')]
            label = f'{key[0]}/{key[1]}/{key[2]}/{key[3]}/{layer} 总表{coordinates}'
            if uncalculated:
                _issue(report, 'MONTHLY_LAYER_SIZE_UNCACHED', label + f' 辅助R公式无缓存{uncalculated}，仍校验项目上限', level='WARN')
            if not layer_amounts and not uncalculated:
                _issue(report, 'MONTHLY_LAYER_SIZE_MISSING', label + ' R未披露，无法核验档位上限', level='WARN')
            if len(layer_amounts) > 1:
                _issue(report, 'MONTHLY_LAYER_SIZE_CONFLICT', label + f' R档位金额冲突{sorted(map(str, layer_amounts))}')
            elif layer_amounts and allocated - next(iter(layer_amounts)) > epsilon:
                _issue(report, 'MONTHLY_ALLOCATION_LAYER_OVER', label + f' 全部U/V（含自持）{allocated} > R{next(iter(layer_amounts))}')
            duplicate_rows = defaultdict(list)
            for r in layer_rows:
                if _text(r.get('认购机构')) not in ('', '-', '认购机构'):
                    duplicate_rows[_text(r['认购机构']), _text(r['认购份额']), _text(r['成本'])].append(r['_row'])
            for duplicates in duplicate_rows.values():
                if len(duplicates) > 1:
                    _issue(report, 'MONTHLY_DUPLICATE_ALLOCATION', label + f' 相同U/V/成本行{duplicates}，保留全部待业务确认', level='WARN')
    report['checks']['sourceConservation'] = {'projects': len(by_project), 'layers': checked_layers,
        'includesSelfHeld': True, 'epsilonBillions': str(epsilon), 'projectTotals': details}


def verify_source(ledger, *, raise_on_fail=True):
    """全年源字段与项目/层守恒；FAIL仅阻断release，可供preview展示报告。"""
    report = _report(ledger)
    try:
        before = file_sha256(ledger)
        scoped = _source_scope(read_source(ledger), ledger, report)
        _source_conservation(scoped, report)
        _expect(report, file_sha256(ledger), before, 'MONTHLY_SOURCE_CHANGED', '源校验期间台账改变')
        report['checks']['sourceSha256'] = before
    except (OSError, ValueError, TypeError, KeyError) as error:
        _issue(report, 'MONTHLY_SCHEMA_OR_SOURCE', str(error))
    return _finish(report, raise_on_fail)


def verify_data(monthly, ledger, *, raise_on_fail=True):
    report = _report(ledger)
    if not isinstance(monthly, dict):
        _issue(report, 'MONTHLY_SCHEMA', '月度数据必须为对象')
        return _finish(report, raise_on_fail)
    try:
        before = file_sha256(ledger)
        provenance = build_provenance(ledger)
        _expect(report, monthly.get('provenance'), provenance, 'MONTHLY_PROVENANCE',
                '月度来源/规则/别名/资源依赖指纹或scope缺失、陈旧')
        for field, path in (('sourceSha256', ledger), ('classificationSha256', CLASSIFICATION),
                            ('fundRulesSha256', FUND_RULES), ('demoAliasesSha256', DEMO_ALIASES)):
            _expect(report, monthly.get(field), file_sha256(path), 'MONTHLY_FINGERPRINT', field + '与实际文件不一致')
        if 'source' in monthly:
            _expect(report, monthly['source'], Path(ledger).name, 'MONTHLY_SOURCE', '月度来源文件名错误')
        if 'scope' in monthly:
            _expect(report, monthly['scope'], SCOPE, 'MONTHLY_SCOPE', '月度scope不一致')
        _expect(report, monthly.get('year'), YEAR, 'MONTHLY_YEAR', '业务年份不一致')
        _expect(report, monthly.get('groups'), GROUPS, 'MONTHLY_GROUPS', '全局板块集合/顺序不一致')
        rows = read_source(ledger)
        aliases, lookup = _aliases(), _classification_lookup()
        from abs_common import normalize_investor_name
        from institution_identity import registry
        for pair in registry()['distinct']:
            if _identity(pair['left'], aliases)[0] == _identity(pair['right'], aliases)[0]:
                _issue(report, 'MONTHLY_DISTINCT_IDENTITY', f"独立机构被合并: {pair['left']} / {pair['right']}")
        scoped = _source_scope(rows, ledger, report)
        _source_conservation(scoped, report)
        report['sourceQC'] = {'status': 'FAIL' if any(i['level'] == 'FAIL' and (
            i['code'].startswith('MONTHLY_SOURCE') or i['code'].startswith('MONTHLY_ALLOCATION') or
            i['code'].startswith('MONTHLY_LAYER') or i['code'].startswith('MONTHLY_PROJECT') or
            i['code'].startswith('MONTHLY_ORPHAN')) for i in report['issues']) else 'PASS'}
        projects, evidence = defaultdict(set), set()
        for row in scoped:
            key = row['_key']
            projects[key[0]].add(key)
            # 原表已有申购证据，即使字段不完整也不能伪称“无WXY”而补U/V。
            if _bid_evidence(row):
                evidence.add(key)
        records, allocated, allocation_records = defaultdict(list), defaultdict(lambda: defaultdict(Decimal)), []
        allocation_missing, unknown_records = defaultdict(list), []
        excluded, changes, fallback_rows = [], {}, []
        for row in scoped:
            asset, date, project, venue = row['_key']
            raw = _text(row['穿透机构'])
            size, rate = _decimal(row['申购规模']), _decimal(row['申购利率'])
            if raw and raw not in ('-', '穿透机构'):
                if size is None or size <= 0 or rate is None:
                    excluded.append({'row': row['_row'], 'institution': raw,
                                     'reason': '无有效申购利率或正申购规模（含自持）'})
                    _issue(report, 'MONTHLY_BID_MISSING', f"总表{row['_row']} 申购字段不完整/非正规模，保留排除记录", level='WARN')
                else:
                    name, reason = _identity(raw, aliases)
                    records[asset, name, row['_month']].append({'row': row['_row'], 'rawInstitution': raw,
                        'project': project, 'date': date, 'rate': rate, 'size': size, 'venue': venue, 'track': 'bids'})
                    previous = normalize_investor_name(raw)
                    if previous != name:
                        change = changes.setdefault(raw, {'raw': raw, 'previousName': previous, 'name': name,
                                                          'reason': reason, 'rows': []})
                        change['rows'].append(row['_row'])
            allocation_raw, amount = _text(row['认购机构']), _decimal(row['认购份额'])
            if (allocation_raw in ('', '-', '认购机构') or any(
                    mark in allocation_raw + _text(row['成本']) for mark in ('自持', '不出表'))):
                continue
            name, _ = _identity(allocation_raw, aliases)
            if amount is None:
                allocation_missing[asset, name].append(row['_row'])
                unknown_records.append({'row': row['_row'], 'asset': asset, 'name': name,
                    'rawInstitution': allocation_raw, 'date': date, 'project': project, 'venue': venue,
                    'size': None, 'reason': 'U已披露、V未披露，配售金额未知'})
                _issue(report, 'MONTHLY_ALLOCATION_UNKNOWN', f'总表{row["_row"]} {asset}/{name} U已披露但V缺失，分母不完整/份额未知', level='WARN')
                continue
            if amount <= 0:
                continue
            allocated[asset][name] += amount
            allocation_records.append({'row': row['_row'], 'asset': asset, 'name': name,
                'rawInstitution': allocation_raw, 'date': date, 'size': amount})
            if row['_month'] < 3 and row['_key'] not in evidence:
                records[asset, name, row['_month']].append({'row': row['_row'], 'rawInstitution': allocation_raw,
                    'project': project, 'date': date, 'rate': _decimal(row['成本']), 'size': amount,
                    'venue': venue, 'track': 'allocations'})
                fallback_rows.append(row['_row'])
        assets = monthly.get('assets')
        if not isinstance(assets, dict):
            raise ValueError('assets必须为对象')
        _expect(report, set(assets), set(projects), 'MONTHLY_ASSET_SCOPE', '资产集合/别名合并范围不一致')
        all_bids = all_fallback = all_institutions = 0
        for asset, project_keys in projects.items():
            actual = assets.get(asset, {})
            if not isinstance(actual, dict):
                raise ValueError('资产条目必须为对象')
            names = {name for a, name, _ in records if a == asset}
            visible_groups = GROUPS[:3] if asset in THREE_GROUP_ASSETS else GROUPS
            expected_insts = []
            denominator = sum(allocated[asset].values(), Decimal(0))
            missing_rows = sorted(r for (a, _), rows in allocation_missing.items() if a == asset for r in rows)
            allocation_complete = not missing_rows
            for name in names:
                months = [sorted(records[asset, name, m], key=lambda r: (r['track'] != 'bids', r['row'])) for m in range(12)]
                participation = {(r['date'], r['project'], r['venue']) for month in months for r in month}
                group, classification_reason = _classify(name, lookup)
                amount = allocated[asset][name]
                bid_sizes = [sum((r['size'] for r in month if r['track'] == 'bids'), Decimal(0)) for month in months]
                allocation_sizes = [sum((r['size'] for r in month if r['track'] == 'allocations'), Decimal(0)) for month in months]
                expected_insts.append({'name': name, 'group': group, 'classificationReason': classification_reason,
                    'months': months, 'participationCount': len(participation),
                    'monthCount': sum(bool(month) for month in months),
                    'participationRate': Decimal(len(participation)) / len(project_keys),
                    'knownAllocatedSize': amount,
                    'allocatedSize': None if allocation_missing[asset, name] else amount,
                    'allocationShare': amount / denominator if denominator and allocation_complete else None,
                    'allocationComplete': not allocation_missing[asset, name],
                    'unknownAllocationRows': allocation_missing[asset, name],
                    'monthlyBidSizes': bid_sizes, 'monthlyFallbackSizes': allocation_sizes,
                    'monthlySizes': [sum((r['size'] for r in month), Decimal(0)) for month in months],
                    'monthlyTracks': ['mixed' if any(r['track'] == 'bids' for r in month) and any(r['track'] == 'allocations' for r in month) else
                                      'bids' if any(r['track'] == 'bids' for r in month) else 'allocations' if month else '' for month in months]})
            expected_insts.sort(key=lambda i: (GROUPS.index(i['group']), -i['participationCount'], -i['monthCount'], i['name']))
            insts = actual.get('institutions', [])
            _expect(report, [i.get('name') for i in insts], [i['name'] for i in expected_insts],
                    'MONTHLY_INSTITUTION_SCOPE_SORT', asset + ' 机构集合/板块参与率排序不一致')
            inst_map = {i.get('name'): i for i in insts}
            for expected in expected_insts:
                inst = inst_map.get(expected['name'], {})
                for field, value in expected.items():
                    code = 'MONTHLY_RAW_RECONCILIATION' if field == 'months' else 'MONTHLY_INSTITUTION_METRIC'
                    _expect(report, inst.get(field), value, code, f"{asset}/{expected['name']}/{field}与独立源表核对不一致")
            visible = [i for i in expected_insts if i['group'] in visible_groups]
            bid_count = sum(r['track'] == 'bids' for i in expected_insts for m in i['months'] for r in m)
            fallback_count = sum(r['track'] == 'allocations' for i in expected_insts for m in i['months'] for r in m)
            metrics = {'counts': [sum(bool(i['months'][m]) for i in expected_insts) for m in range(12)],
                'visibleCounts': [sum(bool(i['months'][m]) for i in visible) for m in range(12)],
                'visibleGroups': visible_groups, 'projectCount': len(project_keys),
                'projectCounts': [sum(int(key[1][5:7]) == m + 1 for key in project_keys) for m in range(12)],
                'totalKnownAllocatedSize': denominator,
                'totalAllocatedSize': denominator if allocation_complete else None, 'allocationComplete': allocation_complete,
                'unknownAllocationRows': missing_rows, 'bidCount': bid_count, 'fallbackCount': fallback_count,
                'bubbleMaxSize': max((s for i in visible for s in i['monthlySizes']), default=Decimal(1))}
            for field, value in metrics.items():
                _expect(report, actual.get(field), value, 'MONTHLY_ASSET_METRIC', f'{asset}/{field}不一致')
            all_bids += bid_count
            all_fallback += fallback_count
            all_institutions += len(expected_insts)
        _expect(report, monthly.get('fallbackRows'), sorted(fallback_rows), 'MONTHLY_FALLBACK', '1—3月无WXY项目U/V补充行集合错误')
        _expect(report, monthly.get('allocationRecords'), allocation_records, 'MONTHLY_ALLOCATION_RECORDS', '全年U/V原始行、机构或金额不一致（含隐藏板块）')
        _expect(report, monthly.get('unknownAllocationRecords'), unknown_records, 'MONTHLY_ALLOCATION_UNKNOWN_RECORDS', '非自持未知配售U/V证据记录不一致')
        _expect(report, monthly.get('excluded'), excluded, 'MONTHLY_EXCLUDED', '排除申购行集合不一致')
        _expect(report, monthly.get('identityChanges'), list(changes.values()), 'MONTHLY_IDENTITY_CHANGES', '别名变更审计记录不一致')
        _expect(report, file_sha256(ledger), before, 'MONTHLY_SOURCE_CHANGED', '校验期间台账改变')
        _expect(report, build_provenance(ledger), provenance, 'MONTHLY_RULES_CHANGED', '校验期间规则依赖改变')
        report['checks'].update({'assets': len(projects), 'projects': sum(map(len, projects.values())),
            'institutions': all_institutions, 'rawBids': all_bids, 'fallbackAllocations': all_fallback,
            'allocationRecords': len(allocation_records), 'independentSourceReader': 'openpyxl',
            'aggregation': 'Decimal and project sets', 'scope': SCOPE})
    except (OSError, ValueError, TypeError, KeyError, AttributeError, InvalidOperation) as error:
        _issue(report, 'MONTHLY_SCHEMA_OR_SOURCE', str(error))
    return _finish(report, raise_on_fail)


class _HTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=False)
        self.ids = Counter()
        self.blocks = []
        self.block_elements = []
        self.current = None
        self.current_element = None
        self.elements = []
        self.stack = []

    def handle_starttag(self, tag, attrs):
        names = [key for key, _ in attrs]
        if len(set(names)) != len(names):
            raise ValueError('HTML属性重复')
        attrs = dict(attrs)
        element = {'tag': tag, 'attrs': attrs, 'ancestors': tuple(self.stack), 'text': ''}
        self.elements.append(element)
        if attrs.get('id'):
            self.ids[attrs['id']] += 1
        if tag in ('script', 'style'):
            self.current = [tag, attrs, '']
            self.current_element = element
        if tag not in ('area', 'base', 'br', 'col', 'embed', 'hr', 'img', 'input', 'link', 'meta', 'param', 'source', 'track', 'wbr'):
            self.stack.append((tag, attrs, element))

    def handle_data(self, data):
        if self.current is not None:
            self.current[2] += data
        if self.stack:
            self.stack[-1][2]['text'] += data

    def handle_endtag(self, tag):
        if self.current is not None and self.current[0] == tag:
            self.blocks.append(self.current)
            self.block_elements.append(self.current_element)
            self.current = None
            self.current_element = None
        for index in range(len(self.stack) - 1, -1, -1):
            if self.stack[index][0] == tag:
                del self.stack[index:]
                break

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)


def _json(payload):
    def pairs(items):
        result = {}
        for key, value in items:
            if key in result:
                raise ValueError(f'JSON字段重复: {key}')
            result[key] = value
        return result
    def invalid(value):
        raise ValueError(f'JSON非有限数值: {value}')
    return json.loads(payload, object_pairs_hook=pairs, parse_constant=invalid)


def _inactive_context(node, *, allow_own_hidden=False):
    """Tab的class隐藏允许；模板、inert、hidden属性及内联隐藏包装不允许。"""
    contexts = [(tag, attrs) for tag, attrs, _ in node['ancestors']]
    contexts.append((node['tag'], node['attrs']))
    for index, (tag, attrs) in enumerate(contexts):
        if tag == 'template' or 'inert' in attrs:
            return tag + ' template/inert非活动上下文'
        if 'hidden' in attrs and not (allow_own_hidden and index == len(contexts) - 1):
            return tag + ' hidden属性'
        style = re.sub(r'/\*.*?\*/', '', attrs.get('style', ''), flags=re.S).lower()
        if re.search(r'(?:^|;)\s*(?:display\s*:\s*none|visibility\s*:\s*(?:hidden|collapse)|content-visibility\s*:\s*hidden)\s*(?:!\s*important\s*)?(?:;|$)', style):
            return tag + ' 内联隐藏样式'
    return None


def extract_data(content):
    parser = _HTML()
    parser.feed(content)
    parser.close()
    values = []
    for identifier in ('lb-data', 'lb-monthly-data'):
        blocks = [b for b in parser.blocks if b[0] == 'script' and b[1].get('id') == identifier]
        if parser.ids[identifier] != 1 or len(blocks) != 1 or blocks[0][1].get('type') != 'application/json' or 'src' in blocks[0][1]:
            raise ValueError(f'{identifier}必须为唯一内联JSON脚本块')
        node = next(e for e in parser.block_elements if e['attrs'].get('id') == identifier)
        inactive = _inactive_context(node)
        if inactive:
            raise ValueError(f'{identifier}位于非活动/隐藏上下文: {inactive}')
        values.append(_json(blocks[0][2]))
    return values[0], values[1]


def _verify_resources(parser, report):
    expected_js = [(SCRIPTS / 'latest_bookkeeping_panel.js').read_text(encoding='utf-8'),
                   (SCRIPTS / 'lab/monthly_bid_institutions_demo.js').read_text(encoding='utf-8')]
    scripts = [b for b in parser.blocks if b[0] == 'script' and b[1].get('type', '') not in ('application/json', 'application/ld+json')]
    actual_js = [b[2] for b in scripts]
    for block, node in zip(parser.blocks, parser.block_elements):
        if (block[0] == 'script' and block[2] in expected_js) or (
                block[0] == 'style' and 'latest-bookkeeping' in block[2]):
            inactive = _inactive_context(node)
            if inactive:
                _issue(report, 'HTML_INACTIVE_RESOURCE', '实际组件资源位于非活动/隐藏上下文: ' + inactive)
    for script, name in zip(expected_js, ('latest', 'monthly')):
        if actual_js.count(script) != 1:
            _issue(report, 'HTML_INLINE_JS', f'{name}实际内联JS缺失/篡改/重复')
    if all(actual_js.count(script) == 1 for script in expected_js) and actual_js.index(expected_js[0]) > actual_js.index(expected_js[1]):
        _issue(report, 'HTML_INLINE_JS_ORDER', '最近项目与月度JS执行顺序改变')
    for block in scripts:
        if block[2] in expected_js:
            if 'src' in block[1] or block[1].get('type', '') not in ('', 'text/javascript', 'application/javascript'):
                _issue(report, 'HTML_INLINE_JS', '组件JS属性导致脚本不执行或被替换')
        elif any(token in block[2] + block[1].get('src', '') for token in ('lb-', 'latest-bookkeeping', 'monthly_bid')):
            _issue(report, 'HTML_EXTRA_COMPONENT_JS', '存在额外可覆盖簿记组件的JS')
    latest_css = (SCRIPTS / 'latest_bookkeeping_panel.css').read_text(encoding='utf-8')
    monthly_css = (SCRIPTS / 'lab/monthly_bid_institutions_demo.css').read_text(encoding='utf-8')
    styles = [b[2] for b in parser.blocks if b[0] == 'style']
    for css, name in ((latest_css, 'latest'), (monthly_css, 'monthly')):
        if sum(s.count(css) for s in styles) != 1:
            _issue(report, 'HTML_INLINE_CSS', f'{name}实际内联CSS缺失/篡改/重复')
    for style in styles:
        remainder = style.replace(latest_css, '').replace(monthly_css, '')
        if any(token in remainder for token in ('lb-', 'latest-bookkeeping', 'monthly_bid')):
            _issue(report, 'HTML_EXTRA_COMPONENT_CSS', '存在额外可覆盖簿记组件的CSS')
    for tag, attrs, content in parser.blocks:
        if tag == 'style' and (latest_css in content or monthly_css in content):
            if attrs.get('type', 'text/css') != 'text/css' or attrs.get('media', 'all') not in ('all', 'screen') or 'disabled' in attrs:
                _issue(report, 'HTML_INLINE_CSS', '组件CSS属性导致当前样式失效')
    report['checks']['inlineResources'] = {'latestJS': file_sha256(SCRIPTS / 'latest_bookkeeping_panel.js'),
        'monthlyJS': file_sha256(SCRIPTS / 'lab/monthly_bid_institutions_demo.js'),
        'latestCSS': file_sha256(SCRIPTS / 'latest_bookkeeping_panel.css'),
        'monthlyCSS': file_sha256(SCRIPTS / 'lab/monthly_bid_institutions_demo.css')}


def _verify_structure(parser, report):
    """核对真实body的节点/控件契约，允许Demo调整布局与添加包装。"""
    template = _HTML()
    template.feed((SCRIPTS / 'latest_bookkeeping_panel_body.html').read_text(encoding='utf-8'))
    contracts = {e['attrs']['id']: e for e in template.elements if e['attrs'].get('id')}
    monthly_contracts = {
        'latest-bookkeeping': ('div', {}), 'lb-monthly-title': ('h2', {}),
        'lb-monthly-search': ('input', {'type': 'search'}),
        'lb-monthly-expand': ('button', {'type': 'button', 'aria-expanded': 'false'}),
        'lb-monthly-scroll': ('div', {'tabindex': '0'}),
        'lb-monthly-head': ('thead', {}), 'lb-monthly-rows': ('tbody', {})}
    contracts.update({name: {'tag': tag, 'attrs': attrs} for name, (tag, attrs) in monthly_contracts.items()})
    by_id = {e['attrs'].get('id'): e for e in parser.elements if e['attrs'].get('id')}
    for identifier, contract in contracts.items():
        node = by_id.get(identifier)
        if parser.ids[identifier] != 1 or node is None:
            _issue(report, 'HTML_COMPONENT_STRUCTURE', f'{identifier}节点缺失/重复')
            continue
        if node['tag'] != contract['tag']:
            _issue(report, 'HTML_COMPONENT_STRUCTURE', f'{identifier}节点类型错误')
        inactive = _inactive_context(node, allow_own_hidden='hidden' in contract['attrs'])
        if inactive:
            _issue(report, 'HTML_INACTIVE_COMPONENT', f'{identifier}位于非活动/隐藏上下文: {inactive}')
        for attr in ('type', 'role', 'aria-expanded', 'tabindex'):
            if attr in contract['attrs'] and node['attrs'].get(attr) != contract['attrs'][attr]:
                _issue(report, 'HTML_COMPONENT_STRUCTURE', f'{identifier}/{attr}控件契约不一致')
        if identifier != 'latest-bookkeeping' and not any(a[1].get('id') == 'latest-bookkeeping' for a in node['ancestors']):
            _issue(report, 'HTML_COMPONENT_STRUCTURE', f'{identifier}不在簿记容器内')
        if 'hidden' in node['attrs'] and 'hidden' not in contract['attrs']:
            _issue(report, 'HTML_COMPONENT_STRUCTURE', f'{identifier}被隐藏')
        if 'disabled' in node['attrs'] or 'style' in node['attrs']:
            _issue(report, 'HTML_COMPONENT_STRUCTURE', f'{identifier}存在额外禁用/内联样式')
    for identifier, count in parser.ids.items():
        if (identifier.startswith('lb-') or identifier == 'latest-bookkeeping') and count != 1:
            _issue(report, 'HTML_COMPONENT_STRUCTURE', f'{identifier}节点重复')
    views = [e for e in parser.elements if 'data-view' in e['attrs']]
    actual_views = [(e['tag'], e['attrs'].get('type'), e['attrs']['data-view'], e['text'].strip()) for e in views]
    _expect(report, actual_views, [('button', 'button', 'bids', '申购明细'),
                                 ('button', 'button', 'allocations', '最终配售')],
            'HTML_DETAIL_CONTROLS', '实际申购/最终配售控件缺失、顺序或标签改变')
    for identifier in ('lb-monthly-head', 'lb-monthly-rows'):
        node = by_id.get(identifier)
        if node and not any(a[0] == 'table' and 'lb-monthly-table' in a[1].get('class', '').split() for a in node['ancestors']):
            _issue(report, 'HTML_COMPONENT_STRUCTURE', f'{identifier}缺少实际月度table容器')


def verify_html(content, ledger, *, raise_on_fail=True):
    report = _report(ledger)
    try:
        parser = _HTML()
        parser.feed(content)
        parser.close()
        latest, monthly = extract_data(content)
        report = verify_data(monthly, ledger, raise_on_fail=False)
        _verify_structure(parser, report)
        if not isinstance(latest, dict):
            raise ValueError('lb-data必须为对象')
        meta = latest.get('provenance', {})
        if (meta.get('sourceSha256') != monthly.get('sourceSha256') or
                latest.get('source') != Path(ledger).name or
                monthly.get('provenance', {}).get('source') != latest.get('source')):
            _issue(report, 'HTML_SOURCE_MISMATCH', '最近项目lb-data与monthly来源不一致')
        latest_names = [a.get('name') for a in latest.get('assets', [])]
        source_years = defaultdict(set)
        for row in read_source(ledger):
            if _text(row.get('发行场所')) in VENUES and row['_date'] and _text(row.get('项目名称')) and _text(row.get('资产类型')):
                name = ASSET_ALIASES.get(_text(row['资产类型']), _text(row['资产类型']))
                source_years[name].add(row['_date'][:4])
        monthly_names = set(monthly.get('assets', {}))
        history_only = {name for name, years in source_years.items() if str(YEAR) not in years}
        if (not monthly_names <= set(latest_names) or len(latest_names) != len(set(latest_names))
                or not (set(latest_names) - monthly_names) <= history_only):
            _issue(report, 'HTML_ASSET_SELECTOR_SCOPE', '选择器须含全部当年月表资产；额外资产必须由源表证明只存在于其他年份')
        _verify_resources(parser, report)
        # 延迟导入只使用renderer，绝不调用其聚合函数。支持组件被综合框架包装。
        from lab.monthly_bid_institutions_demo import render_component, component_css
        expected_components = [render_component(latest, monthly, preview=True)]
        if monthly.get('auditQc', {}).get('status') == 'PASS':
            expected_components.append(render_component(latest, monthly, preview=False))
        if not any(component in content for component in expected_components):
            _issue(report, 'HTML_CURRENT_RENDERER', '实际月度body/data/JS与当前render_component不一致')
        combined_css = component_css()
        if sum(b[2].count(combined_css) for b in parser.blocks if b[0] == 'style') != 1:
            _issue(report, 'HTML_CURRENT_COMPONENT_CSS', '实际组合样式与当前component_css不一致')
        # render_body被Demo重排包装，保留原QC子数据检查与真实运行时契约。
        import latest_bookkeeping_qc as latest_qc
        from latest_bookkeeping_runtime import verify_runtime
        recent = latest_qc.check_data(latest, ledger, raise_on_fail=False)
        report['latestQC'] = recent
        for item in recent['issues']:
            _issue(report, item['code'], item['message'], level=item['level'], origin='latest')
        if recent['status'] == 'PASS' and not any(i['level'] == 'FAIL' and i['code'].startswith('HTML_') for i in report['issues']):
            report['checks']['latestRuntime'] = verify_runtime(latest)
        _verify_embedded_qc(monthly, parser, report)
        if any(i['level'] == 'FAIL' or i.get('blocksRelease') for i in report['issues']):
            banners = [e for e in parser.elements if 'lb-monthly-qc-banner' in e['attrs'].get('class', '').split()]
            if len(banners) != 1 or banners[0]['attrs'].get('data-qc-status') != 'FAIL' or banners[0]['attrs'].get('role') != 'alert':
                _issue(report, 'HTML_RELEASE_BLOCK_BANNER', '阻断预览必须有唯一FAIL警示横幅及role=alert')
    except (OSError, ValueError, TypeError, KeyError, AttributeError) as error:
        _issue(report, 'HTML_GATE', str(error))
    return _finish(report, raise_on_fail)


def _verify_embedded_qc(monthly, parser, report):
    """嵌入阻断不能被独立健康重算洗成PASS；正常源BLOCK预览不重复报错。"""
    recalculated_block = any(i['level'] == 'FAIL' or i.get('blocksRelease') for i in report['issues'])
    embedded = monthly.get('auditQc')
    banners = [e for e in parser.elements if 'lb-monthly-qc-banner' in e['attrs'].get('class', '').split()]
    if embedded is not None:
        if not isinstance(embedded, dict):
            _issue(report, 'HTML_EMBEDDED_QC_STATE', 'auditQc必须是报告对象')
            return
        status = embedded.get('status')
        issues = embedded.get('issues', [])
        if not isinstance(issues, list) or any(not isinstance(i, dict) for i in issues):
            _issue(report, 'HTML_EMBEDDED_QC_STATE', 'auditQc.issues结构错误')
            return
        embedded_block = (status == 'FAIL' or embedded.get('releaseStatus') == 'BLOCK'
                          or embedded.get('releaseBlocked') is True
                          or any(i.get('level') == 'FAIL' or i.get('blocksRelease') for i in issues))
        contradictions = (status not in ('PASS', 'FAIL')
            or (status == 'PASS' and embedded_block)
            or (status == 'FAIL' and (embedded.get('releaseStatus') == 'PASS' or embedded.get('releaseBlocked') is False))
            or ('releaseBlocked' in embedded and type(embedded['releaseBlocked']) is not bool)
            or ('releaseStatus' in embedded and embedded['releaseStatus'] not in ('PASS', 'BLOCK')))
        if contradictions or embedded_block != recalculated_block:
            _issue(report, 'HTML_EMBEDDED_QC_STATE', '嵌入auditQc状态/阻断标记与独立重算不一致；保守阻断')
        if status == 'PASS' and any(e['attrs'].get('data-qc-status') == 'FAIL' for e in banners):
            _issue(report, 'HTML_EMBEDDED_QC_STATE', '嵌入PASS与实际FAIL横幅不一致')
    elif not recalculated_block and any(e['attrs'].get('data-qc-status') == 'FAIL' for e in banners):
        _issue(report, 'HTML_EMBEDDED_QC_STATE', '实际FAIL横幅与独立健康结果不一致；保守阻断')


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--dashboard', type=Path, required=True)
    parser.add_argument('--ledger', type=Path, required=True)
    parser.add_argument('--report', type=Path)
    args = parser.parse_args(argv)
    try:
        content = args.dashboard.read_text(encoding='utf-8')
        report = verify_html(content, args.ledger, raise_on_fail=False)
    except (OSError, UnicodeError) as error:
        report = _report(args.ledger)
        _issue(report, 'CLI_INPUT', str(error))
        _finish(report, False)
    report['dashboard'] = str(args.dashboard.resolve())
    encoded = json.dumps(report, ensure_ascii=False, allow_nan=False, indent=2) + '\n'
    if args.report:
        if args.report.resolve() in (args.dashboard.resolve(), args.ledger.resolve()):
            parser.error('--report不能覆盖dashboard或ledger')
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(encoded, encoding='utf-8')
    print(encoded, end='')
    return 0 if report['status'] == 'PASS' else 1


if __name__ == '__main__':
    raise SystemExit(main())
