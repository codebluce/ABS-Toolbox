"""独立人工期望的月度数据回归；输入为临时原始总表Excel，不从生产函数推期望。"""
from __future__ import annotations

import copy
from contextlib import redirect_stdout
from decimal import Decimal
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook

LAB = Path(__file__).resolve().parent
sys.path[:0] = [str(LAB), str(LAB.parent)]
import monthly_bid_institutions_demo as demo
import monthly_bookkeeping_qc as gate

# J项目规模、R档位金额、T/U/V及W/X/Y物理位置固定，不消费生产header逻辑。
HEADERS = ['序号', '月份', '类别', '资产类型', '项目名称', '发行场所', '计划管理人',
           '联席承销商', '托管行', '规模', '期限', '簿记时间', '国股CD', 'ALL-IN成本',
           '认购倍数', '分层情况', '分层占比', '对应金额（亿）', '评级', '成本',
           '认购机构', '认购份额', '申购利率', '穿透机构', '申购规模']


def source_row(project='一月甲', date='2026-01-09', venue='上交所', asset='赊销白条',
               bidder='邮储银行', rate=0.02, bid=3, institution='邮储银行', allocated=10,
               cost=0.02):
    return [1, 12, 'ABS', asset, project, venue, '管理人', None, None, 100, '21+3', date,
            None, None, None, '优先A', None, None, 'AAA', cost, institution, allocated,
            rate, bidder, bid]


class MonthlyDataTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='monthly-data-')
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.ledger = self.directory / '2026年ABS发行台账-1016-定稿.xlsx'

    def write_source(self, rows):
        book = Workbook()
        sheet = book.active
        sheet.title = '总表'
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(row)
        book.save(self.ledger)
        book.close()

    def compute(self, rows):
        self.write_source(rows)
        return demo.compute_monthly(self.ledger)

    def institution(self, monthly, name='邮储银行', asset='赊销白条'):
        return next(i for i in monthly['assets'][asset]['institutions'] if i['name'] == name)

    def assert_gate_pass(self, monthly):
        report = gate.verify_data(monthly, self.ledger, raise_on_fail=False)
        self.assertEqual(report['status'], 'PASS', report['issues'])

    def test_partial_wxy_and_explicit_zero_never_unlock_uv(self):
        for row in (source_row(rate=None), source_row(bid=None),
                    source_row(rate=0, bid=0), source_row(rate=0.02, bid=0)):
            with self.subTest(row=row[-3:]):
                monthly = self.compute([row, source_row(project='四月后', date='2026-04-09')])
                asset = monthly['assets']['赊销白条']
                bank = self.institution(monthly)
                self.assertEqual(monthly['fallbackRows'], [])
                self.assertEqual(asset['bidCount'], 1)
                self.assertEqual(asset['fallbackCount'], 0)
                self.assertEqual(bank['monthlySizes'][0], 0)
                self.assertEqual(bank['allocatedSize'], 20)
                self.assertEqual(asset['projectCount'], 2)
                self.assertEqual(bank['participationCount'], 1)
                self.assertTrue(monthly['dataWarnings'])
                report = gate.verify_data(monthly, self.ledger, raise_on_fail=False)
                self.assertEqual(report['status'], 'FAIL')
                self.assertTrue(any(i.get('blocksRelease') for i in report['issues']))
                self.assertFalse(any(i['code'] == 'MONTHLY_FALLBACK' for i in report['issues']))

    def test_entire_year_rejects_boolean_nonfinite_negative_bad_error_and_rate_range(self):
        later = [source_row(project=f'后{m}', date=f'2026-{m:02d}-09') for m in (8, 9, 10)]
        for column in (21, 22, 24):  # V / W / Y
            for value in (True, False, 'NaN', 'nan', 'inf', '-inf', -1, '#VALUE!', '错误'):
                with self.subTest(column=column + 1, value=value):
                    row = source_row()
                    row[column] = value
                    self.write_source([row, *later])
                    with self.assertRaises(demo.MonthlySourceError) as ctx:
                        demo.compute_monthly(self.ledger)
                    self.assertIn('MONTHLY_SOURCE_VALUE', [i['code'] for i in ctx.exception.issues])
                    self.assertEqual(ctx.exception.issues[0]['row'], 2)
        row = source_row(rate=20)
        self.write_source([row, *later])
        with self.assertRaises(demo.MonthlySourceError):
            demo.compute_monthly(self.ledger)

    def test_nan_inf_float_rejected_before_decimal_arithmetic(self):
        # Excel不能保留IEEE NaN为数值；临时表用字符串，另直接核验严格解析器。
        for value in (float('nan'), float('inf'), float('-inf'), Decimal('NaN'), True):
            errors = []
            self.assertIsNone(demo._number(value, '认购份额', 2, errors))
            self.assertEqual(errors[0]['code'], 'MONTHLY_SOURCE_VALUE')

    def test_orphan_zero_is_evidence_and_is_rejected(self):
        for row, code in ((source_row(bidder=None), 'MONTHLY_ORPHAN_BID'),
                          (source_row(bidder=None, rate=0, bid=0), 'MONTHLY_ORPHAN_BID'),
                          (source_row(institution=None, allocated=10), 'MONTHLY_ORPHAN_ALLOCATION'),
                          (source_row(institution=None, allocated=0), 'MONTHLY_ORPHAN_ALLOCATION')):
            self.write_source([row])
            with self.assertRaises(demo.MonthlySourceError) as ctx:
                demo.compute_monthly(self.ledger)
            self.assertIn(code, [i['code'] for i in ctx.exception.issues])

    def test_partial_header_is_not_absence_and_does_not_fallback(self):
        for row in (source_row(rate='申购利率', bidder=None, bid=None),
                    source_row(rate=None, bidder='穿透机构', bid=None),
                    source_row(rate=None, bidder=None, bid='申购规模')):
            self.write_source([row])
            with self.assertRaises(demo.MonthlySourceError) as ctx:
                demo.compute_monthly(self.ledger)
            self.assertIn('MONTHLY_LOCAL_HEADER', [i['code'] for i in ctx.exception.issues])

    def test_completely_absent_wxy_blank_dash_and_complete_headers_only_early_uv(self):
        for rate, bidder, size in ((None, None, None), ('-', '-', '-'),
                                  ('申购利率', '穿透机构', '申购规模')):
            monthly = self.compute([source_row(rate=rate, bidder=bidder, bid=size),
                                    source_row(project='四月无', date='2026-04-09',
                                               rate=None, bidder=None, bid=None)])
            bank = self.institution(monthly)
            self.assertEqual(monthly['fallbackRows'], [2])
            self.assertEqual(bank['monthlySizes'][0], 10)
            self.assertEqual(bank['monthlyTracks'][0], 'allocations')
            self.assertEqual(bank['monthlySizes'][3], 0)
            self.assertEqual(bank['participationRate'], 1 / 2)
            self.assert_gate_pass(monthly)

    def test_project_evidence_on_another_row_blocks_fallback(self):
        monthly = self.compute([source_row(rate=None),
                                source_row(rate=None, bid=None, bidder=None, allocated=5)])
        self.assertEqual(monthly['fallbackRows'], [])
        self.assertEqual(monthly['assets']['赊销白条']['institutions'], [])
        self.assertEqual(monthly['assets']['赊销白条']['totalAllocatedSize'], 15)
        report = gate.verify_data(monthly, self.ledger, raise_on_fail=False)
        self.assertEqual(report['status'], 'FAIL')
        self.assertTrue(any(i.get('blocksRelease') for i in report['issues']))
        self.assertFalse(any(i['code'] == 'MONTHLY_FALLBACK' for i in report['issues']))

    def test_trim_identity_and_decimal_numeric_strings_and_zero_rate(self):
        monthly = self.compute([source_row(venue=' 上交所 ', allocated=' 10.125 ', bid=' 3.25 ', rate='0'),
                                source_row(venue='上交所', allocated='.875', bid='.75', rate='.02')])
        asset = monthly['assets']['赊销白条']
        bank = self.institution(monthly)
        self.assertEqual(asset['projectCount'], 1)
        self.assertEqual(asset['totalAllocatedSize'], 11)
        self.assertEqual(bank['monthlySizes'][0], 4)
        self.assertEqual(bank['participationCount'], 1)
        self.assertEqual(bank['months'][0][0]['venue'], '上交所')
        self.assertEqual(bank['months'][0][0]['rate'], 0)
        self.assertEqual(bank['allocationShare'], 1)
        self.assert_gate_pass(monthly)

    def test_mixed_preserves_components_denominator_bubble_and_sort(self):
        monthly = self.compute([source_row(project='一月无', rate=None, bidder=None, bid=None),
                                source_row(project='一月有', date='2026-01-22'),
                                source_row(project='四月有', date='2026-04-09', bidder='华夏基金',
                                           institution='华夏基金', bid=2, allocated=5)])
        asset = monthly['assets']['赊销白条']
        bank = self.institution(monthly)
        self.assertEqual(bank['monthlySizes'][0], 13)
        self.assertEqual(bank['monthlyBidSizes'][0], 3)
        self.assertEqual(bank['monthlyFallbackSizes'][0], 10)
        self.assertEqual(bank['monthlyTracks'][0], 'mixed')
        self.assertEqual(bank['participationCount'], 2)
        self.assertEqual(bank['participationRate'], 2 / 3)
        self.assertEqual(bank['allocationShare'], 20 / 25)
        self.assertEqual(asset['bubbleMaxSize'], 13)
        self.assertEqual([i['name'] for i in asset['institutions']], ['邮储银行', '华夏基金'])
        self.assert_gate_pass(monthly)

    def test_missing_numerator_and_denominator_unknown_with_known_partial(self):
        monthly = self.compute([source_row(institution='华夏基金', bidder='华夏基金', allocated=None),
                                source_row(project='二月有', date='2026-02-09', institution='华夏基金',
                                           bidder='华夏基金', allocated=2),
                                source_row(project='三月有', date='2026-03-09', allocated=10)])
        asset = monthly['assets']['赊销白条']
        fund = self.institution(monthly, '华夏基金')
        bank = self.institution(monthly)
        self.assertIsNone(asset['totalAllocatedSize'])
        self.assertEqual(asset['totalKnownAllocatedSize'], 12)
        self.assertFalse(asset['allocationComplete'])
        self.assertEqual(asset['unknownAllocationRows'], [2])
        self.assertIsNone(fund['allocatedSize'])
        self.assertEqual(fund['knownAllocatedSize'], 2)
        self.assertFalse(fund['allocationComplete'])
        self.assertEqual(bank['allocatedSize'], 10)
        self.assertTrue(bank['allocationComplete'])
        self.assertTrue(all(i['allocationShare'] is None for i in asset['institutions']))
        self.assertEqual(monthly['unknownAllocationRecords'][0]['row'], 2)
        self.assertIsNone(monthly['unknownAllocationRecords'][0]['size'])
        self.assert_gate_pass(monthly)

    def test_unknown_hidden_institution_poison_share_and_selfheld_missing_excluded(self):
        monthly = self.compute([source_row(institution='某信托', bidder=None, rate=None, bid=None, allocated=None),
                                source_row(project='二月有', date='2026-02-09')])
        self.assertEqual(monthly['assets']['赊销白条']['unknownAllocationRows'], [2])
        self.assertIsNone(self.institution(monthly)['allocationShare'])
        self.assert_gate_pass(monthly)
        monthly = self.compute([source_row(institution='京东自持', allocated=None),
                                source_row(project='二月有', date='2026-02-09')])
        self.assertEqual(monthly['unknownAllocationRecords'], [])
        self.assertEqual(monthly['assets']['赊销白条']['totalAllocatedSize'], 10)
        self.assertEqual(self.institution(monthly)['allocationShare'], 1)
        self.assert_gate_pass(monthly)

    def test_no_uv_and_explicit_zero_are_known_zero(self):
        monthly = self.compute([source_row(institution=None, allocated=None, bidder='华夏基金'),
                                source_row(bidder='某银行', institution='某银行', allocated=0),
                                source_row(project='二月有', date='2026-02-09')])
        for name in ('华夏基金', '某银行'):
            inst = self.institution(monthly, name)
            self.assertEqual(inst['allocatedSize'], 0)
            self.assertEqual(inst['allocationShare'], 0)
            self.assertTrue(inst['allocationComplete'])
        self.assert_gate_pass(monthly)

    def test_funds_and_explicit_categories_without_bare_group_guessing(self):
        rows = [source_row(project=name, institution=name, bidder=name) for name in
                ('平安', '华夏', '中银', '建信', '平安基金', '华夏基金', '中银基金', '建信基金',
                 '华泰保兴基金', '前海开源', '国投资管投行', '华泰资管（自营）')]
        monthly = self.compute(rows)
        for name in ('平安', '华夏', '中银', '建信'):
            self.assertEqual(self.institution(monthly, name)['group'], '其他')
        for name in ('平安基金', '华夏基金', '中银基金', '建信基金', '华泰保兴基金', '前海开源基金'):
            self.assertEqual(self.institution(monthly, name)['group'], '公募基金')
        for name in ('国投资管（投行）', '华泰资管（自营）'):
            self.assertEqual(self.institution(monthly, name)['group'], '券商')
        self.assert_gate_pass(monthly)

    def test_config_schema_unique_nonempty_and_fingerprint(self):
        original = json.loads(demo.FUND_RULES.read_text(encoding='utf-8'))
        self.assertEqual(len(original['publicFunds']), 16)
        variants = []
        for field, value in (('schema', 'wrong'), ('publicFunds', []),
                             ('publicFunds', original['publicFunds'] + [original['publicFunds'][0]]),
                             ('publicFundAliases', {'平安': '平安基金'}), ('source', {})):
            variant = copy.deepcopy(original)
            variant[field] = value
            variants.append(json.dumps(variant, ensure_ascii=False))
        variants.append('{"schema":"monthly-bid-classification/v1","schema":"wrong"}')
        path = self.directory / 'classification.json'
        for encoded in variants:
            path.write_text(encoded, encoding='utf-8')
            with mock.patch.object(demo, 'FUND_RULES', path), self.assertRaises(ValueError):
                demo.fund_classification()
        monthly = self.compute([source_row()])
        dependencies = monthly['provenance']['dependencies']
        self.assertIn('scripts/lab/monthly_bid_classification.json', dependencies)
        self.assertNotIn(original['source']['path'], dependencies)
        self.assertEqual(monthly['fundRulesSha256'], gate.file_sha256(demo.FUND_RULES))

    def test_source_and_classification_change_during_read_detected(self):
        self.write_source([source_row()])
        original = demo.qc.read_source

        def mutate_ledger(ledger):
            rows = original(ledger)
            book = load_workbook(ledger)
            book['总表']['V2'] = 11
            book.save(ledger)
            book.close()
            return rows

        with mock.patch.object(demo.qc, 'read_source', side_effect=mutate_ledger), self.assertRaisesRegex(ValueError, '变化'):
            demo.compute_monthly(self.ledger)
        self.write_source([source_row()])
        path = self.directory / 'classification.json'
        path.write_bytes(demo.FUND_RULES.read_bytes())

        def mutate_config(ledger):
            rows = original(ledger)
            value = json.loads(path.read_text(encoding='utf-8'))
            value['source']['note'] = '读后变化'
            path.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
            return rows

        with mock.patch.object(demo, 'FUND_RULES', path), mock.patch.object(demo.qc, 'read_source', side_effect=mutate_config), self.assertRaisesRegex(ValueError, '变化'):
            demo.compute_monthly(self.ledger)

    def test_component_css_inline_contract_banner_and_strict_no_output(self):
        self.write_source([source_row()])
        monthly = demo.compute_monthly(self.ledger)
        report = gate.verify_data(monthly, self.ledger, raise_on_fail=False)
        demo._attach_qc(monthly, report)
        data = demo.panel.compute_data(self.ledger)
        component = demo.render_component(data, monthly)
        self.assertNotIn('<!doctype', component)
        self.assertEqual(component.count('id="lb-monthly-data"'), 1)
        self.assertEqual(component.count('id="lb-data"'), 1)
        self.assertIn((LAB / 'monthly_bid_institutions_demo.js').read_text(encoding='utf-8'), component)
        self.assertIn((LAB / 'monthly_bid_institutions_demo.css').read_text(encoding='utf-8'), demo.component_css())
        self.assertEqual(gate.verify_html(demo.build_document(data, monthly), self.ledger, raise_on_fail=False)['status'], 'PASS')
        fail = {'status': 'FAIL', 'issues': [{'level': 'FAIL', 'code': 'SOURCE_CONFLICT',
                                            'message': '金额冲突 <源表> & 待确认'}]}
        demo._attach_qc(monthly, fail)
        content = demo.build_document(data, monthly)
        self.assertIn('data-qc-status="FAIL"', content)
        self.assertIn('金额冲突 &lt;源表&gt; &amp; 待确认', content)
        self.assertTrue(monthly['auditQc']['releaseBlocked'])
        with self.assertRaises(ValueError):
            demo.render_component(data, monthly, preview=False)
        strict_output = self.directory / 'strict.html'
        preview_output = self.directory / 'preview.html'
        with mock.patch.object(gate, 'verify_data', return_value=fail), mock.patch.object(gate, 'verify_html', return_value=fail):
            with redirect_stdout(io.StringIO()) as stream:
                self.assertEqual(demo.main(['--ledger', str(self.ledger), '--output', str(strict_output), '--strict']), 1)
            self.assertFalse(strict_output.exists())
            self.assertEqual(json.loads(stream.getvalue())['auditQc']['status'], 'FAIL')
            with redirect_stdout(io.StringIO()):
                self.assertEqual(demo.main(['--ledger', str(self.ledger), '--output', str(preview_output)]), 0)
            self.assertIn('data-qc-status="FAIL"', preview_output.read_text(encoding='utf-8'))

    def test_real_source_manual_totals_after_user_correction_pass(self):
        ledger = LAB.parents[1] / 'deliverables/ledger/03_final/2026年ABS发行台账-1016-定稿.xlsx'
        source_hash = gate.file_sha256(ledger)
        monthly = demo.compute_monthly(ledger)
        totals = {'赊销白条': 326.475, '信托金条': 153.8, '信托企业主贷': 104.74,
                  '小贷企业主贷': 118, '京企贷': 18, '保理': 140, '动产质押': 20,
                  '白条取现': 30, '金采': 113, '汽车融资租赁': 20.61, '金条白取混发': 143.0}
        self.assertEqual({name: a['totalAllocatedSize'] for name, a in monthly['assets'].items()}, totals)
        self.assertEqual(sum(Decimal(str(a['totalAllocatedSize'])) for a in monthly['assets'].values()), Decimal('1187.625'))
        self.assertEqual(sum(a['projectCount'] for a in monthly['assets'].values()), 121)
        self.assertEqual(sum(a['bidCount'] for a in monthly['assets'].values()), 1997)
        self.assertEqual(len(monthly['fallbackRows']), 373)
        self.assertEqual(len(monthly['allocationRecords']), 1886)
        self.assertEqual(monthly['unknownAllocationRecords'], [])
        self.assertTrue(all(a['allocationComplete'] for a in monthly['assets'].values()))
        report = gate.verify_data(monthly, ledger, raise_on_fail=False)
        self.assertEqual(report['status'], 'PASS', report['issues'])
        self.assertEqual(report['releaseStatus'], 'PASS')
        self.assertFalse(any(i['level'] == 'FAIL' for i in report['issues']))
        source_report = gate.verify_source(ledger, raise_on_fail=False)
        self.assertEqual(source_report['status'], 'PASS', source_report['issues'])
        self.assertFalse(any(i['level'] == 'FAIL' for i in source_report['issues']))
        self.assertEqual(gate.file_sha256(ledger), source_hash)


if __name__ == '__main__':
    unittest.main()
