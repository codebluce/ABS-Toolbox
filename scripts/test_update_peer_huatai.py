"""华泰增量主体边界与原始表解析验证。"""
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
import openpyxl
from update_peer_huatai import huatai_projects, asset_labels, PARTIES, decimal
from peer_issuance_panel import compare_weekly_snapshot

class HuataiUpdateTest(unittest.TestCase):
    def test_verified_rename_is_not_addition_and_fact_changes_still_block(self):
        old={'product':'华泰资管诚昕3号4期资产支持专项计划','week':'2026.10.5-2026.10.11',
             'product_type':'ABS','venue':'上海','originator':'华能信托','base_asset':'微众银行微粒贷',
             'asset_type':'消金提现类','amount':Decimal('10'),'date':date(2026,10,9),'aaa':Decimal('.0157'),'term':Decimal('1.04')}
        new={**old,'product':'华泰资管诚新3号4期资产支持专项计划'}
        summary,qc=compare_weekly_snapshot([old],[new])
        self.assertEqual(summary['additions'],[])
        self.assertEqual(summary['deletions'],[])
        self.assertFalse(any(x['level']=='FAIL' for x in qc))
        _,qc=compare_weekly_snapshot([old],[{**new,'amount':Decimal('11')}])
        self.assertTrue(any(x['code']=='HISTORICAL_REVISION' and x['level']=='FAIL' for x in qc))

    def test_scope_and_jd_assets(self):
        self.assertNotIn('宁银消金',PARTIES)
        self.assertNotIn('顺丰小贷',PARTIES)
        self.assertEqual(asset_labels({'party':'京东','asset':'动产质押'}),('京东动产质押','小微toB'))
        self.assertEqual(asset_labels({'party':'蚂蚁','asset':'借呗'}),('蚂蚁借呗','消金提现类'))

    def test_huatai_multi_tranche_is_one_project(self):
        with tempfile.TemporaryDirectory() as temp:
            p=Path(temp)/'h.xlsx';w=openpyxl.Workbook();w.remove(w.active)
            s=w.create_sheet('交易所-2026年');s.append(['设立时间','专项计划名称','交易所','规模','原始权益人','管理人','分层','期限','规模','评级','利率','资产方','资产类型'])
            s.append(['2026/10/23','项目1','上海',10,'信托','资管','优先A',1,9,' AAAsf ','1,65','京东','白条'])
            s.append([None,None,None,None,None,None,'次级',2,1,'-',None,None,None])
            b=w.create_sheet('交易商协会-2026年');b.append(['设立时间','项目','规模','发起','发行','分层','期限','规模','评级','利率','资产方','资产'])
            w.save(p);w.close()
            rows=huatai_projects(p)
            self.assertEqual(len(rows),1)
            self.assertEqual(rows[0]['established'],date(2026,10,23))
            self.assertEqual(rows[0]['amount'],Decimal('10'))
            self.assertEqual(rows[0]['tiers'][0]['rating'],'AAAsf')
            self.assertEqual(rows[0]['tiers'][0]['rate'],Decimal('1.65'))
            self.assertEqual(len(rows[0]['tiers']),2)

    def test_invalid_number_is_not_silently_zero(self):
        with self.assertRaises(Exception):decimal('尚未披露')
        self.assertIsNone(decimal('-'))

if __name__=='__main__':unittest.main()
