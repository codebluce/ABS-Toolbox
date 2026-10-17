"""在J/Q和逐笔配售均支持金额时，修复总表R引用并写入可核验缓存。

仅处理简单J*Q乘式或数值常量，不运行通用Excel公式计算器。
历史文件不回写；原始ZIP中只有目标工作表R单元格被修改。
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from copy import copy
from datetime import datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import tempfile
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook

EPSILON = Decimal('0.0001')
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def number(value):
    if isinstance(value, bool) or value is None:
        raise ValueError('金额/比例缺失或类型错误')
    try:
        result = Decimal(str(value))
    except InvalidOperation:
        raise ValueError('金额/比例不是数值') from None
    if not result.is_finite() or result < 0:
        raise ValueError('金额/比例须为非负有限值')
    return result


def inspect(ledger):
    """公式修改依据按完整项目档位核定，含自持的V逐笔保留。"""
    book = load_workbook(ledger, data_only=False)
    cached = load_workbook(ledger, data_only=True)
    try:
        sheet = book['总表']
        headers = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}
        required = {'规模': 10, '分层占比': 17, '对应金额（亿）': 18,
                    '认购机构': 21, '认购份额': 22}
        if any(headers.get(name) != column for name, column in required.items()):
            raise ValueError('只支持已核定的标准总表J/Q/R/U/V列模板')
        project_cols = [headers[n] for n in ('资产类型', '项目名称', '发行场所', '簿记时间')]
        project_rows = defaultdict(list)
        for row in range(2, sheet.max_row + 1):
            key = tuple(sheet.cell(row, column).value for column in project_cols)
            if all(value is not None for value in key):
                project_rows[key].append(row)
            elif sheet.cell(row, 18).data_type == 'f':
                raise ValueError(f'总表R{row}项目身份不完整')
        changes, evidence = [], []
        for project, rows in project_rows.items():
            if not any(sheet.cell(row, 18).data_type == 'f' for row in rows):
                continue
            sizes = {number(sheet.cell(row, 10).value) for row in rows}
            if len(sizes) != 1 or not next(iter(sizes)):
                raise ValueError(f'{project[1]}项目规模不唯一或为0')
            size = next(iter(sizes))
            tiers = defaultdict(list)
            for row in rows:
                layer = sheet.cell(row, headers['分层情况']).value
                if not layer:
                    raise ValueError(f'{project[1]}缺少档位')
                tiers[layer].append(row)
            ratio_total = Decimal(0)
            project_total = Decimal(0)
            for layer, tier_rows in tiers.items():
                ratios = {number(sheet.cell(row, 17).value) for row in tier_rows}
                if len(ratios) != 1:
                    raise ValueError(f'{project[1]}/{layer}比例不唯一')
                ratio = next(iter(ratios))
                if ratio > 1:
                    raise ValueError(f'{project[1]}/{layer}比例超过100%')
                expected = size * ratio
                allocated = Decimal(0)
                for row in tier_rows:
                    institution, amount = sheet.cell(row, 21).value, sheet.cell(row, 22).value
                    if institution in (None, '', '-', '认购机构') and amount in (None, '', '-'):
                        continue  # 独立申购行，无配售记录，不算未知配售。
                    if institution in (None, '', '-', '认购机构') or amount in (None, '', '-'):
                        raise ValueError(f'总表U{row}/V{row}配售字段不完整，不能自动确认R')
                    allocated += number(amount)
                if abs(allocated - expected) > EPSILON:
                    raise ValueError(f'{project[1]}/{layer} J×Q={expected}与配售合计{allocated}有争议')
                ratio_total += ratio
                project_total += expected
                anchor = tier_rows[0]
                evidence.append({'project': str(project[1]), 'asset': str(project[0]),
                                 'date': str(project[3]), 'venue': str(project[2]), 'layer': str(layer),
                                 'anchor': anchor, 'rows': tier_rows, 'J': str(size), 'Q': str(ratio),
                                 'expectedR': str(expected), 'allocationSumV': str(allocated)})
                for row in tier_rows:
                    cell = sheet.cell(row, 18)
                    if cell.data_type != 'f':
                        if abs(number(cell.value) - expected) > EPSILON:
                            raise ValueError(f'{cell.coordinate}现存金额有争议，不能覆盖')
                        continue
                    formula = cell.value
                    constant = re.fullmatch(r'=\s*([0-9]+(?:\.[0-9]+)?)\s*', formula)
                    product = re.fullmatch(r'=\s*\$?([JQ])\$?(\d+)\s*\*\s*\$?([JQ])\$?(\d+)\s*', formula)
                    if constant:
                        if abs(number(constant[1]) - expected) > EPSILON:
                            raise ValueError(f'{cell.coordinate}常量金额与J/Q/V不一致')
                        fixed = formula
                    elif product and {product[1], product[3]} == {'J', 'Q'}:
                        fixed = f'=$J${anchor}*$Q${anchor}'
                    else:
                        raise ValueError(f'{cell.coordinate}不是受支持的简单公式，需人工核对')
                    changes.append({'cell': cell.coordinate, 'project': str(project[1]), 'layer': str(layer),
                                    'beforeFormula': formula, 'afterFormula': fixed,
                                    'beforeCache': cached['总表'][cell.coordinate].value,
                                    'afterCache': str(expected)})
            if abs(ratio_total - 1) > EPSILON or abs(project_total - size) > EPSILON:
                raise ValueError(f'{project[1]}全部档位与项目规模不勾稽')
        return changes, evidence
    finally:
        book.close()
        cached.close()


def repair(ledger, *, backup, report, expected_sha):
    ledger, backup, report = map(Path, (ledger, backup, report))
    before = digest(ledger)
    if before != expected_sha:
        raise ValueError('源台账指纹已变，必须重新核定')
    if backup.exists():
        raise FileExistsError('备份路径已存在')
    changes, evidence = inspect(ledger)
    backup.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(ledger, backup)
    with zipfile.ZipFile(ledger) as source:
        infos = [copy(info) for info in source.infolist()]
        originals = {info.filename: source.read(info.filename) for info in infos}
    workbook = ET.fromstring(originals['xl/workbook.xml'])
    relationships = ET.fromstring(originals['xl/_rels/workbook.xml.rels'])
    targets = {item.attrib['Id']: item.attrib['Target'] for item in relationships}
    sheet = next(item for item in workbook.find('m:sheets', NS) if item.attrib['name'] == '总表')
    target = targets[sheet.attrib['{' + NS['r'] + '}id']]
    entry = target.lstrip('/') if target.startswith('/') else 'xl/' + target
    xml = originals[entry].decode('utf-8')
    for change in changes:
        pattern = r'(<c\b[^>]*\br="' + change['cell'] + r'"[^>]*>)(.*?)(</c>)'
        matches = list(re.finditer(pattern, xml, re.S))
        if len(matches) != 1:
            raise ValueError('目标R单元格XML不唯一')
        inner = matches[0][2]
        formula_match = re.search(r'<f(?:\s[^>]*)?>(.*?)</f>', inner, re.S)
        if not formula_match or formula_match[1] != change['beforeFormula'][1:]:
            raise ValueError('XML公式与读取快照不一致')
        inner = re.sub(r'<f(?:\s[^>]*)?>.*?</f>', '<f>' + change['afterFormula'][1:] + '</f>', inner, count=1, flags=re.S)
        inner = re.sub(r'<v(?:\s[^>]*)?\s*/>|<v(?:\s[^>]*)?>.*?</v>', '', inner, flags=re.S)
        inner += '<v>' + change['afterCache'] + '</v>'
        xml = re.sub(pattern, lambda match: match[1] + inner + match[3], xml, count=1, flags=re.S)
    descriptor, temporary = tempfile.mkstemp(suffix='.xlsx', dir=ledger.parent)
    os.close(descriptor)
    try:
        with zipfile.ZipFile(temporary, 'w') as output:
            for info in infos:
                output.writestr(copy(info), xml.encode('utf-8') if info.filename == entry else originals[info.filename])
        with zipfile.ZipFile(temporary) as output:
            if output.testzip() is not None:
                raise ValueError('写入台账ZIP校验失败')
            for name, value in originals.items():
                if name != entry and output.read(name) != value:
                    raise ValueError('非目标工作表被修改')
        raw, values = load_workbook(temporary, data_only=False), load_workbook(temporary, data_only=True)
        try:
            for change in changes:
                if raw['总表'][change['cell']].value != change['afterFormula'] or number(values['总表'][change['cell']].value) != number(change['afterCache']):
                    raise ValueError('公式/缓存读回核对失败')
        finally:
            raw.close()
            values.close()
        if digest(ledger) != before:
            raise ValueError('修复期间原台账被其他进程改变')
        os.replace(temporary, ledger)
        result = {'checkedAt': datetime.now().astimezone().isoformat(timespec='seconds'),
                  'source': str(ledger.resolve()), 'backup': str(backup.resolve()),
                  'beforeSha256': before, 'afterSha256': digest(ledger),
                  'formulaChangedCells': sum(c['beforeFormula'] != c['afterFormula'] for c in changes),
                  'cacheWrittenCells': len(changes), 'changes': changes, 'evidence': evidence,
                  'cacheMethod': '已核定单档J×Q，V合计/各档Q总和/项目总额独立勾稽；仅简单乘式与已核定常量',
                  'untouched': '除总表目标R节点外全部ZIP项原始字节保留，J/Q/T/U/V/W/X/Y与历史文件未改'}
        report.parent.mkdir(parents=True, exist_ok=True)
        report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        return result
    finally:
        Path(temporary).unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--ledger', type=Path, required=True)
    parser.add_argument('--backup', type=Path, required=True)
    parser.add_argument('--report', type=Path, required=True)
    parser.add_argument('--expected-sha', required=True)
    args = parser.parse_args()
    result = repair(args.ledger, backup=args.backup, report=args.report, expected_sha=args.expected_sha)
    print(json.dumps({k: v for k, v in result.items() if k not in ('changes', 'evidence')}, ensure_ascii=False))


if __name__ == '__main__':
    main()
