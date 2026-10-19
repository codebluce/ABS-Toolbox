"""发行定价 > 最新簿记 Demo：从定稿台账读取各资产最近项目，不虚构报价。

T/U/V（最终配售）与 W/X/Y（申购报价）独立输出，不按 Excel 行号配对。
"""
from __future__ import annotations

import argparse
from datetime import datetime
import json
import math
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

import latest_bookkeeping_qc as qc

SCRIPT_DIR = Path(__file__).resolve().parent
BOOKKEEPING_CSS = (SCRIPT_DIR / 'latest_bookkeeping_panel.css').read_text(encoding='utf-8')
_BODY = (SCRIPT_DIR / 'latest_bookkeeping_panel_body.html').read_text(encoding='utf-8')
_JS = (SCRIPT_DIR / 'latest_bookkeeping_panel.js').read_text(encoding='utf-8')
RECENT_PERIODS = 3
ALLOWED_VENUES = frozenset({'上交所', '深交所', '银行间'})
ASSET_ORDER = ('赊销白条', '金条白取混发', '信托金条', '白条取现', '信托企业主贷',
               '小贷企业主贷', '保理', '金采', '京企贷', '动产质押', '汽车融资租赁')
ASSET_ALIASES = {'信托白取': '白条取现', '金条&白取混发': '金条白取混发'}
REQUIRED = {'规模', '发行场所', '资产类型', '项目名称', '簿记时间', '分层情况', '成本', '认购机构',
            '认购份额', '申购利率', '穿透机构', '申购规模', '期限'}


def text(value):
    if value is None or pd.isna(value):
        return ''
    value = str(value).strip()
    return '' if value in {'nan', 'None'} else value


def number(value):
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def booking_date(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            value = from_excel(value)
        except (ValueError, OverflowError):
            return pd.NaT
    parsed = pd.to_datetime(value, errors='coerce')
    return parsed.normalize() if pd.notna(parsed) else parsed


def find_latest_final(root: Path | None = None) -> Path:
    root = root or SCRIPT_DIR.parent / 'deliverables' / 'ledger' / '03_final'
    # 按文件标明的业务日期，而不是 mtime；备份、加工中及 field_sync 不参与。
    import re
    candidates = []
    for path in root.glob('*-定稿.xlsx'):
        match = re.fullmatch(r'(\d{4})年ABS发行台账-(\d{4})-定稿.xlsx', path.name)
        if match:
            try:
                stamp = datetime.strptime(''.join(match.groups()), '%Y%m%d')
            except ValueError:
                continue
            candidates.append((stamp, path))
    if not candidates:
        raise FileNotFoundError(f'未找到定稿台账: {root}')
    return max(candidates, key=lambda item: item[0])[1]


def compute_data(xlsx_path, *, preprocessed_path=None, source_name=None):
    # 不消费跨行填充的共享tmp：按原始总表读取，防止成本跨档或跨项目串入。
    # 参数仅为旧调用兼容保留；明确合并的字段按合并区域展开，明细不猜测补值。
    before = qc.file_sha256(xlsx_path)
    book = load_workbook(xlsx_path, data_only=True)
    try:
        if '总表' not in book.sheetnames:
            raise ValueError('簿记面板必须读取“总表”工作表')
        sheet = book['总表']
        headers = [text(cell.value) for cell in sheet[1]]
        missing = REQUIRED - set(headers)
        if missing or len(set(headers)) != len(headers):
            raise ValueError(f'簿记面板列缺失或重复: {sorted(missing)}')
        inherited = {}
        for area in sheet.merged_cells.ranges:
            for r in range(area.min_row, area.max_row + 1):
                for c in range(area.min_col, area.max_col + 1):
                    inherited[(r, c)] = sheet.cell(area.min_row, area.min_col).value
        records = []
        for r, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            row = {name: inherited.get((r, c), values[c - 1]) for c, name in enumerate(headers, 1)}
            complete_detail_header = all(text(row.get(name)) == name for name in ('成本','认购机构','认购份额','申购利率','穿透机构','申购规模'))
            if complete_detail_header or not any(value is not None for value in row.values()) or text(row.get('项目名称')) == '项目名称':
                continue
            if isinstance(row.get('簿记时间'), (int, float)):
                try:
                    row['簿记时间'] = from_excel(row['簿记时间'], book.epoch)
                except (TypeError, ValueError, OverflowError):
                    row['簿记时间'] = None
            records.append(row)
    finally:
        book.close()
    result = compute_frame(pd.DataFrame(records, columns=headers), (source_name or Path(xlsx_path).name))
    if before != qc.file_sha256(xlsx_path):
        raise ValueError('读取期间台账发生变化，请重试生成')
    result['provenance'] = {'schema':qc.SCHEMA_VERSION, 'sourceSha256':before,
                            'rulesSha256':qc.rules_sha256()}
    result['qc'] = qc.check_data(result, xlsx_path)
    return result


def compute_frame(frame, source_name):
    """保留台账原名和两套原始明细；不做机构过滤或年份推测纠错。"""
    frame = frame.copy()
    frame['_date'] = frame['簿记时间'].map(booking_date)
    frame['_asset'] = frame['资产类型'].map(text).map(lambda value: ASSET_ALIASES.get(value, value))
    frame['_project'] = frame['项目名称'].map(text)
    frame['_venue'] = frame['发行场所'].map(text)
    valid = frame[frame['_date'].notna() & frame['_asset'].ne('') & frame['_project'].ne('')]
    if valid.empty:
        raise ValueError('台账没有有效的资产、项目和簿记日期，无法选择最近一期')
    as_of = valid['_date'].max().strftime('%Y-%m-%d')
    # 先限定发行场所，再按资产选最新；非标/保登不遮蔽更早的合规项目。
    valid = valid[valid['发行场所'].map(text).isin(ALLOWED_VENUES)]
    if valid.empty:
        raise ValueError('台账没有发行场所为上交所、深交所或银行间的有效项目')
    assets = []
    for asset, group in valid.groupby('_asset', sort=False):
        ordered = group.sort_values('_date', ascending=False, kind='stable')
        # 以完整项目事件去重后取前三个，不能按日期把同日项目全部扩进来。
        recent_keys = list(dict.fromkeys(zip(ordered['_date'], ordered['_project'], ordered['_venue'])))[:RECENT_PERIODS]
        recent_dates = list(dict.fromkeys(key[0] for key in recent_keys))
        latest_date = recent_dates[0]
        keys = set(recent_keys)
        latest = ordered[[key in keys for key in zip(ordered['_date'], ordered['_project'], ordered['_venue'])]]
        projects = []
        for (project_date, project_name, project_venue), rows in latest.groupby(['_date', '_project', '_venue'], sort=False):
            tiers = {}
            for _, row in rows.iterrows():
                layer = text(row.get('分层情况')) or '未分层'
                if layer == '-':
                    layer = '未分层'
                tier = tiers.setdefault(layer, {'name': layer, 'costs': [], 'allocations': [], 'bids': []})
                cost = number(row.get('成本'))
                cost_value = cost if cost is not None else text(row.get('成本'))
                if cost_value != '' and cost_value not in tier['costs']:
                    tier['costs'].append(cost_value)
                investor = text(row.get('认购机构'))
                if investor and investor not in {'-', '认购机构'}:
                    tier['allocations'].append({'institution': investor, 'share': number(row.get('认购份额')),
                                                'rate': None, 'bidSize': None, 'cost': cost_value})
                bidder = text(row.get('穿透机构'))
                if bidder and bidder not in {'-', '穿透机构'}:
                    quote = number(row.get('申购利率'))
                    tier['bids'].append({'institution': bidder, 'share': None,
                                         'rate': quote if quote is not None else text(row.get('申购利率')),
                                         'bidSize': number(row.get('申购规模'))})
            first = rows.iloc[0]
            projects.append({'name': project_name, 'date': project_date.strftime('%Y-%m-%d'),
                             'size': number(first.get('规模')), 'manager': text(first.get('计划管理人')),
                             'venue': text(first.get('发行场所')), 'tenor': text(first.get('期限')),
                             'tiers': list(tiers.values())})
        assets.append({'name': asset, 'date': latest_date.strftime('%Y-%m-%d'),
                       'dates': [stamp.strftime('%Y-%m-%d') for stamp in recent_dates], 'projects': projects})
    order = {name: index for index, name in enumerate(ASSET_ORDER)}
    # 未约定类型保持源表首次出现的顺序，排在指定资产之后。
    assets.sort(key=lambda item: order.get(item['name'], len(order)))
    return {'source': source_name, 'asOf': as_of,
            'assets': assets, 'invalidDateRows': int(frame['_date'].isna().sum())}


def render_body(data, *, preview=False):
    encoded = json.dumps(data, ensure_ascii=False, separators=(',', ':'), allow_nan=False)
    for value, escaped in [('&', '\\u0026'), ('<', '\\u003c'), ('>', '\\u003e'),
                           ('\u2028', '\\u2028'), ('\u2029', '\\u2029')]:
        encoded = encoded.replace(value, escaped)
    body = _BODY
    if not preview:
        body = body.replace('DEMO · 真实台账预览', '定稿台账 · 自动校验')
        body = body.replace('Demo 仅供本地评审 · 数据来自定稿台账，不实时更新 · 不代表实时发行情况', '数据来自本地定稿台账 · 随台账生成更新 · 不代表实时发行情况')
    return (f'<div id="latest-bookkeeping">{body}'
            f'<script id="lb-data" type="application/json">{encoded}</script></div>'
            f'<script>{_JS}</script>')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--ledger', type=Path, default=None)
    parser.add_argument('--base', type=Path, default=None, help='已有综合看板，Demo保留全部其他面板')
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    ledger = args.ledger or find_latest_final()
    body = render_body(compute_data(ledger), preview=True)
    if args.base:
        document = args.base.read_text(encoding='utf-8')
        anchor = '<div class="sub-tabs-pane" data-module="pricing">'
        if anchor not in document or 'id="latest-bookkeeping"' in document:
            raise ValueError('基础看板缺少发行定价模块或已存在最新簿记面板')
        button = '<button class="sub-tab-button" data-sub="bookkeeping" onclick="selectSub(\'bookkeeping\')">最新簿记</button>'
        document = document.replace(anchor, anchor + button, 1)
        document = document.replace('</style>', BOOKKEEPING_CSS + '\n</style>', 1)
        document = document.replace('<div class="panel-container">',
            '<div class="panel-container"><div class="panel" data-module="pricing" data-sub="bookkeeping">' + body + '</div>', 1)
    else:
        from gen_integrated_dashboard import build_integrated_html
        document = build_integrated_html([('pricing', 'bookkeeping', body)], BOOKKEEPING_CSS)
    # 仅预览默认进入新面板，不影响基础看板的导航默认值。
    document = document.replace('<body>', '<body class="lb-demo-preview">', 1)
    document = document.replace('</style>', '\n@media(max-width:700px){body.lb-demo-preview .top-tabs,body.lb-demo-preview .sub-tabs-pane.active{overflow-x:auto;flex-wrap:nowrap;}body.lb-demo-preview .tab-button,body.lb-demo-preview .sub-tab-button{white-space:nowrap;flex-shrink:0;}body.lb-demo-preview .top-tabs{padding:0 10px;}body.lb-demo-preview .tab-button{padding:12px 14px;}body.lb-demo-preview .sub-tabs-wrapper{padding:8px 10px 0;}}\n</style>', 1)
    document = document.replace('</body>', "<script>document.addEventListener('DOMContentLoaded',()=>{selectModule('pricing');selectSub('bookkeeping');});</script></body>")
    if args.output.exists():
        raise FileExistsError(f'请使用新的 Demo 输出路径，避免覆盖已有产物: {args.output}')
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(document, encoding='utf-8')
    print(f'Demo: {args.output}\n台账: {ledger.name}')


if __name__ == '__main__':
    main()
