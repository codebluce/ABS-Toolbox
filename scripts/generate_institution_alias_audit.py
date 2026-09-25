"""从当前发行台账和次级快照重建机构名称映射审计清单（不修改源表）。"""
from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

import openpyxl

from abs_common import _INVESTOR_ALIAS_MAP, normalize_investor_name
from entity_alias import ENTITY_MERGE_MAP, BANK_NORM_MAP, normalize_entity, normalize_bank
from institution_identity import registry, canonical as canonical_institution

ROOT = Path(__file__).resolve().parents[1]
LEDGER = ROOT / 'deliverables/ledger/03_final/2026年ABS发行台账-0924-定稿.xlsx'
SNAPSHOT = ROOT / 'data/secondary_allocation_snapshot.json'
SECONDARY = ROOT / 'data/secondary_allocation_aliases.json'
PROFILE = ROOT / 'data/机构画像数据.json'
CREDIT = ROOT / 'deliverables/dashboards/04_reference/20260630额度盘点.xlsx'
CSV = ROOT / 'data/机构名称映射核对.csv'
DOC = ROOT / 'docs/机构名称映射审计清单.md'
FIELDS = {6: '计划管理人', 7: '联席承销商', 8: '托管行', 20: '认购机构', 23: '穿透机构'}


def audit_rows():
    counts = Counter()
    wb = openpyxl.load_workbook(LEDGER, read_only=True, data_only=True)
    try:
        for row in wb['总表'].iter_rows(min_row=2, max_col=24, values_only=True):
            for col, field in FIELDS.items():
                val = row[col]
                if val is None or not str(val).strip():
                    continue
                # 承销商为多个机构时分段计数，避免把一整串当作一家机构。
                for raw in (str(val).split('/') if col == 7 else [str(val)]):
                    if raw.strip() and raw.strip() != '-':
                        counts[('发行台账0924/总表', field, raw.strip())] += 1
    finally:
        wb.close()
    snapshot = json.loads(SNAPSHOT.read_text(encoding='utf-8'))
    for record in snapshot['records']:
        for field, key in [('次级投资人', 'investor'), ('次级管理人', 'manager')]:
            raw = str(record.get(key) or '').strip()
            if raw:
                counts[('次级速查快照', field, raw)] += 1
    for record in json.loads(PROFILE.read_text(encoding='utf-8')):
        name = str(record.get('name') or '').strip()
        if name:
            counts[('机构进展画像快照', '进展名称', name)] += 1
    credit_wb = openpyxl.load_workbook(CREDIT, read_only=True, data_only=True)
    try:
        for sheet in ('总授信', '非标授信'):
            for row in credit_wb[sheet].iter_rows(min_row=2, max_col=1, values_only=True):
                name = str(row[0] or '').strip()
                if name:
                    counts[(f'额度盘点/{sheet}', '授信机构', name)] += 1
        for sheet in ('全口径历史发行', '非标历史发行'):
            for row in credit_wb[sheet].iter_rows(min_row=2, max_col=21, values_only=True):
                name = str(row[20] or '').strip()
                if name:
                    counts[(f'额度盘点/{sheet}', '历史认购机构', name)] += 1
    finally:
        credit_wb.close()
    global_aliases = {a['alias']: a for a in registry()['aliases']}
    result = []
    for (source, field, raw), count in sorted(counts.items()):
        if field == '次级投资人':
            target = canonical_institution(raw)
            rule = '跨面板精确别名' if target != raw else '原名保留'
            evidence = 'data/institution_alias_registry.json' if target != raw else ''
        elif field in ('进展名称', '授信机构'):
            target = canonical_institution(raw)
            rule = '跨面板精确别名' if target != raw else '原名保留'
            evidence = 'data/institution_alias_registry.json' if target != raw else ''
        elif field == '托管行':
            target = normalize_bank(raw)
            rule = ('跨面板精确别名' if raw in global_aliases else '托管行分行归并') if target != raw else '原名保留'
            evidence = ('data/institution_alias_registry.json' if raw in global_aliases else
                        'scripts/entity_alias.py:BANK_NORM_MAP/银行名提取') if target != raw else ''
        elif field in ('计划管理人', '联席承销商', '次级管理人'):
            target = normalize_entity(raw)
            rule = '机构统计显式/跨面板别名' if target != raw else '原名保留'
            evidence = ('data/institution_alias_registry.json' if raw in global_aliases else
                        'scripts/entity_alias.py:ENTITY_MERGE_MAP') if target != raw else ''
        else:
            target = normalize_investor_name(raw)
            rule = ('跨面板精确别名' if raw in global_aliases else
                    '投资者旧版显式别名' if raw in _INVESTOR_ALIAS_MAP else
                    '投资者格式/后缀规则') if target != raw else '原名保留'
            evidence = ('data/institution_alias_registry.json' if raw in global_aliases else
                        'scripts/abs_common.py:normalize_investor_name') if target != raw else ''
        result.append((source, field, raw, target, count, rule, evidence))
    return result


def main():
    rows = audit_rows()
    CSV.parent.mkdir(parents=True, exist_ok=True)
    with CSV.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['来源', '字段', '原始名称', '统计名称', '出现次数', '归并类型', '规则来源'])
        writer.writerows(rows)
    global_rows = registry()['aliases']
    global_aliases = {row['alias']: row for row in global_rows}
    secondary = json.loads(SECONDARY.read_text(encoding='utf-8'))
    lines = [
        '# 机构名称映射审计清单', '',
        '数据版本：2026-09-24 定稿发行台账「总表」、次级速查/机构进展画像快照及 20260630 额度盘点（机构与历史认购列）；生成日期：2026-09-25。',
        '这里只定义报表**统计/检索口径**，不是法律主体一致性认定；不得据此推断授信主体同一。',
        '', '## 跨面板精确别名（仅登记项目归并）', '',
        '| 原始名称 | 统计名称 | 依据 / 用途 | 来源 |', '| --- | --- | --- | --- |',
    ]
    for item in global_rows:
        lines.append(f"| {item['alias']} | {item['canonical']} | {item['reason']} | {item['source']} |")
    lines.extend(['', '## 明确保留独立', '', '| A | B | 原因 |', '| --- | --- | --- |'])
    for item in registry()['distinct']:
        lines.append(f"| {item['left']} | {item['right']} | {item['reason']} |")
    lines.extend(['', '## 原次级速查别名的全局推广对照', '',
                  '原始明细不改写；次级速查与其他面板现在统一使用全局标准名。历史目标名如「中金自营」按台账既有规则展示为「中金公司（自营）」。',
                  '', '| 原始名称 | 原次级速查目标 | 全局统计名称 |', '| --- | --- | --- |'])
    for src, dst in secondary.items():
        lines.append(f'| {src} | {dst} | {global_aliases[src]["canonical"]} |')
    lines.extend(['', '## 待人工复核（暂不合并）', '',
                  '- `广发证券:自营/资管` 等进展快照中的组合标签：拆分依据未确认，不自动并入任何单家机构的投资或授信。',
                  '- 对授信精准匹配后出现的未关联历史机构逐条核查授信表；不得通过银行→理财子、证券→资管的名称包含关系补全。',
                  '', '## 旧规则及逐条核对', '',
                  '- 已确认 `广发证券（资管）→广发资管`；与 `广发证券` 保持独立。原次级 18 条全部纳入全局登记，历史展示名与统一名差异见上表。',
                  '- 投资者历史显式简称见 `scripts/abs_common.py:_INVESTOR_ALIAS_MAP`；机构统计旧简称见 `scripts/entity_alias.py:ENTITY_MERGE_MAP`。仅适用其各自使用场景。',
                  '- 托管行按分行→银行口径归并（`BANK_NORM_MAP`，另有“XX银行”提取）；不能据此合并证券或资管机构。',
                  '- 投资者旧规则会删除“有限公司”等法定后缀及“投行上报”尾注，并规范投行/自营/投顾括号。此类自动变化应按 CSV 中“投资者格式/后缀规则”逐条复核。',
                  '- 搜索候选允许对完整机构名称做关键字筛选，但点选后只查询该机构；问答只用完整名称、承销商 `/` 分段或登记的别名，不用前缀包含/去“资管”推断合并。',
                  '- 授信与机构画像的投资/授信关联仅按完整名称或登记别名精确关联；原机构画像手工部门组合表及银行→理财子、证券→资管、前缀兜底已停用。历史展示可能变为“暂无”，不得当作零认购。',
                  '- 原机构画像中带 `/` 的名称仍代表一条进展快照；分段各自按完整名称关联，不能合并成单一机构身份。带冒号的部门组合标签没有明确拆分依据时不自动关联。',
                  '', f'逐条实录见 `data/机构名称映射核对.csv`（{len(rows)} 条“来源×字段×原名”组合；含出现次数、标准名和规则来源）；未登记的相似名称保持原名，需人工审核后才能新增跨面板映射。',
                  '台账总表中每个非空原始名称按行计数；联席承销商按 `/` 分段；次级/进展快照、授信表及历史认购列按记录计数。不代表机构唯一数或认购规模。', ''])
    DOC.parent.mkdir(parents=True, exist_ok=True)
    DOC.write_text('\n'.join(lines), encoding='utf-8')
    print(f'{CSV}: {len(rows)} 行; {DOC}')


if __name__ == '__main__':
    main()
