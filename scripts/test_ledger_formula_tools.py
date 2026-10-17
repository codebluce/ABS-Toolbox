"""Synthetic workbooks only; run: .venv/bin/python scripts/test_ledger_formula_tools.py"""
import ast
from copy import copy
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
import ledger_formula_tools as tools
import increment_merge as merge
from abs_common import preprocess_xlsx_for_pandas, load_and_filter, COLNAMES


def identity(ws, row, project='P', tier='A', j=10, q=.8):
    for col, value in {4:'asset', 5:project, 6:'venue', 12:'2026-10-16',
                       10:j, 16:tier, 17:q}.items():
        ws.cell(row, col).value = value


class MovementTests(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = '总表'

    def test_translate_relative_and_absolute(self):
        self.assertEqual(tools.formula_translate('=Q2*J2', 'R2', 'R20'), '=Q20*J20')
        self.assertEqual(tools.formula_translate('=Q2*$J$2', 'R2', 'R20'), '=Q20*$J$2')
        identity(self.ws, 2, 'old')
        identity(self.ws, 20, 'new')
        with self.assertRaisesRegex(tools.LedgerFormulaError, 'crosses'):
            tools.validate_r_formula(self.ws, 20, '=Q20*$J$2')

    def test_same_project_header_j_different_tier_allowed(self):
        identity(self.ws, 2, tier='B')
        identity(self.ws, 20)
        self.assertEqual(tools.validate_r_formula(self.ws, 20, '=Q20*$J$2'), Decimal('8.0'))
        with self.assertRaises(tools.LedgerFormulaError):
            tools.validate_r_formula(self.ws, 20, '=Q2*J20')

    def test_insert_absolute_relative_cross_sheet_and_columns(self):
        self.ws['A1'] = '=J2+$J$2+J1+$J$1+SUM(J1:J3)'
        self.ws['R2'] = '=Q2*J2'
        other = self.wb.create_sheet('other')
        other['A1'] = "='总表'!$J$2+J2"
        self.ws['B1'] = '=other!J2'
        self.ws['C1'] = '=SUM(J:J)+SUM($2:$3)'
        tools.insert_rows_rebase(self.ws, 2, 3)
        self.assertEqual(self.ws['A1'].value, '=J5+$J$5+J1+$J$1+SUM(J1:J6)')
        self.assertEqual(self.ws['R5'].value, '=Q5*J5')
        self.assertEqual(other['A1'].value, "='总表'!$J$5+J2")
        self.assertEqual(self.ws['B1'].value, '=other!J2')
        self.assertEqual(self.ws['C1'].value, '=SUM(J:J)+SUM($5:$6)')

    def test_insert_unsupported_fails_before_mutation(self):
        self.ws['A2'] = 123
        self.wb.create_sheet('other')['A1'] = '=SUM(Table1[Amount])'
        with self.assertRaises(tools.LedgerFormulaError):
            tools.insert_rows_rebase(self.ws, 2)
        self.assertEqual(self.ws['A2'].value, 123)
        for formula in ('=[Book.xlsx]Sheet1!A2', '=SUM(Sheet1:Sheet3!A2)', '=namedRange'):
            with self.assertRaises(tools.LedgerFormulaError):
                tools.formula_translate(formula, 'R2', 'R20')

    def test_metadata_boundary_and_exact_details(self):
        identity(self.ws, 2)
        self.ws['R2'] = 8
        self.ws['S2'] = 'AAA'
        self.ws['T2'] = .02
        self.ws['U2'], self.ws['V2'] = 'x', 2
        self.ws['X3'], self.ws['Y3'] = ' bid ', '=1+1'
        self.ws['P4'] = 'B'
        self.ws['U4'] = 'next tier'
        identity(self.ws, 5, 'new', 'A', q=None)
        self.ws['U6'] = 'new detail'
        before = {(r,c):self.ws.cell(r,c).value for r in range(1,7) for c in range(21,26)}
        merge.unmerge_and_fill_processed(self.ws)
        self.assertEqual(self.ws['P3'].value, 'A')
        self.assertEqual(self.ws['Q3'].value, .8)
        self.assertIsNone(self.ws['R3'].value)
        self.assertIsNone(self.ws['Q4'].value)
        self.assertIsNone(self.ws['T4'].value)
        self.assertIsNone(self.ws['Q5'].value)
        self.assertIsNone(self.ws['R6'].value)
        self.assertEqual(before, {(r,c):self.ws.cell(r,c).value for r in range(1,7) for c in range(21,26)})

    def test_explicit_merge_r_fixed_anchor_and_detail_merges(self):
        identity(self.ws, 2)
        self.ws['R2'] = '=Q2*$J$2'
        self.ws['U2'] = ' x '
        self.ws.merge_cells('E2:E3')
        self.ws.merge_cells('R2:R3')
        self.ws.merge_cells('U2:U3')
        tools.expand_ledger_metadata(self.ws)
        self.assertEqual(self.ws['R3'].value, '=Q2*$J$2')
        self.assertEqual(tools.validate_r_formula(self.ws, 3), Decimal('8.0'))
        self.assertIsNone(self.ws['U3'].value)
        self.assertIn('U2:U3', {str(x) for x in self.ws.merged_cells.ranges})

    def test_first_data_row_partial_detail_header(self):
        self.ws.append(COLNAMES)
        self.ws['T2'] = '成本'
        identity(self.ws, 3)
        self.assertEqual(tools.first_data_row(self.ws), 3)
        self.assertEqual(merge.get_all_projects(self.ws)['P']['start'], 3)
        identity(self.ws, 2, 'real')
        self.assertEqual(tools.first_data_row(self.ws), 2)
        self.assertIn('real', merge.get_all_projects(self.ws))

    def test_j_sync_literal_conflict_atomic_formula_safe(self):
        identity(self.ws, 2)
        self.ws['R2'] = 8
        source = self.wb.create_sheet('raw')
        identity(source, 2, j=20)
        source['B2'] = 'new month'
        projects = {'P': {'start':2,'end':2}}
        with self.assertRaisesRegex(tools.LedgerFormulaError, 'conflicts'):
            merge.sync_existing_project_fields(self.ws, source, projects, projects)
        self.assertEqual(self.ws['J2'].value, 10)
        self.assertIsNone(self.ws['B2'].value)
        self.ws['R2'] = '=Q2*J2'
        merge.sync_existing_project_fields(self.ws, source, projects, projects)
        self.assertEqual(self.ws['J2'].value, 20)
        self.assertEqual(self.ws['R2'].value, '=Q2*J2')
        self.assertEqual(tools.validate_r_formula(self.ws, 2), Decimal('16.0'))

    def test_invalid_numbers_and_unknown_formula(self):
        identity(self.ws, 2)
        for q in (None, True, 'unknown', float('inf'), -.1, 1.2):
            self.ws['Q2'] = q
            with self.assertRaises(tools.LedgerFormulaError):
                tools.validate_r_formula(self.ws, 2, '=Q2*J2')
        for formula in ('=8.1', '=V2', '=J2*J2', '=other!J2*Q2'):
            with self.assertRaises(tools.LedgerFormulaError):
                tools.validate_r_formula(self.ws, 2, formula)

    def test_preprocess_cache_fallback_and_wrong_ref_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'synthetic.xlsx'
            self.ws.append(COLNAMES)
            identity(self.ws, 2)
            self.ws['R2'] = '=J2*Q2'
            self.ws['V2'] = '=1+1'
            self.ws['Y2'] = 3
            self.wb.save(path)
            # Inject an original Excel formula cache, including detail V.
            with ZipFile(path) as archive:
                files = {n:archive.read(n) for n in archive.namelist()}
            xml = files['xl/worksheets/sheet1.xml'].decode()
            xml = xml.replace('<f>1+1</f><v />', '<f>1+1</f><v>2</v>')
            files['xl/worksheets/sheet1.xml'] = xml.encode()
            with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
                for n,data in files.items(): archive.writestr(n,data)
            original = path.read_bytes()
            tmp = preprocess_xlsx_for_pandas(path)
            try:
                result = load_workbook(tmp, data_only=True)
                self.assertEqual(result.active['R2'].value, 8)
                self.assertEqual(result.active['V2'].value, 2)
                self.assertEqual(result.active['Y2'].value, 3)
                result.close()
                df, _, _, _ = load_and_filter(path, preprocessed_path=tmp, fix_year=False)
                self.assertEqual(len(df), 1)  # real row 2 is retained
            finally:
                Path(tmp).unlink()
            self.assertEqual(original, path.read_bytes())
            identity(self.ws, 3, 'wrong')
            self.ws['R2'] = '=J3*Q3'
            self.wb.save(path)
            with self.assertRaisesRegex(tools.LedgerFormulaError, 'crosses'):
                preprocess_xlsx_for_pandas(path)

    def test_constant_r_original_cache_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'constant.xlsx'
            self.ws.append(COLNAMES)
            identity(self.ws, 2, j=20)
            self.ws['R2'] = '=8.1'
            self.wb.save(path)
            with ZipFile(path) as archive:
                files = {n:archive.read(n) for n in archive.namelist()}
            files['xl/worksheets/sheet1.xml'] = files['xl/worksheets/sheet1.xml'].replace(
                b'<f>8.1</f><v />', b'<f>8.1</f><v>8.1</v>')
            with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
                for n, data in files.items(): archive.writestr(n, data)
            tmp = preprocess_xlsx_for_pandas(path)
            try:
                wb = load_workbook(tmp, data_only=True)
                self.assertEqual(wb.active['R2'].value, 8.1)  # JQ is 16, never substituted
                wb.close()
            finally:
                Path(tmp).unlink()

    def test_blank_and_full_project_key_boundaries(self):
        identity(self.ws, 2)
        self.ws['R2'] = 8
        self.ws['T2'] = .02
        identity(self.ws, 3, q=None)
        self.ws['L3'] = '2026-10-17'
        self.ws['U5'] = 'orphan after blank'
        tools.expand_ledger_metadata(self.ws)
        self.assertIsNone(self.ws['Q3'].value)
        self.assertIsNone(self.ws['T3'].value)
        self.assertIsNone(self.ws['E5'].value)

    def test_delete_relative_absolute_other_sheet_and_values(self):
        identity(self.ws, 2)
        identity(self.ws, 4)
        self.ws['R4'] = '=Q4*$J$2'
        self.ws['V4'] = 3
        other = self.wb.create_sheet('other')
        other['A1'] = "='总表'!$Q$4+Q4"
        self.ws['A1'] = '=other!J4+J2'
        tools.delete_rows_rebase(self.ws, 3)
        self.assertEqual(self.ws['R3'].value, '=Q3*$J$2')
        self.assertEqual(other['A1'].value, "='总表'!$Q$3+Q4")
        self.assertEqual(self.ws['A1'].value, '=other!J4+J2')
        self.assertEqual(self.ws['V3'].value, 3)
        self.assertEqual(self.ws['J3'].value, 10)
        self.assertEqual(self.ws['Q3'].value, .8)
        self.assertEqual(tools.validate_r_formula(self.ws, 3), Decimal('8.0'))

    def test_delete_target_or_range_intersection_fails_before_mutation(self):
        self.ws['J3'] = 10
        self.ws['R4'] = '=$J$3*Q4'
        with self.assertRaisesRegex(tools.LedgerFormulaError, 'deleted target'):
            tools.delete_rows_rebase(self.ws, 3)
        self.assertEqual(self.ws['J3'].value, 10)
        self.assertEqual(self.ws['R4'].value, '=$J$3*Q4')
        self.ws['R4'] = '=SUM(J2:J4)'
        with self.assertRaisesRegex(tools.LedgerFormulaError, 'deleted target'):
            tools.delete_rows_rebase(self.ws, 3)
        self.assertEqual(self.ws['J3'].value, 10)

    def test_same_label_explicit_q_r_subsegments_reset_fill(self):
        identity(self.ws, 2, q=.75)
        self.ws['R2'] = 7.5
        self.ws['S2'] = 'AAA'
        self.ws['T2'] = .02
        identity(self.ws, 3, q=.25)
        self.ws['R3'] = 2.5
        self.ws['U4'] = 'detail'
        tools.expand_ledger_metadata(self.ws)
        self.assertEqual(self.ws['Q2'].value, .75)
        self.assertEqual(self.ws['Q3'].value, .25)
        self.assertEqual(self.ws['R3'].value, 2.5)
        self.assertIsNone(self.ws['S3'].value)
        self.assertIsNone(self.ws['T3'].value)
        self.assertEqual(self.ws['Q4'].value, .25)
        self.assertIsNone(self.ws['R4'].value)
        self.assertIsNone(self.ws['S4'].value)
        self.assertIsNone(self.ws['T4'].value)
        # R alone can mark another explicit subsegment, without carrying old Q.
        self.ws['R5'] = 1
        self.ws['U5'] = 'next subsegment'
        tools.expand_ledger_metadata(self.ws)
        self.assertIsNone(self.ws['Q5'].value)
        self.assertEqual(self.ws['R5'].value, 1)

    def test_stale_direct_formula_cache_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'stale.xlsx'
            self.ws.append(COLNAMES)
            identity(self.ws, 2)
            self.ws['R2'] = '=J2*Q2'
            self.wb.save(path)
            with ZipFile(path) as archive:
                files = {n:archive.read(n) for n in archive.namelist()}
            files['xl/worksheets/sheet1.xml'] = files['xl/worksheets/sheet1.xml'].replace(
                b'<f>J2*Q2</f><v />', b'<f>J2*Q2</f><v>7</v>')
            with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
                for n, data in files.items(): archive.writestr(n, data)
            with self.assertRaisesRegex(tools.LedgerFormulaError, 'stale R cache'):
                preprocess_xlsx_for_pandas(path)

    def test_actual_append_branch_translates_source_coordinates(self):
        identity(self.ws, 19, project='old')
        source = self.wb.create_sheet('raw')
        identity(source, 2, project='new')
        source['R2'] = '=Q2*$J$2'
        data = {c:{'value':source.cell(2,c).value, 'origin':source.cell(2,c).coordinate,
                   'font':copy(source.cell(2,c).font), 'border':copy(source.cell(2,c).border),
                   'alignment':copy(source.cell(2,c).alignment), 'number_format':'General'}
                for c in range(1,26)}
        tree = ast.parse(Path(merge.__file__).read_text())
        branch = next(n for n in ast.walk(tree) if isinstance(n, ast.For)
                      and isinstance(n.target, ast.Name) and n.target.id == 'proj_name'
                      and any(isinstance(child, ast.Name) and child.id == 'row_data' for child in ast.walk(n))
                      and any(isinstance(child, ast.Name) and child.id == 'current_row' for child in ast.walk(n)))
        env = dict(vars(merge), ws_a=self.ws, ws_b=source, increment_ordered=['new'], increment_rows={'new':[data]},
                   current_row=20, append_start=20)
        exec(compile(ast.Module(body=[branch], type_ignores=[]), merge.__file__, 'exec'), env)
        self.assertEqual(self.ws['R20'].value, '=Q20*$J$20')
        self.assertEqual(tools.validate_r_formula(self.ws,20), Decimal('8.0'))

    def test_case_insensitive_sheet_insert_delete_and_quoted_name(self):
        self.ws.title = "Ledger O'Brien"
        self.ws['J3'] = 10
        other = self.wb.create_sheet('Other')
        other['A1'] = "='ledger o''brien'!$J$3+J3"
        tools.insert_rows_rebase(self.ws, 3, 2)
        self.assertEqual(other['A1'].value, "='ledger o''brien'!$J$5+J3")
        with self.assertRaisesRegex(tools.LedgerFormulaError, 'deleted target'):
            tools.delete_rows_rebase(self.ws, 5)
        self.assertEqual(self.ws['J5'].value, 10)
        tools.delete_rows_rebase(self.ws, 3, 2)
        self.assertEqual(other['A1'].value, "='ledger o''brien'!$J$3+J3")

    def test_review_literal_ledger_casefold_counterexample(self):
        self.ws.title = 'Ledger'
        self.ws['J3'] = 10
        other = self.wb.create_sheet('Other')
        other['A1'] = "='ledger'!$J$3"
        with self.assertRaisesRegex(tools.LedgerFormulaError, 'deleted target'):
            tools.delete_rows_rebase(self.ws, 3)
        self.assertEqual(self.ws['J3'].value, 10)
        tools.insert_rows_rebase(self.ws, 3, 2)
        self.assertEqual(other['A1'].value, "='ledger'!$J$5")

    def test_project_relocation_absolute_ranges_and_external_refusal(self):
        mapping = {2:20, 3:21}
        self.assertEqual(tools.relocate_project_formula('=Q2*$J$2', mapping, 'raw', 'target'),
                         '=Q20*$J$20')
        self.assertEqual(tools.relocate_project_formula("=SUM('RAW'!$J$2:J3)", mapping, 'raw', "Target O'Brien"),
                         "=SUM('Target O''Brien'!$J$20:J21)")
        for formula in ('=J1', '=other!J2', '=SUM(J:J)', '=SUM(J2:J4)'):
            with self.assertRaises(tools.LedgerFormulaError):
                tools.relocate_project_formula(formula, mapping, 'raw', 'target')
        # Ordinary cell copy preserves absolute source coordinates as before.
        self.assertEqual(tools.formula_translate('=Q2*$J$2', 'R2', 'R20'), '=Q20*$J$2')

    def test_actual_insert_branch_uses_live_tier_anchor(self):
        identity(self.ws, 2)
        identity(self.ws, 3)
        identity(self.ws, 4, tier='B', q=.2)
        self.ws['R2'] = '=Q2*$J$2'
        self.ws['R3'] = '=Q2*$J$2'
        self.ws['R4'] = '=Q4*$J$2'
        self.ws['U3'], self.ws['V3'] = 'keep', 2
        tree = ast.parse(Path(merge.__file__).read_text())
        branch = next(n for n in ast.walk(tree) if isinstance(n, ast.For)
                      and isinstance(n.target, ast.Name) and n.target.id == 'anchor_row')
        module = ast.Module(body=[branch], type_ignores=[])
        env = dict(vars(merge), ws_a=self.ws,
                   inserts_by_anchor={3:[({'rate':.02,'name':'new bid','size':1}, {'_row':2}, 'A')]})
        exec(compile(module, merge.__file__, 'exec'), env)
        self.assertEqual(self.ws['R4'].value, '=Q2*$J$2')
        self.assertEqual(tools.validate_r_formula(self.ws, 4), Decimal('8.0'))
        self.assertEqual(self.ws['R5'].value, '=Q5*$J$2')
        self.assertEqual(self.ws['U3'].value, 'keep')
        self.assertIsNone(self.ws['U4'].value)
        self.assertEqual(self.ws['X4'].value, 'new bid')


if __name__ == '__main__':
    unittest.main(verbosity=2)
