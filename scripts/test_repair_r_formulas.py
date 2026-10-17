from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from repair_r_formulas import digest, inspect, repair


class RepairFormulaTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.ledger = self.root / 'ledger.xlsx'

    def workbook(self, *, allocation=8, ratio=.8, formula='=J2*Q2'):
        book = Workbook()
        sheet = book.active
        sheet.title = '总表'
        sheet.append(['序号','月份','类别','资产类型','项目名称','发行场所','计划管理人',
                      '联席承销商','托管行','规模','期限','簿记时间','国股CD','ALL-IN成本',
                      '认购倍数','分层情况','分层占比','对应金额（亿）','评级','成本',
                      '认购机构','认购份额','申购利率','穿透机构','申购规模'])
        for layer, q, v, expression in [('优先A',ratio,allocation,formula),('次级',1-ratio,2,'=J3*Q3')]:
            row = [None]*25
            for column,value in {4:'白条',5:'测试项目',6:'上交所',10:10,12:'2026-01-08',
                                 16:layer,17:q,18:expression,21:'测试机构',22:v,23:.02,24:'申购机构',25:3}.items():
                row[column-1]=value
            sheet.append(row)
        other = book.create_sheet('其他')
        other['A1']='其他内容保持'
        book.save(self.ledger)
        book.close()

    def test_wrong_reference_repaired_cache_written_raw_fields_preserved(self):
        self.workbook(formula='=J3*Q3')
        before = digest(self.ledger)
        backup, report = self.root/'backup.xlsx', self.root/'report.json'
        result=repair(self.ledger,backup=backup,report=report,expected_sha=before)
        self.assertEqual(result['formulaChangedCells'],2)
        self.assertEqual(result['cacheWrittenCells'],2)
        self.assertEqual(digest(backup),before)
        raw=load_workbook(self.ledger,data_only=False)
        values=load_workbook(self.ledger,data_only=True)
        try:
            self.assertEqual(raw['总表']['R2'].value,'=$J$2*$Q$2')
            self.assertEqual(values['总表']['R2'].value,8)
            self.assertEqual(values['总表']['R3'].value,2)
            self.assertEqual(values['总表']['V2'].value,8)
            self.assertEqual(values['总表']['Y2'].value,3)
            self.assertEqual(raw['其他']['A1'].value,'其他内容保持')
        finally:
            raw.close(); values.close()

    def test_disputed_values_wrong_fingerprint_and_unsupported_formula_reject(self):
        self.workbook(allocation=7)
        with self.assertRaises(ValueError): inspect(self.ledger)
        self.workbook(formula='=SUM(J2,Q2)')
        with self.assertRaises(ValueError): inspect(self.ledger)
        self.workbook()
        before=digest(self.ledger)
        with self.assertRaises(ValueError):
            repair(self.ledger,backup=self.root/'backup.xlsx',report=self.root/'report.json',expected_sha='0'*64)
        self.assertEqual(digest(self.ledger),before)
        self.assertFalse((self.root/'backup.xlsx').exists())


if __name__ == '__main__':
    unittest.main()
